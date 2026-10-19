"""Task-oriented review analysis UI: start, library, history."""
from __future__ import annotations

import json
import re
import tempfile
from pathlib import Path

from PyQt6.QtCore import Qt, QTimer, pyqtSignal
from PyQt6.QtGui import QDesktopServices, QPainter, QColor
from PyQt6.QtCore import QUrl
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTabWidget, QFormLayout, QLabel,
    QFileDialog, QMessageBox, QComboBox, QSpinBox, QTableWidget,
    QTableWidgetItem, QHeaderView, QDialog, QDialogButtonBox, QListWidget,
    QListWidgetItem, QProgressBar, QScrollArea, QApplication, QSizePolicy,
)

from core.widgets import Card, LogConsole, FileDropLineEdit, make_button, configure_combo, file_row, ActionBar
from core.theme import THEME
from .analysis_dialog import load_product_templates, check_template_fingerprint, template_fingerprint
from .analysis_jobs import AnalysisJobs, inspect_workbook, now
from .panel import ReviewCollectionPanel, _TaskThread
from .paths import EXPORT_DIR


STATUS_NAMES = {"queued": "等待分析", "running": "正在分析", "stopped": "已停止", "failed": "未完成", "done": "已完成"}


class InputTable(QTableWidget):
    files_dropped = pyqtSignal(list)

    def __init__(self):
        super().__init__(0, 4)
        self.setAcceptDrops(True)
        self.setHorizontalHeaderLabels(["评论文件", "分析范围", "品类", "操作"])
        self.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        for column, width in ((1, 165), (2, 180), (3, 80)):
            self.horizontalHeader().setSectionResizeMode(column, QHeaderView.ResizeMode.Fixed)
            self.setColumnWidth(column, width)
        self.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.verticalHeader().hide()
        self.setMinimumHeight(135)
        self.setMaximumHeight(220)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def paintEvent(self, event):
        super().paintEvent(event)
        if not self.rowCount():
            painter = QPainter(self.viewport())
            painter.setPen(QColor(THEME.tokens()["text2"]))
            painter.drawText(self.viewport().rect().adjusted(12, 8, -12, -8),
                             Qt.AlignmentFlag.AlignCenter | Qt.TextFlag.TextWordWrap,
                             "将评论 Excel 拖到这里，或点击上方“导入 Excel”\n支持多个文件，每个文件独立生成报告")

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        paths = [url.toLocalFile() for url in event.mimeData().urls() if url.isLocalFile()]
        self.files_dropped.emit(paths)
        event.acceptProposedAction()


class ModuleWidget(QWidget):
    def __init__(self, config_mgr, runner):
        super().__init__()
        self.setProperty("workspacePage", True)
        self.config_mgr, self.runner = config_mgr, runner
        self.jobs = AnalysisJobs(Path(config_mgr.get_path()).parent / "comment_analysis")
        self.inputs = []
        self._queue = []
        self._active = None
        self._worker = None
        self._preparing = False
        self._stopped = False
        self._report_path = None
        self._stop_file = Path(tempfile.gettempdir()) / "comment_stop.signal"
        self._init_ui()
        self._refresh_history()
        QApplication.instance().aboutToQuit.connect(self._shutdown)
        if self.jobs.warning:
            self.start_hint.setText(self.jobs.warning)

    @staticmethod
    def _plain(text=""):
        label = QLabel(text)
        label.setTextFormat(Qt.TextFormat.PlainText)
        label.setWordWrap(True)
        return label

    def _init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 16, 24, 20)
        settings = make_button("设置")
        settings.clicked.connect(self._show_settings)
        self.tabs = QTabWidget()
        self.tabs.setCornerWidget(settings)
        self.start_tab = self._build_start()
        self.review_collection = ReviewCollectionPanel(config_mgr=self.config_mgr, runner=self.runner, parent=self)
        self.review_collection.analysis_requested.connect(self._use_collected_review_file)
        self.library_tab = self.review_collection
        self.result_tab = self._build_history()
        self.tabs.addTab(self.start_tab, "开始分析")
        self.tabs.addTab(self.library_tab, "评论库")
        self.tabs.addTab(self.result_tab, "分析记录")
        self.tabs.currentChanged.connect(self._tab_changed)
        layout.addWidget(self.tabs, 1)
        self._build_settings()
        self._refresh_templates()

    def _build_start(self):
        container = QWidget()
        outer = QVBoxLayout(container)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(8)
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 12, 0, 0)
        data = Card("1  添加评论")
        data.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Maximum)
        actions = QHBoxLayout()
        self.add_btn = make_button("导入 Excel")
        self.add_btn.clicked.connect(self._browse_inputs)
        self.library_btn = make_button("从评论库选择")
        self.library_btn.clicked.connect(lambda: self.tabs.setCurrentWidget(self.library_tab))
        actions.addWidget(self.add_btn)
        actions.addWidget(self.library_btn)
        actions.addStretch()
        data.content_layout.addLayout(actions)
        self.input_table = InputTable()
        self.input_table.files_dropped.connect(self._add_inputs)
        data.content_layout.addWidget(self.input_table)
        layout.addWidget(data)

        category = Card("2  选择产品品类")
        category.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Maximum)
        row = QHBoxLayout()
        self.product_combo = configure_combo(QComboBox())
        self.product_combo.currentIndexChanged.connect(self._apply_product)
        row.addWidget(self.product_combo, 1)
        self.apply_category_btn = make_button("应用到全部文件")
        self.apply_category_btn.clicked.connect(self._apply_product)
        row.addWidget(self.apply_category_btn)
        category.content_layout.addLayout(row)
        self.dimension_hint = self._plain("选择品类后显示本次报告的分析维度。")
        category.content_layout.addWidget(self.dimension_hint)
        layout.addWidget(category)

        more_btn = make_button("更多选项", "ghost")
        more_btn.setCheckable(True)
        self.more_options = QWidget()
        form = QFormLayout(self.more_options)
        self.output_dir = FileDropLineEdit(want_dir=True, placeholder="默认保存在各原文件旁；评论库报告保存在“评论导出”文件夹")
        form.addRow("报告位置", file_row(self.output_dir, "dir", parent=self))
        self.mode_combo = configure_combo(QComboBox())
        self.mode_combo.addItem("生成完整洞察报告（推荐）", "full")
        self.mode_combo.addItem("仅提取评论标签", "extract")
        self.mode_combo.addItem("仅汇总已有标签文件", "summary")
        form.addRow("处理方式", self.mode_combo)
        more_btn.toggled.connect(self.more_options.setVisible)
        more_btn.toggled.connect(lambda checked: more_btn.setText("收起更多选项" if checked else "更多选项"))
        self.more_options.hide()
        layout.addWidget(more_btn, alignment=Qt.AlignmentFlag.AlignLeft)
        layout.addWidget(self.more_options)
        privacy = self._plain("原始文件保持不变，分析记录自动保存在本机。")
        privacy.setObjectName("FormHint")
        layout.addWidget(privacy)
        layout.addStretch()
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        scroll.setWidget(page)
        outer.addWidget(scroll, 1)
        self.start_hint = self._plain("")
        outer.addWidget(self.start_hint)
        self.start_actions = ActionBar()
        self.scope_hint = self.start_actions.summary
        self.start_btn = make_button("开始分析", "primary")
        self.start_btn.setMinimumWidth(150)
        self.start_btn.clicked.connect(self._start_analysis)
        self.start_actions.row.addWidget(self.start_btn)
        outer.addWidget(self.start_actions)
        return container

    def _build_history(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 12, 0, 0)
        self.history_hint = self._plain("分析记录会自动保存在本机。选择一条记录查看报告或继续任务。")
        layout.addWidget(self.history_hint)
        self.history_table = QTableWidget(0, 5)
        self.history_table.setHorizontalHeaderLabels(["商品 / 文件", "品类", "评论数", "创建时间", "状态"])
        self.history_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.history_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.history_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.history_table.setSelectionMode(QTableWidget.SelectionMode.SingleSelection)
        self.history_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.history_table.setMinimumHeight(145)
        self.history_table.setMaximumHeight(225)
        self.history_table.verticalHeader().hide()
        self.history_table.itemSelectionChanged.connect(self._select_history)
        layout.addWidget(self.history_table)
        self.task_status = self._plain("暂无分析任务。到“开始分析”添加评论即可生成第一份报告。")
        layout.addWidget(self.task_status)
        self.progress = QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.hide()
        layout.addWidget(self.progress)
        actions = QHBoxLayout()
        self.resume_btn = make_button("继续分析", "primary")
        self.resume_btn.clicked.connect(self._resume_selected)
        self.stop_btn = make_button("停止", "danger")
        self.stop_btn.clicked.connect(self._stop)
        self.open_report_btn = make_button("打开完整报告", "primary")
        self.open_report_btn.clicked.connect(self._open_report)
        self.folder_btn = make_button("打开所在文件夹")
        self.folder_btn.clicked.connect(self._open_folder)
        for button in (self.resume_btn, self.stop_btn, self.open_report_btn, self.folder_btn):
            actions.addWidget(button)
            button.hide()
        actions.addStretch()
        layout.addLayout(actions)
        self.result_context = self._plain("")
        layout.addWidget(self.result_context)
        self.result_table = QTableWidget(0, 5)
        self.result_table.setHorizontalHeaderLabels(["维度", "主要类别", "提及次数", "维度内占比", "涉及评论"])
        self.result_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.result_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.result_table.verticalHeader().hide()
        self.result_table.hide()
        layout.addWidget(self.result_table, 1)
        self.result_note = self._plain("每个重点维度展示前三类。占比按该维度提及次数计算；完整明细与原文证据见报告。")
        self.result_note.hide()
        layout.addWidget(self.result_note)
        log_btn = make_button("查看详细日志", "ghost")
        log_btn.setCheckable(True)
        self.log = LogConsole(max_height=190)
        self.log.text.document().setMaximumBlockCount(1200)
        self.log.hide()
        log_btn.toggled.connect(self.log.setVisible)
        log_btn.toggled.connect(lambda checked: log_btn.setText("收起详细日志" if checked else "查看详细日志"))
        layout.addWidget(log_btn, alignment=Qt.AlignmentFlag.AlignLeft)
        layout.addWidget(self.log)
        layout.addStretch()
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        scroll.setWidget(page)
        return scroll

    def _build_settings(self):
        self.settings_dialog = QDialog(self)
        self.settings_dialog.setWindowTitle("评论分析设置")
        self.settings_dialog.resize(790, 640)
        layout = QVBoxLayout(self.settings_dialog)
        tabs = QTabWidget()
        page = QWidget()
        content = QVBoxLayout(page)
        templates = Card("分析模板", "直接新增品类并补充分析维度。导入模板会替换当前配置。")
        self.cfg_status = self._plain()
        templates.content_layout.addWidget(self.cfg_status)
        row = QHBoxLayout()
        self.add_category_btn = make_button("＋ 新增品类", "primary")
        self.add_category_btn.clicked.connect(self._add_category)
        row.addWidget(self.add_category_btn)
        self.edit_prompt_btn = make_button("查看 / 修改提示词")
        self.edit_prompt_btn.clicked.connect(self._edit_prompts)
        row.addWidget(self.edit_prompt_btn)
        import_btn = make_button("导入模板")
        export_btn = make_button("导出模板")
        import_btn.clicked.connect(self._import_config)
        export_btn.clicked.connect(self._export_config)
        row.addWidget(import_btn)
        row.addWidget(export_btn)
        row.addStretch()
        templates.content_layout.addLayout(row)
        content.addWidget(templates)
        speed = Card("处理速度", "默认设置适合日常使用；遇到服务商限流时，可降低并发数。")
        form = QFormLayout()
        self.concurrency_spin = QSpinBox()
        self.concurrency_spin.setRange(1, 12)
        self.concurrency_spin.setValue(int(self.config_mgr.get("comment_concurrency", 8)))
        self.rpm_spin = QSpinBox()
        self.rpm_spin.setRange(0, 100000)
        self.rpm_spin.setSpecialValueText("不限制")
        self.rpm_spin.setValue(int(self.config_mgr.get("comment_rpm", 900)))
        self.tpm_spin = QSpinBox()
        self.tpm_spin.setRange(0, 10000000)
        self.tpm_spin.setSingleStep(10000)
        self.tpm_spin.setSpecialValueText("不限制")
        self.tpm_spin.setValue(int(self.config_mgr.get("comment_tpm", 90000)))
        form.addRow("同时处理的评论数", self.concurrency_spin)
        form.addRow("每分钟请求上限（RPM）", self.rpm_spin)
        form.addRow("每分钟 Token 上限（TPM）", self.tpm_spin)
        speed.content_layout.addLayout(form)
        content.addWidget(speed)
        content.addStretch()
        tabs.addTab(page, "模板与处理速度")
        tabs.addTab(self.review_collection.cloud_page, "数据与备份")
        layout.addWidget(tabs)
        close = make_button("完成", "primary")
        close.clicked.connect(self.settings_dialog.accept)
        layout.addWidget(close, alignment=Qt.AlignmentFlag.AlignRight)

    def _show_settings(self):
        self.settings_dialog.exec()
        try:
            self.config_mgr.update({"comment_concurrency": self.concurrency_spin.value(),
                                    "comment_rpm": self.rpm_spin.value(), "comment_tpm": self.tpm_spin.value()})
        except OSError as exc:
            QMessageBox.warning(self, "设置未保存", str(exc))

    def _tab_changed(self, index):
        if self.tabs.widget(index) is self.review_collection:
            self.review_collection.show_library()

    def _refresh_templates(self):
        self.templates = load_product_templates(self.config_mgr)
        selected = self.product_combo.currentData()
        self.product_combo.blockSignals(True)
        self.product_combo.clear()
        self.product_combo.addItem("请选择品类…", None)
        for key, entry in self.templates.items():
            self.product_combo.addItem(entry.get("display_name", key), key)
        self.product_combo.setCurrentIndex(max(0, self.product_combo.findData(selected)))
        self.product_combo.blockSignals(False)
        self.cfg_status.setText(f"已加载 {len(self.templates)} 个分析模板")
        options = self.review_collection.excel_options
        selected_import = options.category.currentData()
        options.configs = self.templates.copy()
        options.category.blockSignals(True)
        options.category.clear()
        for key, entry in self.templates.items():
            options.category.addItem(entry.get("display_name", key), key)
        options.category.setCurrentIndex(max(0, options.category.findData(selected_import)))
        options.category.blockSignals(False)
        self._render_inputs()

    def _add_category(self):
        from .category_editor import AddCategoryDialog
        dialog = AddCategoryDialog(self.config_mgr, self.settings_dialog)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self._refresh_templates()
            name = self.templates[dialog.created_key]["display_name"]
            self.cfg_status.setText(f"已添加“{name}”，共 {len(self.templates)} 个分析模板。可返回分析页选择。")

    def _edit_prompts(self):
        from .prompt_editor import PromptEditorDialog
        dialog = PromptEditorDialog(self.config_mgr, self.settings_dialog, self.product_combo.currentData())
        dialog.exec()
        if dialog.changed:
            self._refresh_templates()

    def _import_config(self):
        path, _ = QFileDialog.getOpenFileName(self, "导入分析模板", "", "JSON (*.json)")
        if not path:
            return
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8"))
            templates = data.get("products", data)
            if not isinstance(templates, dict) or not templates or any(not isinstance(v, dict) or not v.get("fields") for v in templates.values()):
                raise ValueError("模板须包含品类及 fields 分析字段。")
            # An EXE's bundled resource is temporary: save an override beside it.
            import sys
            destination = Path(sys.executable).parent / "product_configs.json" if getattr(sys, "frozen", False) else Path(self.config_mgr.resolve_config("product_configs.json"))
            destination.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            self._refresh_templates()
        except (OSError, ValueError, AttributeError) as exc:
            QMessageBox.warning(self, "无法导入", str(exc))

    def _export_config(self):
        path, _ = QFileDialog.getSaveFileName(self, "导出分析模板", "product_configs.json", "JSON (*.json)")
        if path:
            try:
                Path(path).write_text(json.dumps(self.templates, ensure_ascii=False, indent=2), encoding="utf-8")
            except OSError as exc:
                QMessageBox.warning(self, "无法导出", str(exc))

    def _browse_inputs(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "添加评论文件", "", "Excel (*.xlsx)")
        self._add_inputs(paths)

    def _background(self, target, done):
        if self._worker is not None:
            return False
        worker = _TaskThread(lambda log: target(), self)
        self._worker = worker
        worker.result.connect(done)
        worker.failed.connect(self._background_failed)
        worker.finished.connect(lambda: self._worker_finished(worker))
        worker.start()
        return True

    def _worker_finished(self, worker):
        if self._worker is worker:
            self._worker = None
        worker.deleteLater()
        self._set_input_enabled(not self._preparing and not self._active and not self._queue)

    def _background_failed(self, message):
        self._preparing = False
        self.start_hint.setText(message)
        self.tabs.setCurrentWidget(self.start_tab)

    def _add_inputs(self, paths):
        if self._worker or self._preparing or self._active or self._queue:
            self.start_hint.setText("请等待当前任务结束后再添加文件。")
            return
        existing = {item["path"] for item in self.inputs}
        paths = list(dict.fromkeys(str(Path(p).resolve()) for p in paths if p))
        paths = [p for p in paths if p not in existing]
        if not paths:
            return
        key = self.product_combo.currentData()
        self.start_hint.setText("正在识别工作表与评论数量…")
        self._set_input_enabled(False)
        def scan():
            valid, errors = [], []
            for path in paths:
                try:
                    sheets = inspect_workbook(path)
                    valid.append({"path": path, "sheets": sheets, "selected": [s["name"] for s in sheets], "product": key})
                except Exception as exc:
                    errors.append(f"{Path(path).name}：{exc}")
            return valid, errors
        def done(result):
            valid, errors = result
            self.inputs.extend(valid)
            self.start_hint.setText("\n".join(errors))
            self._render_inputs()
        self._background(scan, done)

    def _render_inputs(self):
        self.input_table.setRowCount(len(self.inputs))
        self.input_table.setColumnHidden(2, len(self.inputs) <= 1)
        self.input_table.setFixedHeight(min(230, max(100, 40 + 46 * len(self.inputs))))
        total = 0
        for row, item in enumerate(self.inputs):
            label = QTableWidgetItem(Path(item["path"]).name)
            label.setToolTip(item["path"])
            self.input_table.setItem(row, 0, label)
            count = sum(s["count"] for s in item["sheets"] if s["name"] in item["selected"])
            total += count
            scope = make_button(f"{count} 条 · 调整范围", "compact")
            scope.clicked.connect(lambda checked=False, item=item: self._edit_scope(item))
            self.input_table.setCellWidget(row, 1, scope)
            if len(self.inputs) > 1:
                combo = configure_combo(QComboBox())
                combo.addItem("请选择品类…", None)
                for key, entry in getattr(self, "templates", {}).items():
                    combo.addItem(entry.get("display_name", key), key)
                combo.setCurrentIndex(max(0, combo.findData(item.get("product"))))
                combo.currentIndexChanged.connect(lambda index, combo=combo, item=item: item.update(product=combo.currentData()))
                self.input_table.setCellWidget(row, 2, combo)
            else:
                self.input_table.removeCellWidget(row, 2)
            remove = make_button("移除", "ghost")
            remove.clicked.connect(lambda checked=False, item=item: self._remove_input(item))
            self.input_table.setCellWidget(row, 3, remove)
            self.input_table.setRowHeight(row, 46)
        self.scope_hint.setText(f"已添加 {len(self.inputs)} 个文件，共 {total} 条评论；每个文件独立生成报告。" if self.inputs else "尚未添加评论。可将 Excel 拖入上方区域。")
        self.start_btn.setEnabled(bool(self.inputs) and not (
            self._worker or self._preparing or self._active or self._queue))

    def _remove_input(self, item):
        self.inputs.remove(item)
        self._render_inputs()

    def _edit_scope(self, item):
        dialog = QDialog(self)
        dialog.setWindowTitle("选择要分析的工作表")
        dialog.resize(480, 340)
        layout = QVBoxLayout(dialog)
        layout.addWidget(self._plain("已自动识别包含评论正文的工作表；取消勾选可缩小分析范围。"))
        choices = QListWidget()
        for sheet in item["sheets"]:
            row = QListWidgetItem(f"{sheet['name']} · {sheet['count']} 条评论")
            row.setData(Qt.ItemDataRole.UserRole, sheet["name"])
            row.setFlags(row.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            row.setCheckState(Qt.CheckState.Checked if sheet["name"] in item["selected"] else Qt.CheckState.Unchecked)
            choices.addItem(row)
        layout.addWidget(choices)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("确定")
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText("取消")
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            item["selected"] = [choices.item(i).data(Qt.ItemDataRole.UserRole) for i in range(choices.count()) if choices.item(i).checkState() == Qt.CheckState.Checked]
            self._render_inputs()

    def _apply_product(self, *_args):
        key = self.product_combo.currentData()
        if key:
            for item in self.inputs:
                item["product"] = key
            entry = self.templates.get(key, {})
            dims = [str(field.get("excel_col", "")) for field in entry.get("fields", []) if field.get("excel_col") != "翻译后评论"]
            self.dimension_hint.setText("将分析：" + "、".join(dims))
        else:
            self.dimension_hint.setText("选择品类后显示本次报告的分析维度。")
        self._render_inputs()

    def _set_input_enabled(self, enabled):
        for widget in (self.add_btn, self.library_btn, self.input_table, self.product_combo, self.apply_category_btn, self.more_options, self.start_btn):
            widget.setEnabled(enabled)
        self.start_btn.setEnabled(enabled and bool(self.inputs))

    def _can_start(self):
        if self._active or self._queue or self._preparing or self._worker or self.runner.is_running() or self.runner.is_any_running() or self.review_collection._collection_running:
            self.start_hint.setText("当前有任务正在处理，请等待完成后再开始分析。")
            return False
        if self.jobs.warning:
            self.start_hint.setText(self.jobs.warning)
            return False
        env = self.config_mgr.get_env_dict()
        if not (env.get("API_KEY") or env.get(env.get("API_ENV_NAME", "SILICONFLOW_API_KEY"))):
            self.start_hint.setText("请先在软件左下角“设置”中填写 AI 服务的 API Key，再开始分析。")
            return False
        return True

    def _start_analysis(self):
        if not self._can_start():
            return
        if not self.inputs:
            self.start_hint.setText("请先导入 Excel，或从评论库选择评论。")
            return
        templates = load_product_templates(self.config_mgr)
        for item in self.inputs:
            if not item["selected"]:
                self.start_hint.setText(f"请为 {Path(item['path']).name} 至少选择一个工作表。")
                return
            if item.get("product") not in templates:
                self.start_hint.setText(f"请为 {Path(item['path']).name} 选择分析品类。")
                return
        self._prepare_tasks(self.inputs, templates, self.output_dir.text().strip(), self.mode_combo.currentData())

    def _prepare_tasks(self, items, templates, output_dir="", mode="full"):
        import copy
        items, templates = copy.deepcopy(items), copy.deepcopy(templates)
        self._preparing = True
        self._set_input_enabled(False)
        self.start_hint.setText("正在保存本次评论范围，准备分析…")
        def prepare():
            records = []
            for item in items:
                if mode == "summary":
                    self._validate_labels(item["path"], item["selected"], templates[item["product"]])
                records.append(self.jobs.prepare(item["path"], item["selected"], item["product"], templates[item["product"]], output_dir, mode, item.get("title", "")))
            return records
        def done(records):
            self._preparing = False
            self.jobs.records.extend(records)
            if not self._save_jobs():
                return
            self._queue = list(records)
            self._stopped = False
            self.start_hint.setText("已创建分析任务，可在“分析记录”查看进度。")
            self.tabs.setCurrentWidget(self.result_tab)
            self._refresh_history(records[0]["id"])
            QTimer.singleShot(0, self._next_task)
        self._background(prepare, done)

    @staticmethod
    def _validate_labels(path, sheets, template):
        from openpyxl import load_workbook
        required = set(template.get("stats_fields") or [])
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for name in sheets:
                rows = book[name].iter_rows(values_only=True)
                headers = list(next(rows, ()))
                if not required or not required.issubset(headers):
                    raise ValueError(f"{name} 缺少所选品类的标签列，请选择“生成完整洞察报告”。")
                if "_处理状态" in headers:
                    index = headers.index("_处理状态")
                    if any(str(row[index] or "").strip() not in ("", "成功") for row in rows):
                        raise ValueError("标签文件有未完成的评论，请先继续分析或补跑，再进行汇总。")
        finally:
            book.close()

    def _use_collected_review_file(self, payload):
        self.tabs.setCurrentWidget(self.start_tab)
        if not self._can_start():
            return
        templates = load_product_templates(self.config_mgr)
        key = payload["product_key"]
        ok, reason = check_template_fingerprint(key, payload.get("fingerprint"), templates)
        if not ok:
            self.start_hint.setText(f"无法开始：{reason}。请重新选择品类。")
            return
        # The library export is already a frozen, normalized range.
        path = payload["path"]
        self._preparing = True
        self._set_input_enabled(False)
        def prepare():
            sheets = inspect_workbook(path)
            return [self.jobs.prepare(path, [sheet["name"] for sheet in sheets], key, templates[key], str(EXPORT_DIR), "full", "、".join(payload.get("asins", [])) or "评论库分析")]
        def done(records):
            self._preparing = False
            self.jobs.records.extend(records)
            if not self._save_jobs():
                return
            self._queue = records
            self._stopped = False
            self._refresh_history(records[0]["id"])
            self.tabs.setCurrentWidget(self.result_tab)
            QTimer.singleShot(0, self._next_task)
        self._background(prepare, done)

    def _save_jobs(self):
        try:
            self.jobs.save()
            return True
        except OSError as exc:
            self.history_hint.setText(f"任务记录未能保存：{exc}")
            self.start_hint.setText(f"任务记录未能保存：{exc}")
            return False

    def _next_task(self):
        if self._active:
            return
        if not self._queue or self._stopped:
            self._set_input_enabled(self._worker is None)
            self._refresh_history()
            return
        # The previous worker emits completion immediately before its thread exits.
        if self.runner.is_running():
            QTimer.singleShot(50, self._next_task)
            return
        self._active = self._queue.pop(0)
        self._active.update(status="running", updated=now(), message="正在读取评论…")
        self._stop_file.unlink(missing_ok=True)
        self._set_input_enabled(False)
        self._save_jobs()
        self._refresh_history(self._active["id"])
        self._launch_stage()

    def _launch_stage(self):
        if not self._active:
            return
        if self._stopped:
            self._stage_finished(-3)
            return
        if self.runner.is_running():
            QTimer.singleShot(50, self._launch_stage)
            return
        job = self._active
        stage = job["stage"]
        job["message"] = "正在分析评论…" if stage == "extract" else "正在整理分析结论…"
        env = self.config_mgr.get_env_dict()
        env.update({"RAW_INPUT_FILE": job["input"], "STEP1_OUTPUT_FILE": job["step1"],
                    "STEP2_OUTPUT_FILE": job["report"], "CURRENT_PRODUCT": job["product"],
                    "CONFIG_FILE": job["config"], "TARGET_SHEETS": "", "COMMENT_COLUMN": "内容",
                    "STEP1_CONCURRENCY": str(self.concurrency_spin.value()),
                    "STEP1_RPM_LIMIT": str(self.rpm_spin.value()), "STEP1_TPM_LIMIT": str(self.tpm_spin.value()),
                    "STEP1_MAX_FAILURE_RATE": "0", "STEP1_RESUME_ONLY": "0"})
        self._disconnect_runner()
        self.runner.log_signal.connect(self._on_log)
        self.runner.finished_signal.connect(self._stage_finished)
        self._refresh_history(job["id"])
        script = "comments_step_1_fast.py" if stage == "extract" else "comments_step_2.py"
        try:
            self.runner.run_script(str(Path(self.config_mgr.get_script_dir()) / script), env_extra=env)
        except Exception as exc:
            self._on_log(str(exc))
            self._stage_finished(-1)

    def _disconnect_runner(self):
        for signal, slot in ((self.runner.log_signal, self._on_log), (self.runner.finished_signal, self._stage_finished)):
            try:
                signal.disconnect(slot)
            except TypeError:
                pass

    def _stage_finished(self, code):
        if not self._active:
            return
        self._disconnect_runner()
        job = self._active
        if code == 0 and job["stage"] == "extract" and job["mode"] != "extract":
            job["stage"] = "summary"
            job["message"] = "评论分析完成，正在准备报告…"
            self._save_jobs()
            if not self._stopped:
                QTimer.singleShot(0, self._launch_stage)
                return
        elif code == 0 and Path(job["report"]).is_file():
            job.update(status="done", message="分析完成，报告已保存。")
        else:
            job.update(status="failed", message="部分评论未完成，请重试未完成部分。" if code == 2 else "任务未完成，已保存进度。可查看详细日志后重试。")
        if job["status"] != "done" and self._stopped:
            job.update(status="stopped", message="已停止，完成的进度已保存。")
        elif job["status"] == "running":
            job.update(status="stopped", message="已停止，评论标签已完成，可继续生成报告。")
        job["updated"] = now()
        task_id = job["id"]
        self._active = None
        self._save_jobs()
        self._refresh_history(task_id)
        QTimer.singleShot(0, self._next_task)

    def _on_log(self, message):
        if not self._active:
            return
        job = self._active
        try:
            with open(job["log"], "a", encoding="utf-8") as stream:
                stream.write(message + "\n")
        except OSError:
            pass
        match = re.search(r"进度:\s*(\d+)/(\d+)", message)
        if match and not self._stopped:
            complete, total = map(int, match.groups())
            job["message"] = f"已处理 {complete} / {total} 条评论"
            job["progress"] = int(complete * 100 / max(1, total))
        selected = self._selected_job()
        if selected and selected["id"] == job["id"]:
            self.log.append(message)
            self.task_status.setText(job["message"])
            if match and not self._stopped:
                self.progress.setRange(0, 100)
                self.progress.setValue(job["progress"])

    def _stop(self):
        self._stopped = True
        if self._active:
            self._stop_file.write_text("stop", encoding="utf-8")
            self._active["message"] = "正在停止，请等待当前请求结束…" if self._active["stage"] == "extract" else "正在整理的报告会在完成后保存，后续任务已停止。"
        for job in self._queue:
            job.update(status="stopped", message="尚未开始，可稍后继续分析。", updated=now())
        self._queue.clear()
        self._save_jobs()
        self._refresh_history()

    def _resume_selected(self):
        job = self._selected_job()
        if not job or job["status"] not in ("stopped", "failed"):
            return
        if not self._can_start():
            self.history_hint.setText(self.start_hint.text())
            return
        if not Path(job["input"]).exists() or not Path(job["config"]).exists():
            self.history_hint.setText("该任务的本地分析数据已不存在，请重新导入评论创建任务。")
            return
        if job["stage"] == "summary" and not Path(job["step1"]).exists():
            self.history_hint.setText("该任务的评论标签文件已不存在，请重新创建分析任务。")
            return
        self._stopped = False
        job.update(status="queued", message="准备继续分析…")
        self._queue = [job]
        self._next_task()

    def _selected_job(self):
        row = self.history_table.currentRow()
        item = self.history_table.item(row, 0) if row >= 0 else None
        task_id = item.data(Qt.ItemDataRole.UserRole) if item else None
        return next((job for job in self.jobs.records if job["id"] == task_id), None)

    def _refresh_history(self, select_id=None):
        current = self._selected_job()
        select_id = select_id or (current["id"] if current else None)
        self.history_table.blockSignals(True)
        records = list(reversed(self.jobs.records))
        self.history_table.setRowCount(len(records))
        selected_row = 0
        for row, job in enumerate(records):
            values = [job["title"], job["product_name"], str(job["count"]), job["created"].replace("T", " "), STATUS_NAMES.get(job["status"], job["status"])]
            for col, text in enumerate(values):
                cell = QTableWidgetItem(text)
                cell.setData(Qt.ItemDataRole.UserRole, job["id"])
                self.history_table.setItem(row, col, cell)
            if job["id"] == select_id:
                selected_row = row
        if records:
            self.history_table.selectRow(selected_row)
        self.history_table.blockSignals(False)
        self._select_history()

    def _select_history(self):
        job = self._selected_job()
        self.result_table.hide()
        self.result_note.hide()
        self.result_context.clear()
        self._report_path = None
        self.log.clear()
        for button in (self.resume_btn, self.stop_btn, self.open_report_btn, self.folder_btn):
            button.hide()
        self.progress.hide()
        if not job:
            return
        self.task_status.setText(f"{job['title']} · {job['count']} 条评论\n{job.get('message', '')}")
        busy = bool(self._active or self._queue or self._preparing)
        self.resume_btn.setVisible(job["status"] in ("stopped", "failed"))
        self.resume_btn.setEnabled(not busy)
        self.resume_btn.setText("重试未完成部分" if job["status"] == "failed" else "继续分析")
        self.stop_btn.setVisible(job["status"] in ("running", "queued") and busy)
        self.stop_btn.setEnabled(not self._stopped)
        if job["status"] == "running":
            self.progress.show()
            self.progress.setRange(0, 100 if job.get("progress") and job["stage"] == "extract" else 0)
            if self.progress.maximum():
                self.progress.setValue(job["progress"])
        try:
            path = Path(job["log"])
            if path.exists():
                with path.open("rb") as stream:
                    stream.seek(max(0, path.stat().st_size - 100000))
                    self.log.setPlainText(stream.read().decode("utf-8", errors="replace"))
        except OSError:
            pass
        if job["status"] == "done":
            if not Path(job["report"]).exists():
                self.result_context.setText(f"报告文件已被移动或删除：{job['report']}")
                return
            self._report_path = job["report"]
            self.open_report_btn.show()
            self.open_report_btn.setEnabled(True)
            self.open_report_btn.setText("打开标签文件" if job["mode"] == "extract" else "打开完整报告")
            self.folder_btn.show()
            if job["mode"] != "extract":
                self._load_report(job["report"], switch=False)
                self.result_table.show()
                self.result_note.show()
            else:
                self.result_context.setText("评论标签已提取完成。可打开文件查看逐条评论及分析标签。")

    def _load_report(self, path, switch=True):
        from openpyxl import load_workbook
        self.result_table.setRowCount(0)
        self._report_path = str(Path(path).resolve())
        self.open_report_btn.setEnabled(True)
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                title = str(book["VOC看板"]["A1"].value or "评论洞察")
                selected = {}
                for row in book["摘要汇总"].iter_rows(min_row=2, values_only=True):
                    dimension = str(row[0] or "")
                    if any(word in dimension for word in ("负向", "正向", "改进")):
                        selected.setdefault(dimension, []).append(row)
                groups = sorted(selected, key=lambda dim: next((i for i, word in enumerate(("负向", "正向", "改进")) if word in dim), 3))
                rows = [row for dim in groups for row in sorted(selected[dim], key=lambda row: float(row[2] or 0), reverse=True)[:3]]
            finally:
                book.close()
            self.result_context.setText(title + ("\n" + "；".join(f"{dim}：{max(selected[dim], key=lambda row: float(row[2] or 0))[1]}" for dim in groups) if groups else ""))
            self.result_table.setRowCount(len(rows))
            for index, row in enumerate(rows):
                share = f"{row[3]:.1%}" if isinstance(row[3], (int, float)) else "—"
                for col, value in enumerate((row[0], row[1], row[2], share, row[4])):
                    self.result_table.setItem(index, col, QTableWidgetItem(str(value)))
            if switch:
                self.tabs.setCurrentWidget(self.result_tab)
        except Exception as exc:
            self.result_context.setText(f"报告已保存，预览暂不可用，可直接打开完整报告。\n{exc}")

    def _open_report(self):
        if self._report_path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self._report_path))

    def _open_folder(self):
        if self._report_path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(self._report_path).parent)))

    def _shutdown(self):
        if self._active or self._queue:
            self._stop()
        if self._worker is not None:
            self._worker.wait()
