"""独立的拖拽矩阵工作台；不会改写源工作簿。"""
import json
from io import BytesIO
from pathlib import Path

import pandas as pd
from PyQt6.QtCore import Qt, QSize, QUrl, QEvent, QTimer, pyqtSignal
from PyQt6.QtGui import QPixmap, QIcon, QColor
from PyQt6.QtNetwork import QNetworkAccessManager, QNetworkRequest
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QFormLayout, QLineEdit, QComboBox,
    QListWidget, QListWidgetItem, QAbstractItemView, QTableWidget, QTableWidgetItem,
    QHeaderView, QFileDialog, QLabel, QPushButton, QSplitter, QMessageBox,
)
from core.widgets import FileDropLineEdit, file_row


class FieldSlot(QComboBox):
    """既可下拉选择，也可从字段列表拖入。"""
    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)

    def dragEnterEvent(self, event):
        if event.mimeData().hasFormat('application/x-qabstractitemmodeldatalist'):
            event.acceptProposedAction()

    def dragMoveEvent(self, event):
        event.acceptProposedAction()

    def dropEvent(self, event):
        source = event.source()
        if isinstance(source, QListWidget) and source.currentItem():
            index = self.findText(source.currentItem().text())
            if index >= 0:
                self.setCurrentIndex(index)
                event.acceptProposedAction()


def prepare_data(frame):
    frame = frame.copy().reset_index(drop=True)
    frame.columns = [str(c).strip() for c in frame.columns]
    if frame.columns.duplicated().any():
        raise ValueError('表格存在重复列名，请先区分重复字段。')
    if '上架时间' in frame:
        dates = pd.to_datetime(frame['上架时间'], errors='coerce')
        for name, values in [('上架年份', dates.dt.strftime('%Y')),
                             ('上架月份', dates.dt.strftime('%Y-%m')),
                             ('上架季度', dates.dt.to_period('Q').astype(str))]:
            if name not in frame:
                frame[name] = values.replace('NaT', '未知').fillna('未知')
    return frame


def cell_text(value):
    return '' if pd.isna(value) else str(value)


class AxisEditor(QWidget):
    """Ordered levels and per-field visibility/order; values apply across all parents."""
    changed = pyqtSignal()

    def __init__(self, compact=False):
        super().__init__()
        self.available = []
        self.options = {}
        self.active_field = None
        outer = QHBoxLayout(self) if compact else QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(4)
        box = QVBoxLayout()
        outer.addLayout(box)
        self.levels = QListWidget()
        self.levels.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
        self.levels.setMinimumHeight(48)
        self.levels.setMaximumHeight(85 if compact else 120)
        self.levels.setToolTip('从上到下为各级分组，可拖动调整层级顺序')
        box.addWidget(self.levels)
        remove = QPushButton('移除层级')
        remove.clicked.connect(self.remove_level)
        box.addWidget(remove)
        if compact:
            box = QVBoxLayout()
            outer.addLayout(box, 2)
        self.value_title = QLabel('展示值')
        self.value_title.setWordWrap(True)
        box.addWidget(self.value_title)
        self.values = QListWidget()
        self.values.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
        self.values.setMinimumHeight(48 if compact else 80)
        self.values.setMaximumHeight(100 if compact else 240)
        self.values.setToolTip('勾选展示值，拖动调整顺序；该层设置对所有父级生效')
        box.addWidget(self.values)
        buttons = QHBoxLayout()
        for label, checked in [('全选', True), ('全不选', False)]:
            button = QPushButton(label)
            button.clicked.connect(lambda _, checked=checked: self.check_all(checked))
            buttons.addWidget(button)
        box.addLayout(buttons)
        self.levels.currentTextChanged.connect(self.show_values)
        self.levels.model().rowsMoved.connect(lambda *_: self.changed.emit())
        self.values.itemChanged.connect(self.save_values)
        self.values.model().rowsMoved.connect(self.save_values)

    def fields(self):
        return [self.levels.item(i).text() for i in range(self.levels.count())]

    def add_field(self, field):
        if field not in self.available:
            return
        if field not in self.fields():
            self.levels.addItem(field)
        self.levels.setCurrentRow(self.fields().index(field))
        self.changed.emit()

    def remove_level(self):
        self.levels.takeItem(self.levels.currentRow())
        self.changed.emit()

    def show_values(self, field):
        self.active_field = None
        self.values.clear()
        for value, checked in self.options.get(field, []):
            item = QListWidgetItem(value)
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)
            self.values.addItem(item)
        self.active_field = field
        self.value_title.setText(field or '展示值')

    def save_values(self, *_):
        if self.active_field:
            self.options[self.active_field] = [
                [self.values.item(i).text(), self.values.item(i).checkState() == Qt.CheckState.Checked]
                for i in range(self.values.count())]
            self.changed.emit()

    def check_all(self, checked):
        for i in range(self.values.count()):
            self.values.item(i).setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)

    def sync(self, frame):
        for field in self.fields():
            actual = sorted(set(frame[field].map(lambda x: cell_text(x) or '未知')))
            old = self.options.get(field, [])
            self.options[field] = [entry for entry in old if entry[0] in actual]
            known = {entry[0] for entry in old}
            self.options[field].extend([value, True] for value in actual if value not in known)
        self.show_values(self.levels.currentItem().text() if self.levels.currentItem() else '')

    def payload(self):
        return {'levels': self.fields(), 'values': self.options}

    def restore(self, data):
        if isinstance(data, str):  # Older single-field templates.
            data = {'levels': [] if data == '不分组' else [data], 'values': {}}
        levels = data['levels']
        if len(set(levels)) != len(levels) or any(field not in self.available for field in levels):
            raise ValueError('模板包含重复或不存在的分组字段')
        self.levels.clear()
        self.options = json.loads(json.dumps(data.get('values', {})))
        for field in levels:
            self.add_field(field)
        self.changed.emit()

    # Single-field compatibility for startup checks and old callers.
    def clear(self):
        self.levels.clear()
        self.available = []
        self.options = {}

    def addItem(self, field):
        if field != '不分组':
            self.available.append(field)

    def addItems(self, fields):
        for field in fields:
            self.addItem(field)

    def findText(self, field):
        return self.available.index(field) if field in self.available else -1

    def setCurrentText(self, field):
        self.levels.clear()
        self.add_field(field)


def axis_keys(frame, axis):
    fields = axis.fields()
    if not fields:
        return pd.Series([('全部产品',)] * len(frame), index=frame.index), {}
    keys = frame[fields].map(lambda x: cell_text(x) or '未知')
    ranks = {field: {value: i for i, (value, checked) in enumerate(axis.options[field]) if checked}
             for field in fields}
    mask = pd.Series(True, index=frame.index)
    for field in fields:
        mask &= keys[field].isin(ranks[field])
    selected = keys[mask]
    return pd.Series(list(selected.itertuples(index=False, name=None)), index=selected.index, dtype=object), ranks


def axis_label(key, fields):
    return '\n'.join('  ' * i + (f'{fields[i]}：' if fields else '') + value
                     for i, value in enumerate(key))


def hierarchy_spans(keys):
    """Merge only contiguous equal prefixes, never children of different parents."""
    for level in range(len(keys[0]) if keys else 0):
        start = 0
        while start < len(keys):
            end = start + 1
            while end < len(keys) and keys[end][:level + 1] == keys[start][:level + 1]:
                end += 1
            yield level, start, end, keys[start][level]
            start = end


class MatrixWorkspace(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.frame = None
        self.image_cache = {}
        self.network = QNetworkAccessManager(self)
        self.generation = 0
        self.cells = {}
        self.row_names = []
        self.col_names = []
        layout = QVBoxLayout(self)
        source = QHBoxLayout()
        self.path = FileDropLineEdit(placeholder='选择 Excel 后读取 Sheet')
        source.addLayout(file_row(self.path, 'open', parent=self))
        self.sheets = QComboBox()
        source.addWidget(self.sheets)
        read = QPushButton('读取数据')
        read.clicked.connect(self.read_data)
        source.addWidget(read)
        layout.addLayout(source)
        self.path.textChanged.connect(self.read_sheets)
        split = QSplitter()
        self.workspace_split = split
        split.setChildrenCollapsible(False)
        split.setHandleWidth(6)
        layout.addWidget(split, 1)
        controls = QWidget()
        left = QVBoxLayout(controls)
        controls.setMinimumWidth(145)
        controls.setMaximumWidth(210)
        left.setContentsMargins(2, 2, 2, 2)
        left.addWidget(QLabel('字段库'))
        field_search = QLineEdit()
        field_search.setPlaceholderText('搜索字段名称')
        left.addWidget(field_search)
        self.fields = QListWidget()
        self.fields.setDragEnabled(True)
        self.fields.setDragDropMode(QAbstractItemView.DragDropMode.DragOnly)
        self.fields.setDefaultDropAction(Qt.DropAction.CopyAction)
        left.addWidget(self.fields)
        self.fields.setToolTip('拖到右侧设置分组；勾选或取消勾选可展示或隐藏卡片字段')
        field_search.textChanged.connect(self.filter_fields)
        split.addWidget(controls)
        canvas = QWidget()
        right = QVBoxLayout(canvas)
        right.setContentsMargins(0, 0, 0, 0)
        self.horizontal = AxisEditor(compact=True)
        self.vertical = AxisEditor()
        self.sort = FieldSlot()
        self.direction = QComboBox()
        self.direction.addItems(['从旧到新 / 升序', '从新到旧 / 降序'])
        toolbar = QHBoxLayout()
        toolbar.addWidget(QLabel('格内排序'))
        toolbar.addWidget(self.sort)
        toolbar.addWidget(self.direction)
        self.search = QLineEdit()
        self.search.setPlaceholderText('按 ASIN、品牌或任意文字筛选')
        toolbar.addWidget(self.search, 1)
        swap = QPushButton('交换横纵轴')
        swap.clicked.connect(self.swap_axes)
        toolbar.addWidget(swap)
        right.addLayout(toolbar)
        self.drop_targets = {}
        self.canvas_split = QSplitter(Qt.Orientation.Vertical)
        self.canvas_split.setChildrenCollapsible(False)
        self.canvas_split.setHandleWidth(6)
        self.canvas_split.addWidget(self.make_drop_zone('横向分组', 'horizontal', self.horizontal))
        middle = QSplitter(Qt.Orientation.Horizontal)
        self.matrix_split = middle
        middle.setChildrenCollapsible(False)
        middle.setHandleWidth(6)
        vertical_zone = self.make_drop_zone('纵向分组', 'vertical', self.vertical)
        vertical_zone.setMinimumWidth(150)
        middle.addWidget(vertical_zone)
        preview = QWidget()
        center = QVBoxLayout(preview)
        center.setContentsMargins(0, 0, 0, 0)
        self.card_summary = QLabel('卡片字段')
        self.card_summary.setToolTip('拖入字段添加卡片内容，取消左侧勾选可移除')
        self.register_drop_target(self.card_summary, 'cards')
        center.addWidget(self.card_summary)
        self.table = QTableWidget()
        self.table.setMinimumHeight(340)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.register_drop_target(self.table.viewport(), 'cards')
        center.addWidget(self.table, 1)
        middle.addWidget(preview)
        middle.setSizes([180, 1000])
        middle.setStretchFactor(1, 1)
        self.canvas_split.addWidget(middle)
        self.canvas_split.setSizes([170, 650])
        self.canvas_split.setStretchFactor(1, 1)
        right.addWidget(self.canvas_split, 1)
        actions = QHBoxLayout()
        for title, callback in [('保存布局模板', self.save_template),
                                ('加载布局模板', self.load_template),
                                ('导出当前矩阵 Excel', self.export)]:
            button = QPushButton(title)
            button.clicked.connect(callback)
            actions.addWidget(button)
        right.addLayout(actions)
        split.addWidget(canvas)
        split.setSizes([170, 1200])
        split.setStretchFactor(1, 1)
        self.status = QLabel('读取数据后，自由选择维度。格内卡片可拖动排序；导出保留当前顺序。')
        self.status.setWordWrap(True)
        layout.addWidget(self.status)
        self.refresh_timer = QTimer(self)
        self.refresh_timer.setSingleShot(True)
        self.refresh_timer.setInterval(180)
        self.refresh_timer.timeout.connect(self.render)
        for axis in (self.horizontal, self.vertical):
            axis.changed.connect(self.schedule_render)
        for combo in (self.sort, self.direction):
            combo.currentIndexChanged.connect(self.schedule_render)
        self.fields.itemChanged.connect(self.schedule_render)
        self.search.textChanged.connect(self.schedule_render)

    def filter_fields(self, query):
        for i in range(self.fields.count()):
            item = self.fields.item(i)
            item.setHidden(query.casefold() not in item.text().casefold())

    def schedule_render(self, *_):
        self.refresh_timer.start()

    def register_drop_target(self, widget, role):
        widget.setAcceptDrops(True)
        widget.setProperty('matrixDropRole', role)
        widget.installEventFilter(self)

    def make_drop_zone(self, title, role, control):
        zone = QWidget()
        zone.setObjectName('matrixDropZone')
        zone.setStyleSheet('QWidget#matrixDropZone {border: 2px dashed #adb9ee; border-radius: 8px; background: #f3f5ff;}')
        box = QVBoxLayout(zone)
        box.setContentsMargins(6, 4, 6, 4)
        box.setSpacing(4)
        zone.setToolTip('从字段库拖入可追加分组层级；拖动分隔线可调整面板大小')
        label = QLabel(title)
        label.setWordWrap(True)
        box.addWidget(label)
        box.addWidget(control)
        if role == 'vertical':
            box.addStretch()
        for widget in (zone, label, control):
            self.register_drop_target(widget, role)
        if isinstance(control, AxisEditor):
            for child in control.findChildren(QWidget):
                self.register_drop_target(child, role)
        self.drop_targets[role] = zone
        return zone

    def eventFilter(self, watched, event):
        kind = event.type()
        if kind in (QEvent.Type.DragEnter, QEvent.Type.DragMove, QEvent.Type.Drop):
            # Only the field library changes the layout; product drags retain InternalMove.
            if event.source() is self.fields and self.fields.currentItem():
                role = watched.property('matrixDropRole')
                if role:
                    event.setDropAction(Qt.DropAction.CopyAction)
                    event.accept()
                    if kind == QEvent.Type.Drop:
                        self.apply_field_drop(role, self.fields.currentItem().text())
                    return True
        return super().eventFilter(watched, event)

    def apply_field_drop(self, role, field):
        if self.frame is None or field not in self.frame.columns:
            return
        if role in ('horizontal', 'vertical'):
            getattr(self, role).add_field(field)
        elif role == 'cards':
            for i in range(self.fields.count()):
                item = self.fields.item(i)
                if item.text() == field:
                    item.setCheckState(Qt.CheckState.Checked)
                    break
        self.schedule_render()

    def read_sheets(self):
        path = self.path.text().strip()
        self.sheets.clear()
        if not Path(path).is_file():
            return
        try:
            with pd.ExcelFile(path) as book:
                self.sheets.addItems(book.sheet_names)
        except Exception as error:
            self.status.setText(str(error))

    def read_data(self):
        try:
            if not self.sheets.currentText():
                raise ValueError('请先选择 Excel 和 Sheet。')
            self.frame = prepare_data(pd.read_excel(
                self.path.text().strip(), sheet_name=self.sheets.currentText()))
            self.fields.clear()
            for column in self.frame.columns:
                item = QListWidgetItem(column)
                item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                item.setCheckState(Qt.CheckState.Checked if column in
                                   ['ASIN', '品牌', '上架时间', '价格($)']
                                   else Qt.CheckState.Unchecked)
                self.fields.addItem(item)
            for combo in (self.horizontal, self.vertical, self.sort):
                combo.clear()
                combo.addItem('不分组' if combo != self.sort else '原始顺序')
                combo.addItems(list(self.frame.columns))
            for combo, name in [(self.horizontal, '品牌'),
                                (self.vertical, '上架年份'), (self.sort, '上架时间')]:
                if combo.findText(name) >= 0:
                    combo.setCurrentText(name)
            self.render()
        except Exception as error:
            QMessageBox.warning(self, '读取失败', str(error))

    def selected_fields(self):
        return [self.fields.item(i).text() for i in range(self.fields.count())
                if self.fields.item(i).checkState() == Qt.CheckState.Checked]

    def swap_axes(self):
        h, v = self.horizontal.payload(), self.vertical.payload()
        self.horizontal.restore(v)
        self.vertical.restore(h)
        self.render()

    def render(self):
        self.refresh_timer.stop()
        if self.frame is None:
            return
        self.card_summary.setText('卡片字段')
        self.card_summary.setToolTip('当前展示：' + ('、'.join(self.selected_fields()) or '未选择字段') + '\n拖入字段添加，取消左侧勾选可移除')
        frame = self.frame.copy()
        query = self.search.text().strip()
        if query:
            mask = frame.astype(str).apply(
                lambda col: col.str.contains(query, case=False, regex=False)).any(axis=1)
            frame = frame[mask]
        sort = self.sort.currentText()
        if self.sort.currentIndex() > 0:
            key = frame[sort]
            if sort == '上架时间':
                key = pd.to_datetime(key, errors='coerce')
            frame = frame.assign(_sort=key).sort_values(
                '_sort', ascending=self.direction.currentIndex() == 0,
                na_position='last', kind='stable').drop(columns='_sort')
        for axis in (self.horizontal, self.vertical):
            axis.sync(self.frame)
        hk, hr = axis_keys(frame, self.horizontal)
        vk, vr = axis_keys(frame, self.vertical)
        frame = frame.loc[frame.index.intersection(hk.index).intersection(vk.index)].copy()
        frame['_h'], frame['_v'] = hk.reindex(frame.index), vk.reindex(frame.index)
        def ordered(keys, axis, ranks):
            return sorted(set(keys), key=lambda key: tuple(ranks[field][value] for field, value in zip(axis.fields(), key)))
        col_keys = ordered(frame['_h'], self.horizontal, hr)
        row_keys = ordered(frame['_v'], self.vertical, vr)
        self.col_keys, self.row_keys = col_keys, row_keys
        self.header_rows = max(1, len(self.horizontal.fields()))
        self.header_cols = max(1, len(self.vertical.fields()))
        self.col_names = [axis_label(key, self.horizontal.fields()) for key in col_keys]
        self.row_names = [axis_label(key, self.vertical.fields()) for key in row_keys]
        if len(self.col_names) * len(self.row_names) > 10000:
            self.generation += 1
            self.cells = {}
            self.table.setRowCount(0)
            self.table.setColumnCount(0)
            self.status.setText('矩阵超过 10000 格，请使用上方筛选，或更换横向／纵向分组。')
            return
        self.table.clear()
        self.table.clearSpans()
        self.generation += 1
        self.cells = {}
        top, left = self.header_rows, self.header_cols
        self.table.setRowCount(top + len(row_keys))
        self.table.setColumnCount(left + len(col_keys))
        self.table.setHorizontalHeaderLabels((self.vertical.fields() or ['分组']) + [''] * len(col_keys))
        self.table.setVerticalHeaderLabels([''] * self.table.rowCount())
        def header(row, col, text, rows=1, cols=1):
            item = QTableWidgetItem(text)
            item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            item.setBackground(QColor('#eef2ff'))
            item.setForeground(QColor('#273451'))
            self.table.setItem(row, col, item)
            if rows > 1 or cols > 1:
                self.table.setSpan(row, col, rows, cols)
        for level in range(top):
            self.table.setRowHeight(level, 42)
        for level, field in enumerate(self.vertical.fields() or ['分组']):
            self.table.setColumnWidth(level, 130)
            header(0, level, field, top)
        for level, start, end, value in hierarchy_spans(col_keys):
            field = self.horizontal.fields()[level] if self.horizontal.fields() else '分组'
            header(level, left + start, f'{field}：{value}', cols=end - start)
        for level, start, end, value in hierarchy_spans(row_keys):
            header(top + start, level, value, rows=end - start)
        attrs = self.selected_fields()
        groups = {(row, col): group for (row, col), group in frame.groupby(['_v', '_h'], sort=False)}
        for ri, row_key in enumerate(row_keys):
            self.table.setRowHeight(top + ri, 300)
            for ci, col_key in enumerate(col_keys):
                products = groups.get((row_key, col_key))
                self.table.setColumnWidth(left + ci, 250)
                if products is None:
                    continue
                cards = QListWidget()
                cards.setIconSize(QSize(100, 100))
                cards.setWordWrap(True)
                cards.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
                self.register_drop_target(cards.viewport(), 'cards')
                cards.setStyleSheet('QListWidget::item {padding:10px; margin:4px; border:1px solid #dce2ef; border-radius:6px;}')
                for index, product in products.iterrows():
                    label = '\n'.join(f'{field}: {cell_text(product[field])}' for field in attrs)
                    item = QListWidgetItem(label or f'产品 #{index + 1}')
                    item.setData(Qt.ItemDataRole.UserRole, int(index))
                    cards.addItem(item)
                    if '商品主图' in product:
                        self.attach_image(item, cell_text(product['商品主图']), self.generation)
                self.cells[(ri, ci)] = cards
                self.table.setCellWidget(top + ri, left + ci, cards)
        self.preview_attrs = attrs
        self.status.setText(f'{len(frame)} 个产品 · {len(self.row_names)} 行 × {len(self.col_names)} 列。可拖动格内卡片排序、调整行高列宽。')

    def attach_image(self, item, url, generation):
        if not url:
            return
        def show(raw):
            pixmap = QPixmap()
            if pixmap.loadFromData(raw):
                self.image_cache[url] = raw
                if generation == self.generation:
                    item.setIcon(QIcon(pixmap))
        if url in self.image_cache:
            show(self.image_cache[url])
        elif QUrl(url).scheme() in ('https', 'http'):
            request = QNetworkRequest(QUrl(url))
            request.setTransferTimeout(15000)
            reply = self.network.get(request)
            def finished():
                if reply.error() == reply.NetworkError.NoError:
                    show(bytes(reply.readAll()))
                reply.deleteLater()
            reply.finished.connect(finished)
        elif Path(url).is_file():
            show(Path(url).read_bytes())

    def save_template(self):
        path, _ = QFileDialog.getSaveFileName(self, '保存布局', '', 'JSON (*.json)')
        if path:
            payload = dict(horizontal=self.horizontal.payload(), vertical=self.vertical.payload(),
                           sort=self.sort.currentText(), direction=self.direction.currentIndex(),
                           fields=self.selected_fields(), search=self.search.text())
            Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')

    def load_template(self):
        if self.frame is None:
            self.status.setText('请先读取数据，再加载布局模板。')
            return
        path, _ = QFileDialog.getOpenFileName(self, '加载布局', '', 'JSON (*.json)')
        if not path:
            return
        try:
            data = json.loads(Path(path).read_text(encoding='utf-8'))
            for key, combo in [('sort', self.sort)]:
                if combo.findText(data[key]) < 0:
                    raise ValueError(f'当前数据缺少模板字段：{data[key]}')
            # Validate both axes before mutating the displayed layout.
            for key, axis in [('horizontal', self.horizontal), ('vertical', self.vertical)]:
                probe = AxisEditor()
                probe.available = axis.available
                try:
                    probe.restore(data[key])
                finally:
                    probe.deleteLater()
            for key, axis in [('horizontal', self.horizontal), ('vertical', self.vertical)]:
                axis.restore(data[key])
            for key, combo in [('sort', self.sort)]:
                combo.setCurrentText(data[key])
            self.direction.setCurrentIndex(data.get('direction', 0))
            self.search.setText(data.get('search', ''))
            for i in range(self.fields.count()):
                item = self.fields.item(i)
                item.setCheckState(Qt.CheckState.Checked if item.text() in data['fields'] else Qt.CheckState.Unchecked)
            self.render()
        except Exception as error:
            QMessageBox.warning(self, '模板无法加载', str(error))

    def export(self):
        if self.refresh_timer.isActive():
            self.render()
        if not self.cells:
            self.status.setText('请先生成矩阵预览。')
            return
        path, _ = QFileDialog.getSaveFileName(self, '导出矩阵', '自由组合矩阵.xlsx', 'Excel (*.xlsx)')
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image
            from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
            book = Workbook()
            sheet = book.active
            sheet.title = '自由组合矩阵'
            top, left = self.header_rows, self.header_cols
            def header(row, col, text, rows=1, cols=1):
                cell = sheet.cell(row, col, text)
                if rows > 1 or cols > 1:
                    sheet.merge_cells(start_row=row, start_column=col,
                                      end_row=row + rows - 1, end_column=col + cols - 1)
            for level, field in enumerate(self.vertical.fields() or ['分组']):
                header(1, level + 1, field, top)
            for level, start, end, value in hierarchy_spans(self.col_keys):
                field = self.horizontal.fields()[level] if self.horizontal.fields() else '分组'
                header(level + 1, left + start + 1, f'{field}：{value}', cols=end - start)
            details = book.create_sheet('产品明细')
            details.append(['纵向分组', '横向分组'] + list(self.frame.columns))
            output_row = top + 1
            row_offsets = [output_row]
            image_streams = []
            for ri, row_name in enumerate(self.row_names):
                height = max([1] + [cards.count() for (row, _), cards in self.cells.items() if row == ri])
                for ci, col_name in enumerate(self.col_names):
                    cards = self.cells.get((ri, ci))
                    if cards is None:
                        continue
                    for offset in range(cards.count()):
                        item = cards.item(offset)
                        cell = sheet.cell(output_row + offset, left + ci + 1, item.text())
                        product = self.frame.loc[item.data(Qt.ItemDataRole.UserRole)]
                        url = cell_text(product.get('商品主图', ''))
                        if url in self.image_cache:
                            stream = BytesIO(self.image_cache[url])
                            image_streams.append(stream)
                            picture = Image(stream)
                            ratio = min(120 / picture.width, 100 / picture.height)
                            picture.width *= ratio
                            picture.height *= ratio
                            sheet.add_image(picture, cell.coordinate)
                            cell.value = '\n' * 8 + item.text()
                        details.append([row_name, col_name] + [cell_text(x) for x in product])
                for r in range(output_row, output_row + height):
                    sheet.row_dimensions[r].height = 100 + max(70, 16 * len(self.preview_attrs))
                output_row += height
                row_offsets.append(output_row)
            for level, start, end, value in hierarchy_spans(self.row_keys):
                header(row_offsets[start], level + 1, value, rows=row_offsets[end] - row_offsets[start])
            for row in sheet:
                for cell in row:
                    is_header = cell.row <= top or cell.column <= left
                    cell.alignment = Alignment(vertical='center' if is_header else 'top',
                                               horizontal='center' if is_header else 'left', wrap_text=True)
                    if is_header:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill('solid', fgColor='E8EDFF')
                        edge = Side(style='thin', color='DCE2EF')
                        cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            for col in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 18 if col <= left else 32
            for row in range(1, top + 1):
                sheet.row_dimensions[row].height = 32
            sheet.freeze_panes = sheet.cell(top + 1, left + 1).coordinate
            book.save(path)
            self.status.setText(f'已导出：{path}')
        except Exception as error:
            QMessageBox.warning(self, '导出失败', str(error))
