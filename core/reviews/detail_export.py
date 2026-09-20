"""User-facing, single-sheet export; independent of legacy AI input files."""
import html
import os
import re
import tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .metadata import latest_analysis,raw_of,KEY,as_dict

def display_text(value):
    text=html.unescape(str(value))
    text=re.sub(r"<br\s*/?>","\n",text,flags=re.I)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]","",text)

def export_review_details(rows,output_path):
    items=list(rows)
    if not items:raise ValueError("没有可导出的评论")
    if len(items)>1048575:raise ValueError("超过 Excel 单表行数上限，请缩小筛选范围")
    analyses=[latest_analysis(r) for r in items]
    fields={}
    for a in analyses:
        if a:
            for key in a.get("labels",{}):
                label=str(a.get("field_labels",{}).get(key,key))
                fields[(key,label)]="AI · "+label
    base=[("marketplace","站点"),("asin","目标 ASIN"),("review_id","评论ID"),("rating","星级"),
          ("review_date","日期"),("reviewer_name","评论者"),("title","标题"),("content","内容"),
          ("verified_purchase","已验证购买"),("helpful_votes","点赞数"),("review_url","评论链接")]
    headers=[label for _,label in base]+["分析状态"]+list(fields.values())+["分析品类","模型","模板版本","分析时间","来源变体 ASIN","来源文件"]
    # Keep distinct fields visible if templates reuse the same display label.
    for i,h in enumerate(headers):
        if headers[:i].count(h):headers[i]=h+" ("+str(i+1)+")"
    book=Workbook();sheet=book.active;sheet.title="评论明细"
    sheet.append(headers);sheet.freeze_panes="A2"
    for r,a in zip(items,analyses):
        sources=as_dict(as_dict(raw_of(r).get(KEY)).get("imports"))
        sources=[v for v in sources.values() if isinstance(v,dict) and v.get("target_asin")==r.get("asin")]
        old=as_dict(as_dict(raw_of(r).get(KEY)).get("analyses"))
        if a:state="已分析" if any(a.get("labels",{}).values()) else "已分析，无相关标签"
        else:state="需重新分析" if old else "未分析"
        values=[r.get(k,"") if k!="review_date" else r.get(k) or r.get("review_date_raw","") for k,_ in base]+[state]
        for key,label in fields:
            values.append(a.get("labels",{}).get(key,"") if a and str(a.get("field_labels",{}).get(key,key))==label else "")
        values += [a.get(k,"") if a else "" for k in ("category","model","template_hash","created_at")]
        values += ["；".join(sorted({v.get("linked_asin","") for v in sources if v.get("linked_asin")})),
                   "；".join(sorted({v.get("file","") for v in sources if v.get("file")}))]
        sheet.append(["；".join(map(str,v)) if isinstance(v,list) else v for v in values])
        n=sheet.max_row
        for cell in sheet[n]:
            if isinstance(cell.value,str):
                clean=display_text(cell.value)
                if len(clean)>32767:raise ValueError("单元格内容超过 Excel 限制，请缩小分析版本或来源范围")
                cell.value=clean;cell.data_type="s"  # Untrusted review text is never an Excel formula.
            cell.alignment=Alignment(vertical="top",wrap_text=True)
        url=str(r.get("review_url") or "")
        if url.startswith("https://"):
            cell=sheet.cell(n,11);cell.hyperlink=url;cell.value="打开评论";cell.style="Hyperlink";cell.data_type="s"
        sheet.row_dimensions[n].height=72
    for cell in sheet[1]:
        cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="3559B7")
        cell.alignment=Alignment(vertical="center",wrap_text=True)
    sheet.row_dimensions[1].height=30
    for i,h in enumerate(headers,1):
        width=42 if h in ("内容","标题") or h.startswith("AI ·") else 22
        if h=="内容":width=70
        if h=="来源文件":width=50
        if h in ("星级","已验证购买","点赞数"):width=12
        sheet.column_dimensions[get_column_letter(i)].width=width
    sheet.auto_filter.ref=sheet.dimensions
    path=Path(output_path).resolve();path.parent.mkdir(parents=True,exist_ok=True)
    temporary=None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent,suffix=".xlsx",delete=False) as f:temporary=f.name
        book.save(temporary)
        os.replace(temporary,path)
    finally:
        book.close()
        if temporary and os.path.exists(temporary):os.unlink(temporary)
    return str(path)
