"""Adapters between canonical review rows and PM Stack's Excel AI pipeline."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


EXPORT_FIELDS: List[Tuple[str, str]] = [
    ("asin", "ASIN"),
    ("review_id", "评论ID"),
    ("rating", "星级"),
    ("star_filter", "来源筛选"),
    ("review_date", "日期"),
    ("reviewer_name", "评论者"),
    ("title", "标题"),
    ("content", "内容"),
    ("verified_purchase", "已验证购买"),
    ("helpful_votes", "点赞数"),
    ("review_url", "评论链接"),
]


def fetch_all_reviews(
    repository,
    *,
    asin: Optional[str] = None,
    rating: Optional[int] = None,
    keyword: str = "",
    page_size: int = 1000,
) -> Tuple[List[Dict[str, Any]], str]:
    """Read all matching rows without silently truncating at the RPC page limit."""
    rows: List[Dict[str, Any]] = []
    offset = 0
    source = ""
    while True:
        page, source = repository.query_reviews(
            asin=asin,
            rating=rating,
            keyword=keyword,
            limit=page_size,
            offset=offset,
        )
        rows.extend(page)
        if len(page) < page_size:
            break
        offset += len(page)
    return rows, source


def export_reviews_for_analysis(rows: Iterable[Dict[str, Any]], output_path: str) -> str:
    """Write Chinese headers and star sheets expected by comments_step_1.py."""
    items = list(rows)
    if not items:
        raise ValueError("没有可导出的评论")

    destination = Path(output_path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for item in items:
        try:
            star = int(round(float(item.get("rating"))))
        except (TypeError, ValueError):
            star = 0
        sheet = f"{star} star" if 1 <= star <= 5 else "未分级"
        grouped.setdefault(sheet, []).append(item)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in sorted(
        grouped,
        key=lambda name: -(int(name.split()[0])) if name[0].isdigit() else 1,
    ):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append([label for _, label in EXPORT_FIELDS])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for item in grouped[sheet_name]:
            worksheet.append([item.get(key, "") for key, _ in EXPORT_FIELDS])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.column_dimensions["G"].width = 34
        worksheet.column_dimensions["H"].width = 72
        for cells in worksheet.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(destination)
    return str(destination)


__all__ = ["EXPORT_FIELDS", "export_reviews_for_analysis", "fetch_all_reviews"]

