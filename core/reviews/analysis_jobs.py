"""Local analysis records and immutable workbook/template snapshots.

Each task owns its input, checkpoint and report. Retrying a summary never
re-extracts reviews, and a later edit to the original file/template cannot
silently change an existing task.
"""
from __future__ import annotations

import json
import os
import re
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


CONTENT_HEADERS = {"内容", "评论内容", "评论正文", "content", "review content", "review text", "body"}


def now():
    return datetime.now().isoformat(timespec="seconds")


def _content_index(headers):
    for index, value in enumerate(headers):
        if str(value or "").strip().lower() in CONTENT_HEADERS:
            return index
    return None


def inspect_workbook(path):
    book = load_workbook(path, read_only=True, data_only=True)
    result = []
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            index = _content_index(headers)
            if index is None:
                continue
            count = sum(1 for row in rows if index < len(row) and str(row[index] or "").strip())
            if count:
                result.append({"name": sheet.title, "count": count})
    finally:
        book.close()
    if not result:
        raise ValueError("没有找到可分析的评论。请检查首行是否包含“内容”或“评论内容”等正文列。")
    return result


def snapshot_workbook(source, destination, sheet_names):
    """Copy selected nonempty reviews into one normalized sheet (stable count)."""
    book = load_workbook(source, read_only=True, data_only=True)
    output = Workbook(write_only=True)
    sheet = output.create_sheet("评论")
    headers = []
    count = 0
    try:
        for name in sheet_names:
            if name not in book.sheetnames:
                raise ValueError(f"工作表“{name}”已不存在，请重新添加文件。")
            raw = next(book[name].iter_rows(values_only=True), ())
            for value in raw:
                label = str(value or "").strip()
                if label.lower() in CONTENT_HEADERS:
                    label = "内容"
                if label and label not in headers:
                    headers.append(label)
        for label in ("星级", "来源工作表", "来源行号"):
            if label not in headers:
                headers.append(label)
        sheet.append(headers)
        for name in sheet_names:
            rows = book[name].iter_rows(values_only=True)
            raw = next(rows, ())
            index = _content_index(raw)
            if index is None:
                raise ValueError(f"工作表“{name}”的正文列已变化，请重新添加文件。")
            labels = ["内容" if str(v or "").strip().lower() in CONTENT_HEADERS else str(v or "").strip() for v in raw]
            for row_number, row in enumerate(rows, start=2):
                if index >= len(row) or not str(row[index] or "").strip():
                    continue
                values = dict(zip(labels, row))
                if not values.get("星级"):
                    rating = re.search(r"([1-5])\s*(?:star|星)", name, re.IGNORECASE)
                    values["星级"] = int(rating.group(1)) if rating else None
                values.setdefault("来源工作表", name)
                values.setdefault("来源行号", row_number)
                sheet.append([values.get(label) for label in headers])
                count += 1
        if not count:
            raise ValueError("所选范围没有有效评论。")
        output.save(destination)
    finally:
        book.close()
        output.close()
    return count


class AnalysisJobs:
    def __init__(self, root):
        self.root = Path(root)
        self.path = self.root / "tasks.json"
        self.records = []
        self._reserved_reports = set()
        self.warning = ""
        if self.path.exists():
            try:
                data = json.loads(self.path.read_text(encoding="utf-8"))
                if not isinstance(data, list) or any(not isinstance(row, dict) or "id" not in row for row in data):
                    raise ValueError("记录格式不正确")
                self.records = data
            except (OSError, ValueError) as exc:
                self.warning = f"分析记录读取失败：{exc}。原文件已保留，请先处理后再创建任务。"
        for row in self.records:
            if row.get("status") in ("running", "queued"):
                row.update(status="stopped", message="上次任务未完成，可继续分析。")

    def save(self):
        if self.warning:
            raise OSError(self.warning)
        self.root.mkdir(parents=True, exist_ok=True)
        pending = self.path.with_suffix(".tmp")
        pending.write_text(json.dumps(self.records, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(pending, self.path)

    def prepare(self, source, sheets, key, template, output_dir="", mode="full", title=""):
        task_id = uuid.uuid4().hex
        folder = self.root / task_id
        folder.mkdir(parents=True)
        input_path = folder / "input.xlsx"
        count = snapshot_workbook(source, input_path, sheets)
        config = folder / "product_configs.json"
        config.write_text(json.dumps({key: template}, ensure_ascii=False, indent=2), encoding="utf-8")
        target_dir = Path(output_dir) if output_dir else Path(source).parent
        target_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d")
        suffix = "评论标签" if mode == "extract" else "评论分析"
        source_name = Path(source).stem
        # Imported review exports include a collection timestamp; keep only
        # their product/market prefix so the report has a single readable date.
        source_name = re.sub(r"-Reviews-\d{6,8}-\d+$", "", source_name, flags=re.IGNORECASE)
        report_stem = f"{source_name}_{suffix}_{stamp}"
        report = target_dir / f"{report_stem}.xlsx"
        reserved = self._reserved_reports | {os.path.normcase(os.path.abspath(row["report"]))
                    for row in self.records if row.get("report")}
        number = 2
        while report.exists() or os.path.normcase(os.path.abspath(report)) in reserved:
            report = target_dir / f"{report_stem}_{number}.xlsx"
            number += 1
        self._reserved_reports.add(os.path.normcase(os.path.abspath(report)))
        return {
            "id": task_id, "title": title or Path(source).stem,
            "source": str(Path(source).resolve()), "sheets": list(sheets),
            "input": str(input_path), "config": str(config),
            "step1": str(report if mode == "extract" else input_path if mode == "summary" else folder / "labels.xlsx"),
            "report": str(report), "log": str(folder / "analysis.log"),
            "product": key, "product_name": template.get("display_name", key),
            "count": count, "mode": mode, "stage": "summary" if mode == "summary" else "extract",
            "status": "queued", "created": now(), "updated": now(), "message": "等待分析",
        }
