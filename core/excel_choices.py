"""Asynchronous Excel metadata choices; existing manual values stay intact."""
from pathlib import Path
from PyQt6.QtCore import QObject, QRunnable, QThreadPool, QTimer, Qt, pyqtSignal
from PyQt6.QtWidgets import QComboBox, QCompleter
from core.widgets import configure_combo


class ChoiceEdit(QComboBox):
    textChanged = pyqtSignal(str)

    def __init__(self, text='', parent=None):
        super().__init__(parent)
        configure_combo(self)
        self.setEditable(True)
        self.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self.lineEdit().textChanged.connect(self.textChanged)
        self.setText(text)
        self.completer().setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        self.completer().setFilterMode(Qt.MatchFlag.MatchContains)
        self.completer().setCompletionMode(QCompleter.CompletionMode.PopupCompletion)

    def text(self):
        return self.currentText()

    def setText(self, text):
        self.setEditText(text)

    def setPlaceholderText(self, text):
        self.lineEdit().setPlaceholderText(text)

    def clear(self):
        super().clear()
        self.setEditText('')

    def set_choices(self, choices):
        text = self.text()
        self.blockSignals(True)
        self.lineEdit().blockSignals(True)
        self.clear()
        self.addItem('')
        self.addItems(list(dict.fromkeys(str(x) for x in choices if x is not None and str(x))))
        self.setEditText(text)
        self.lineEdit().blockSignals(False)
        self.blockSignals(False)


class _Signals(QObject):
    ready = pyqtSignal(int, object, str)


class _ReadMetadata(QRunnable):
    def __init__(self, path, generation):
        super().__init__()
        self.path, self.generation = path, generation
        self.signals = _Signals()

    def run(self):
        try:
            if Path(self.path).suffix.lower() in ('.xlsx', '.xlsm'):
                from openpyxl import load_workbook
                book = load_workbook(self.path, read_only=True, data_only=True)
                try:
                    result = {sheet.title: [str(v) for v in next(sheet.iter_rows(max_row=1, values_only=True), ()) if v is not None]
                              for sheet in book.worksheets}
                finally:
                    book.close()
            else:
                import pandas as pd
                with pd.ExcelFile(self.path) as book:
                    result = {name: list(pd.read_excel(book, sheet_name=name, nrows=0).columns) for name in book.sheet_names}
            self.signals.ready.emit(self.generation, result, '')
        except Exception as error:
            self.signals.ready.emit(self.generation, {}, str(error))


class ExcelChoices(QObject):
    """Bind file → sheets → columns without blocking the GUI or overwriting text."""
    metadata_changed = pyqtSignal()
    def __init__(self, path_edit, sheets, columns=(), parent=None):
        super().__init__(parent or path_edit)
        self.path_edit, self.sheets, self.columns = path_edit, list(sheets), list(columns)
        self.metadata = {}
        self.generation = 0
        self._jobs = {}
        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.setInterval(350)
        self.timer.timeout.connect(self._read)
        path_edit.textChanged.connect(self._schedule)
        for field, sheet in self.columns:
            if sheet is not None:
                sheet.textChanged.connect(self._update_columns)
        if path_edit.text().strip():
            self._schedule()

    def _schedule(self, *_):
        self.generation += 1
        self.metadata = {}
        for field in self.sheets:
            field.set_choices([])
        self._update_columns()
        self.metadata_changed.emit()
        self.timer.start()

    def _read(self):
        path = self.path_edit.text().strip().strip('"')
        if not Path(path).is_file():
            return
        job = _ReadMetadata(path, self.generation)
        self._jobs[self.generation] = job
        job.signals.ready.connect(self._loaded)
        QThreadPool.globalInstance().start(job)

    def _loaded(self, generation, result, error):
        self._jobs.pop(generation, None)
        if generation != self.generation:
            return
        self.metadata = result
        for field in self.sheets:
            field.set_choices(result.keys())
            field.setToolTip('读取失败，可手动填写。\n' + error if error else '可输入文字搜索工作表，也可手动填写。')
        self._update_columns()
        self.metadata_changed.emit()

    def _update_columns(self, *_):
        for field, sheet in self.columns:
            name = sheet.text().strip() if sheet is not None else ''
            if not name and self.metadata:
                name = next(iter(self.metadata))
            field.set_choices(self.metadata.get(name, []))
