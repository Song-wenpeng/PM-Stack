"""Workbook-preserving output for market field extraction (no model calls)."""
from copy import copy
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook, LXML
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor, AbsoluteAnchor
from .field_values import is_empty_field

LOG_SHEET = '字段提取日志'
LOG_FIELDS = ['文本提取字段', '图片识别提取字段', '图片识别批次日志',
              '最终缺失字段', '提取完整状态', '_处理状态']
LOG_HEADERS = ['来源工作表', '原行号', 'ASIN'] + LOG_FIELDS
ENRICHMENT = ['商品标题', '详细参数', 'SKU', '五点描述'] + [f'image_{i}' for i in range(1, 10)]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def source_fingerprint(path, fields):
    """Ignore generated columns, but include ordered source rows and formulas."""
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        headers = next(rows)
        excluded = set(fields) | set(LOG_FIELDS)
        cols = sorted((str(h), i) for i, h in enumerate(headers) if h and h not in excluded)
        canonical = sorted(set(str(h) for h in headers if h) | set(fields) | set(LOG_FIELDS) | set(ENRICHMENT))
        mapping = {i + 1: canonical.index(str(h)) + 1 for i, h in enumerate(headers) if h}
        data = []
        for row in rows:
            values = {h: rewrite_formula(None if row[i] == '' else row[i], ws.title, ws.title, mapping) for h, i in cols}
            for name in ENRICHMENT:
                values.setdefault(name, None)
            data.append(sorted(values.items()))
        return hashlib.sha256(json.dumps([ws.title, data], default=str,
                                        ensure_ascii=False).encode('utf-8')).hexdigest()
    finally:
        wb.close()


def rewrite_formula(value, context, sheet, mapping):
    if not isinstance(value, str) or not value.startswith('='):
        return value
    tokens = Tokenizer(value).items
    for token in tokens:
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        parts = token.value.rsplit('!', 1)
        qualifier, ref = (parts[0], parts[1]) if len(parts) == 2 else (None, parts[0])
        target = qualifier.strip("'").replace("''", "'") if qualifier else context
        if target != sheet:
            continue
        match = re.fullmatch(r'(\$?)([A-Z]+)(\$?\d+)?(?::(\$?)([A-Z]+)(\$?\d+)?)?', ref)
        if not match:
            continue  # named ranges and structured references are handled separately
        def endpoint(dollar, col, row):
            old = column_index_from_string(col)
            new = mapping.get(old, old)
            if new is None:
                raise ValueError('公式引用了将移到日志表的列，请先调整该公式。')
            return dollar + get_column_letter(new) + (row or '')
        first = endpoint(*match.group(1, 2, 3))
        if match.group(5):
            left, right = column_index_from_string(match.group(2)), column_index_from_string(match.group(5))
            mapped = [mapping.get(i, i) for i in range(left, right + 1)]
            if any(v is None for v in mapped) or mapped != list(range(mapped[0], mapped[0] + len(mapped))):
                raise ValueError('公式范围跨越本次移动的列，无法安全调整；请先拆分该公式范围。')
            first += ':' + endpoint(*match.group(4, 5, 6))
        token.value = (qualifier + '!' if qualifier else '') + first
    return '=' + ''.join(t.value for t in tokens)


class FieldWorkbookWriter:
    def __init__(self, source, output, fields):
        self.source = Path(source).resolve()
        self.output = Path(output or source).resolve()
        self.fields = list(dict.fromkeys(fields))
        if self.source.suffix.lower() != '.xlsx' or self.output.suffix.lower() != '.xlsx':
            raise ValueError('写回保留格式目前支持 .xlsx，请先将文件另存为 .xlsx。')
        self.expected = digest(self.source)
        self.backup = None
        wb = load_workbook(self.source)
        try:
            self._prepare(wb)  # validate before starting expensive extraction
        finally:
            wb.close()

    def _prepare(self, wb):
        ws = wb.worksheets[0]
        headers = [c.value for c in ws[1]]
        asins = [i + 1 for i, h in enumerate(headers) if str(h).strip().upper() == 'ASIN']
        if len(asins) != 1:
            raise ValueError('首张工作表第一行必须包含唯一的 ASIN 列。')
        names = [h for h in headers if h is not None]
        if len(names) != len(set(names)):
            raise ValueError('源表存在同名列，请先消除重复表头后提取。')
        if LOG_SHEET in wb and [c.value for c in wb[LOG_SHEET][1]] != LOG_HEADERS:
            raise ValueError('已有同名“字段提取日志”表但格式不同，请重命名该表后重试。')
        if (ws.tables or ws.merged_cells or ws.data_validations.count or
                len(ws.conditional_formatting) or wb.defined_names or
                any(s._charts or s._pivots for s in wb)):
            raise ValueError('源文件含合并单元格、Excel 表格、验证规则、条件格式、命名区域、图表或透视表，暂不支持安全移列（普通嵌入图片已支持）。')
        asin = asins[0]
        remaining = [i for i, h in enumerate(headers, 1) if h not in self.fields + LOG_FIELDS]
        pos = remaining.index(asin) + 1
        order = remaining[:pos] + self.fields + remaining[pos:]
        mapping = {i: None for i in range(1, len(headers) + 1)}
        for new, item in enumerate(order, 1):
            old = item if isinstance(item, int) else (headers.index(item) + 1 if item in headers else None)
            if old:
                mapping[old] = new
        # Images belong to their original cells. Move anchors, retain size/offsets.
        # Other worksheets' image anchors are deliberately untouched.
        for image in ws._images:
            anchor = image.anchor
            if isinstance(anchor, AbsoluteAnchor):
                continue  # absolute-position images do not follow cells
            if not isinstance(anchor, (OneCellAnchor, TwoCellAnchor)):
                raise ValueError('图片锚点类型暂不支持，无法安全移动图片。')
            start = anchor._from.col + 1
            end = anchor.to.col + 1 if isinstance(anchor, TwoCellAnchor) else start
            moves = [mapping.get(c, c) for c in range(start, end + 1)]
            if any(c is None for c in moves) or len({new - old for old, new in zip(range(start, end + 1), moves)}) != 1:
                raise ValueError('图片跨越本次插入或移走的列，无法保持原尺寸；请将图片放在单一原数据列后重试。')
            anchor._from.col = moves[0] - 1
            if isinstance(anchor, TwoCellAnchor):
                anchor.to.col = moves[-1] - 1
        # Move cells with their style; adjust formulas in every worksheet before relocation.
        for other in wb:
            for row in other:
                for cell in row:
                    if cell.data_type == 'f':
                        cell.value = rewrite_formula(cell.value, other.title, ws.title, mapping)
        saved = {(c.row, c.column): copy(c) for row in ws for c in row}
        row_count = ws.max_row
        dims = {i: copy(ws.column_dimensions[get_column_letter(i)]) for i in range(1, len(headers) + 1)}
        old_filter = ws.auto_filter.ref
        # topLeftCell is the scrolled viewport, not the freeze boundary.
        # Setting freeze_panes from it would turn A189 + ySplit=1 into
        # 188 frozen rows. Preserve each pane and transform columns only.
        for view in ws.views.sheetView:
            pane = view.pane
            if pane is None:
                continue
            if pane.state in ('frozen', 'frozenSplit') and pane.xSplit:
                boundary = get_column_letter(int(pane.xSplit) + 1) + '1'
                shifted = rewrite_formula('=' + boundary, ws.title, ws.title, mapping)[1:]
                pane.xSplit = column_index_from_string(re.match(r'[A-Z]+', shifted).group()) - 1
            if pane.topLeftCell:
                pane.topLeftCell = rewrite_formula('=' + pane.topLeftCell, ws.title, ws.title, mapping)[1:]
        ws.delete_cols(1, ws.max_column)
        ws.column_dimensions.clear()
        template_col = next((i for i in remaining if i > asin), asin)
        for new, item in enumerate(order, 1):
            old = item if isinstance(item, int) else (headers.index(item) + 1 if item in headers else None)
            dim = copy(dims[old or template_col])
            dim.index = get_column_letter(new)
            dim.min = dim.max = new
            if isinstance(item, str):
                dim.width = max(14, min(28, len(item) * 2 + 4))
            ws.column_dimensions[dim.index] = dim
            for row in range(1, row_count + 1):
                original = saved.get((row, old)) if old else None
                cell = ws.cell(row, new)
                if original:
                    cell.value = original.value
                    cell.data_type = original.data_type
                    cell._style = copy(original._style)
                    cell.comment = copy(original.comment)
                    if original.hyperlink:
                        cell.hyperlink = copy(original.hyperlink)
                if isinstance(item, str):
                    template = saved.get((row, template_col)) or saved[(1, template_col)]
                    cell._style = copy(template._style)
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    cell.alignment = alignment
                    cell.number_format = '0' if item == '包装数' and row > 1 else 'General'
                    if item == '包装数' and row > 1 and re.fullmatch(r'\d+(?:\.0+)?', str(cell.value).strip()):
                        cell.value = int(float(cell.value))
                    if row == 1:
                        cell.value = item
        if old_filter:
            ws.auto_filter.ref = f'A1:{get_column_letter(len(order))}{ws.max_row}'
        return ws

    def restore_logs(self, df):
        wb = load_workbook(self.source, read_only=True)
        try:
            if LOG_SHEET in wb:
                for values in wb[LOG_SHEET].iter_rows(min_row=2, values_only=True):
                    if values[0] == wb.worksheets[0].title and isinstance(values[1], int):
                        index = values[1] - 2
                        if 0 <= index < len(df):
                            for name, value in zip(LOG_FIELDS, values[3:]):
                                df.at[index, name] = value or ''
        finally:
            wb.close()

    def save(self, df):
        if digest(self.source) != self.expected:
            raise RuntimeError('源文件在提取期间被其他程序修改，已停止写回；请保留断点并重新运行。')
        wb = load_workbook(self.source)
        temp = None
        try:
            ws = self._prepare(wb)
            headers = {c.value: c.column for c in ws[1] if c.value is not None}
            for name in ENRICHMENT:
                if name in df.columns and name not in headers:
                    col = ws.max_column + 1
                    ws.cell(1, col, name)._style = copy(ws.cell(1, 1)._style)
                    headers[name] = col
            for index, row in df.iterrows():
                excel_row = int(index) + 2
                for name in self.fields + ENRICHMENT:
                    if name not in headers or name not in df.columns:
                        continue
                    cell = ws.cell(excel_row, headers[name])
                    value = row[name]
                    if is_empty_field(value) or not is_empty_field(cell.value):
                        continue
                    if hasattr(value, 'item'):
                        value = value.item()
                    if name == '包装数' and re.fullmatch(r'\d+(?:\.0+)?', str(value).strip()):
                        value = int(float(value))
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = 's'  # model text must never become a formula
            log = wb[LOG_SHEET] if LOG_SHEET in wb else wb.create_sheet(LOG_SHEET)
            if log.max_row == 1 and log.cell(1, 1).value is None:
                for col, name in enumerate(LOG_HEADERS, 1):
                    log.cell(1, col, name)._style = copy(ws.cell(1, 1)._style)
            existing = {(log.cell(r, 1).value, log.cell(r, 2).value): r for r in range(2, log.max_row + 1)}
            asin_name = next(h for h in headers if str(h).strip().upper() == 'ASIN')
            for index, row in df.iterrows():
                source_row = int(index) + 2
                target = existing.get((ws.title, source_row), log.max_row + 1)
                values = [ws.title, source_row, row.get(asin_name, '')] + [row.get(f, '') for f in LOG_FIELDS]
                for col, value in enumerate(values, 1):
                    if str(value) in ('nan', '<NA>'):
                        value = ''
                    cell = log.cell(target, col)
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = 's'
                    cell._style = copy(ws.cell(source_row, headers[asin_name])._style)
                    cell.number_format = 'General'
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    cell.alignment = alignment
            for col in range(1, len(LOG_HEADERS) + 1):
                log.column_dimensions[get_column_letter(col)].width = 42 if col > 3 else 18
            log.freeze_panes = 'D2'
            log.auto_filter.ref = log.dimensions
            self.output.parent.mkdir(parents=True, exist_ok=True)
            fd, temp = tempfile.mkstemp(prefix=self.output.stem + '.pending-', suffix='.xlsx', dir=self.output.parent)
            os.close(fd)
            wb.save(temp)
            # ElementTree emits literal CRs, which XML readers normalize to LF.
            # Escape them before a later save can silently alter source text.
            with ZipFile(temp) as archive:
                has_cr = any(b'\r' in archive.read(n) for n in archive.namelist() if n.endswith('.xml'))
            if has_cr:
                normalized = temp + '.normalized'
                try:
                    with ZipFile(temp) as original, ZipFile(normalized, 'w') as fixed:
                        for entry in original.infolist():
                            data = original.read(entry.filename)
                            if entry.filename.endswith('.xml'):
                                # et_xmlfile's filename writer uses Windows text
                                # mode (LF -> CRLF); undo that before escaping CR.
                                if os.name == 'nt' and not LXML and entry.filename.startswith('xl/worksheets/'):
                                    data = data.replace(b'\r\n', b'\n')
                                data = data.replace(b'\r', b'&#13;')
                            fixed.writestr(entry, data)
                    os.replace(normalized, temp)
                finally:
                    if os.path.exists(normalized):
                        os.unlink(normalized)
            if digest(self.source) != self.expected:
                raise RuntimeError('源文件保存期间发生变化；结果已保存在 ' + temp)
            if self.backup is None and self.output == self.source:
                self.backup = self.source.with_name(self.source.stem + '.backup_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.xlsx')
                shutil.copy2(self.source, self.backup)
                print(f'原文件备份：{self.backup}', flush=True)
            try:
                os.replace(temp, self.output)
            except OSError as exc:
                raise RuntimeError(f'无法写回，请关闭 Excel/WPS 后重试。已完成结果保留在：{temp}') from exc
            self.expected = digest(self.source)
        finally:
            wb.close()
