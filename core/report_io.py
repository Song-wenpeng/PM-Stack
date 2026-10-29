"""Read-only source guards and validated atomic report/checkpoint writes."""
import hashlib
import json
import math
import os
from datetime import date, datetime
from numbers import Real
from pathlib import Path
import tempfile
import time
from zipfile import ZipFile


RETRY_DELAYS = (0.05, 0.1, 0.2, 0.4)


def normalized_path(path):
    return Path(os.path.abspath(os.path.expanduser(os.fspath(path)))).resolve()


def same_file(left, right):
    """Include Windows case aliases, relative paths, symlinks and hardlinks."""
    a, b = normalized_path(left), normalized_path(right)
    if str(a).replace('\\', '/').casefold() == str(b).replace('\\', '/').casefold():
        return True
    try:
        return os.path.samefile(a, b)
    except (OSError, ValueError):
        return False


def require_distinct(source, *destinations):
    paths = (source,) + destinations
    for i, left in enumerate(paths):
        for right in paths[i + 1:]:
            if same_file(left, right):
                raise ValueError('输入、输出和断点必须是不同文件，禁止覆盖输入原件。')


def category_output_path(source, output=''):
    source = normalized_path(source)
    return normalized_path(output) if str(output).strip() else source.with_name(source.stem + '_类目筛选结果.xlsx')


def category_state_dir():
    override = os.environ.get('CATEGORY_STATE_DIR', '').strip()
    if override:
        return normalized_path(override)
    base = os.environ.get('LOCALAPPDATA')
    root = Path(base) / 'PM Stack' if base else Path.home() / '.pm-stack'
    return root / 'tasks' / 'category-screening'


def category_checkpoint_path(source, checkpoint=''):
    if str(checkpoint).strip():
        return normalized_path(checkpoint)
    identity = str(normalized_path(source)).casefold().encode('utf-8')
    return category_state_dir() / (hashlib.sha256(identity).hexdigest() + '.json')


REFERENCE_DATE_FORMATS = ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d')


def resolve_reference_date(raw='', today=None, label='分析基准日', env_name=''):
    """解析显式基准日，留空则回落到运行当日。

    返回 (datetime, 来源说明)。写死一个过期的基准日会让「上架天数」之类的
    相对指标静默变成负数，所以无法解析时直接报错，不做模糊猜测。
    """
    text = str(raw or '').strip()
    if not text:
        day = date.today() if today is None else today
        stamp = datetime(day.year, day.month, day.day)
        return stamp, f'默认取运行当日 {stamp:%Y-%m-%d}'
    for fmt in REFERENCE_DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        return datetime(parsed.year, parsed.month, parsed.day), f'显式指定 {text}'
    where = f'{env_name} 或界面「{label}」' if env_name else f'界面「{label}」'
    raise ValueError(f'{label}无法解析：{text}。请在 {where} 中使用 YYYY-MM-DD 格式'
                     f'（如 2026-07-31），留空表示使用运行当日。')


def retry_permission(operation):
    for attempt in range(len(RETRY_DELAYS) + 1):
        try:
            return operation()
        except PermissionError:
            if attempt == len(RETRY_DELAYS):
                raise
            time.sleep(RETRY_DELAYS[attempt])


class SafeWriteError(OSError):
    def __init__(self, target, temporary, validated):
        self.target = str(target)
        self.temporary = str(temporary) if temporary else None
        self.validated = validated
        detail = ('已校验的可恢复结果' if validated else '未校验的临时文件（不可直接当作结果）')
        location = f'{detail}：{temporary}' if temporary else '尚未创建临时文件'
        super().__init__(f'安全保存失败，未替换目标：{target}；{location}。请关闭 Excel/WPS 或检查磁盘后重试。')


def atomic_write(path, writer, validator, protected_paths=()):
    """Never open the destination for writing. Keep failed artifacts for recovery."""
    target = normalized_path(path)
    for source in protected_paths:
        require_distinct(source, target)
    temporary = None
    validated = False
    try:
        retry_permission(lambda: target.parent.mkdir(parents=True, exist_ok=True))
        fd, name = retry_permission(lambda: tempfile.mkstemp(
            prefix=target.stem + '.pending-', suffix=target.suffix, dir=target.parent))
        os.close(fd)
        temporary = Path(name)
        retry_permission(lambda: writer(temporary))
        retry_permission(lambda: validator(temporary))
        validated = True

        def replace():
            # Check again immediately before replacement (including samefile aliases).
            for source in protected_paths:
                require_distinct(source, target)
            os.replace(temporary, target)

        retry_permission(replace)
    except Exception:
        raise SafeWriteError(target, temporary, validated) from None
    return target


def write_json(path, payload, protected_paths=()):
    encoded = json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2)

    def write(temporary):
        with temporary.open('w', encoding='utf-8', newline='\n') as stream:
            stream.write(encoded)
            stream.flush()
            os.fsync(stream.fileno())

    def validate(temporary):
        with temporary.open(encoding='utf-8') as stream:
            if json.load(stream) != payload:
                raise ValueError('JSON 校验不一致')

    return atomic_write(path, write, validate, protected_paths)


def write_excel(path, sheets, protected_paths=()):
    """Write plain report sheets, validate every XML/row and the expected shape."""
    import pandas as pd
    from openpyxl import load_workbook

    def write(temporary):
        # An explicit file context also closes the handle if ExcelWriter.close fails.
        with temporary.open('wb') as stream:
            with pd.ExcelWriter(stream, engine='openpyxl') as writer:
                # Ensure close() can cleanly close its ZIP even if the first sheet fails.
                placeholder = writer.book.create_sheet()
                while placeholder.title in sheets:
                    placeholder.title += '_'
                for name, frame in sheets.items():
                    frame.to_excel(writer, sheet_name=name, index=False)
                    # Report/model text must not become executable Excel formulas.
                    for row in writer.sheets[name]:
                        for cell in row:
                            if isinstance(cell.value, str):
                                cell.data_type = 's'
                if len(writer.book.worksheets) > 1:
                    writer.book.remove(placeholder)
            stream.flush()
            os.fsync(stream.fileno())

    def matches(expected, actual):
        if pd.isna(expected) or (isinstance(expected, str) and expected == ''):
            return actual is None
        if isinstance(expected, str):
            return expected == actual  # Detect Excel's silent 32767-character truncation.
        if isinstance(expected, bool):
            return type(actual) is bool and expected == actual
        if isinstance(expected, Real):
            if not math.isfinite(expected):
                return actual == str(expected)  # pandas represents infinity as text.
            return (isinstance(actual, (int, float)) and not isinstance(actual, bool)
                    and math.isclose(expected, actual, rel_tol=1e-14, abs_tol=1e-14))
        if isinstance(expected, (date, datetime)):
            return (isinstance(actual, (date, datetime))
                    and abs(pd.Timestamp(expected) - pd.Timestamp(actual)) <= pd.Timedelta(milliseconds=1))
        return expected == actual

    def validate(temporary):
        with ZipFile(temporary) as archive:
            if archive.testzip() is not None:
                raise ValueError('工作簿 ZIP 校验失败')
        with temporary.open('rb') as stream:
            book = load_workbook(stream, read_only=True, data_only=False)
            try:
                if book.sheetnames != list(sheets):
                    raise ValueError('结果工作表不完整')
                for name, frame in sheets.items():
                    sheet = book[name]
                    if sheet.max_row != len(frame) + 1 or sheet.max_column != len(frame.columns):
                        raise ValueError('结果行列数不一致')
                    rows = sheet.iter_rows(values_only=True)
                    if tuple(next(rows)) != tuple(str(c) for c in frame.columns):
                        raise ValueError('结果列名不一致')
                    # Verify every cell, not just dimensions or the ZIP header.
                    for expected in frame.itertuples(index=False, name=None):
                        actual = next(rows, None)
                        if actual is None or not all(matches(a, b) for a, b in zip(expected, actual)):
                            raise ValueError('结果单元格内容校验不一致')
                    if next(rows, None) is not None:
                        raise ValueError('结果行数不一致')
            finally:
                book.close()

    return atomic_write(path, write, validate, protected_paths)
