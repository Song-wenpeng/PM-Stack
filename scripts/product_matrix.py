# -*- coding: utf-8 -*-
"""
产品矩阵生成脚本

从竞品数据 Excel 中按品牌/场景分列，下载商品主图并生成可视化产品矩阵。

用法:
    python product_matrix.py

通过环境变量配置参数：
    MATRIX_EXCEL_FILE       源 Excel 文件路径
    MATRIX_SOURCE_SHEET     源数据 Sheet 名
    MATRIX_OUTPUT_FILE      输出文件路径
    MATRIX_IMAGE_DIR        图片缓存目录
    MATRIX_BRAND_COL        分组列名（默认: 使用场景）
    MATRIX_TARGET_BRANDS    目标品牌，逗号分隔（留空=全部）
"""

import os
import re
import sys
import time
import requests
import pandas as pd
from pathlib import Path
from io import BytesIO
from PIL import Image as PILImage

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side


# ======================
# 配置（从环境变量读取）
# ======================

EXCEL_FILE = os.getenv("MATRIX_EXCEL_FILE", "")
SOURCE_SHEET = os.getenv("MATRIX_SOURCE_SHEET", "")
OUTPUT_FILE = os.getenv("MATRIX_OUTPUT_FILE", "")
IMAGE_DIR = os.getenv("MATRIX_IMAGE_DIR", os.path.join(os.path.expanduser("~"), "AppData", "Local", "Temp", "matrix_images"))
BRAND_COL_NAME = os.getenv("MATRIX_BRAND_COL", "使用场景")
ASIN_COL_NAME = "ASIN"
DATE_COL_NAME = "上架时间"
IMAGE_URL_COL_NAME = "商品主图"

_target = os.getenv("MATRIX_TARGET_BRANDS", "").strip()
TARGET_BRANDS = [b.strip() for b in _target.split(",") if b.strip()] if _target else []

IMG_HEIGHT_PX = 70
IMAGE_ROW_HEIGHT = 60
TEXT_ROW_HEIGHT = 26
BRAND_COL_WIDTH = 16
SHOW_INFO_UNDER_IMAGE = True
DELETE_TEMP_IMAGES_AFTER_SAVE = True


# ======================
# 工具函数
# ======================

def clean_filename(name):
    name = str(name).strip()
    return re.sub(r'[\\/:*?"<>|]', "_", name)


def safe_str(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def download_and_resize_image(url, save_path, target_height=70, max_retries=3):
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Referer": "https://www.amazon.co.uk/"
    }
    last_error = None
    for attempt in range(1, max_retries + 1):
        try:
            response = requests.get(url, headers=headers, timeout=30)
            response.raise_for_status()
            img = PILImage.open(BytesIO(response.content)).convert("RGB")
            w, h = img.size
            if h <= 0:
                raise ValueError("图片高度异常")
            ratio = target_height / h
            target_width = max(1, int(w * ratio))
            img = img.resize((target_width, target_height), PILImage.LANCZOS)
            img.save(save_path, format="JPEG", quality=85)
            return save_path
        except Exception as e:
            last_error = e
            print(f"图片下载失败，第 {attempt}/{max_retries} 次重试：{url}，错误：{e}")
            time.sleep(2 * attempt)
    raise last_error


def get_or_download_image(asin, url, image_dir):
    image_dir = Path(image_dir)
    image_dir.mkdir(parents=True, exist_ok=True)
    img_path = image_dir / f"{clean_filename(asin)}.jpg"
    if img_path.exists() and img_path.stat().st_size > 0:
        return img_path
    if not url or not url.startswith("http"):
        return None
    download_and_resize_image(url, img_path, target_height=IMG_HEIGHT_PX)
    time.sleep(0.2)
    return img_path


def format_date(value):
    if pd.isna(value):
        return ""
    try:
        dt = pd.to_datetime(value, errors="coerce")
        if pd.isna(dt):
            return str(value)
        return dt.strftime("%Y-%m-%d")
    except Exception:
        return str(value)


def delete_used_images(image_paths):
    deleted = 0
    for p in image_paths:
        try:
            p = Path(p)
            if p.exists() and p.is_file():
                p.unlink()
                deleted += 1
        except Exception as e:
            print(f"删除图片失败：{p}，错误：{e}")
    print(f"已删除本次使用的临时图片：{deleted} 张")


# ======================
# 主程序
# ======================

def main():
    if not EXCEL_FILE:
        print("错误: 未指定源 Excel 文件（MATRIX_EXCEL_FILE）")
        sys.exit(1)
    if not SOURCE_SHEET:
        print("错误: 未指定源数据 Sheet（MATRIX_SOURCE_SHEET）")
        sys.exit(1)
    if not OUTPUT_FILE:
        print("错误: 未指定输出文件路径（MATRIX_OUTPUT_FILE）")
        sys.exit(1)

    used_image_paths = []
    image_dir = Path(IMAGE_DIR)
    image_dir.mkdir(parents=True, exist_ok=True)

    # 1. 读取源数据
    print(f"读取: {EXCEL_FILE} / {SOURCE_SHEET}")
    df = pd.read_excel(EXCEL_FILE, sheet_name=SOURCE_SHEET)
    df.columns = df.columns.astype(str).str.strip()

    required_cols = [BRAND_COL_NAME, ASIN_COL_NAME, DATE_COL_NAME, IMAGE_URL_COL_NAME]
    for col in required_cols:
        if col not in df.columns:
            raise ValueError(f"源表缺少列：{col}，当前列名为：{df.columns.tolist()}")

    df[BRAND_COL_NAME] = df[BRAND_COL_NAME].apply(safe_str)
    df[ASIN_COL_NAME] = df[ASIN_COL_NAME].apply(safe_str)
    df[IMAGE_URL_COL_NAME] = df[IMAGE_URL_COL_NAME].apply(safe_str)
    df["_上架时间_dt"] = pd.to_datetime(df[DATE_COL_NAME], errors="coerce")

    df = df[(df[BRAND_COL_NAME] != "") & (df[ASIN_COL_NAME] != "")].copy()

    # 2. 筛选品牌
    if TARGET_BRANDS:
        target_set = set(TARGET_BRANDS)
        df = df[df[BRAND_COL_NAME].isin(target_set)].copy()

    if df.empty:
        print("筛选后没有可处理的数据，请检查品牌名或表格内容。")
        sys.exit(1)

    brands = sorted(df[BRAND_COL_NAME].dropna().unique().tolist())
    print(f"将生成产品矩阵，品牌数量：{len(brands)} — {brands}")

    # 3. 创建 Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "产品矩阵"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, size=12)
    info_font = Font(size=9, color="666666")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    failed_records = []

    # 4. 每个品牌一列
    for brand_idx, brand in enumerate(brands, start=1):
        col_idx = brand_idx
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = BRAND_COL_WIDTH

        brand_df = df[df[BRAND_COL_NAME] == brand].copy()
        brand_df = brand_df.sort_values(by="_上架时间_dt", ascending=True, na_position="last").reset_index(drop=True)

        ws.cell(row=1, column=col_idx).value = f"{brand}\n共{len(brand_df)}款"
        ws.cell(row=1, column=col_idx).fill = header_fill
        ws.cell(row=1, column=col_idx).font = header_font
        ws.cell(row=1, column=col_idx).alignment = center_align
        ws.cell(row=1, column=col_idx).border = border

        current_row = 2
        for _, item in brand_df.iterrows():
            asin = safe_str(item[ASIN_COL_NAME])
            img_url = safe_str(item[IMAGE_URL_COL_NAME])
            launch_date = format_date(item[DATE_COL_NAME])

            try:
                img_path = get_or_download_image(asin, img_url, image_dir)

                ws.row_dimensions[current_row].height = IMAGE_ROW_HEIGHT
                ws.cell(row=current_row, column=col_idx).alignment = center_align
                ws.cell(row=current_row, column=col_idx).border = border

                if img_path and Path(img_path).exists():
                    xl_img = XLImage(str(img_path))
                    ws.add_image(xl_img, f"{col_letter}{current_row}")
                    used_image_paths.append(Path(img_path))
                else:
                    ws.cell(row=current_row, column=col_idx).value = "无图"

                current_row += 1

                if SHOW_INFO_UNDER_IMAGE:
                    ws.row_dimensions[current_row].height = TEXT_ROW_HEIGHT
                    ws.cell(row=current_row, column=col_idx).value = f"{asin}\n{launch_date}"
                    ws.cell(row=current_row, column=col_idx).font = info_font
                    ws.cell(row=current_row, column=col_idx).alignment = center_align
                    ws.cell(row=current_row, column=col_idx).border = border
                    current_row += 1

                print(f"插入成功：品牌={brand}，ASIN={asin}，上架时间={launch_date}")

            except Exception as e:
                print(f"插入失败：品牌={brand}，ASIN={asin}，错误={e}")
                failed_records.append({
                    "品牌": brand, "ASIN": asin, "上架时间": launch_date,
                    "商品主图链接": img_url, "错误信息": str(e),
                })
                ws.row_dimensions[current_row].height = IMAGE_ROW_HEIGHT
                ws.cell(row=current_row, column=col_idx).value = f"失败\n{asin}"
                ws.cell(row=current_row, column=col_idx).alignment = center_align
                ws.cell(row=current_row, column=col_idx).border = border
                current_row += 1
                if SHOW_INFO_UNDER_IMAGE:
                    ws.row_dimensions[current_row].height = TEXT_ROW_HEIGHT
                    ws.cell(row=current_row, column=col_idx).value = launch_date
                    ws.cell(row=current_row, column=col_idx).alignment = center_align
                    ws.cell(row=current_row, column=col_idx).border = border
                    current_row += 1

    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"

    # 失败记录 Sheet
    if failed_records:
        fail_ws = wb.create_sheet("图片下载失败记录")
        headers = ["品牌", "ASIN", "上架时间", "商品主图链接", "错误信息"]
        for ci, h in enumerate(headers, start=1):
            fail_ws.cell(row=1, column=ci).value = h
        for ri, rec in enumerate(failed_records, start=2):
            for ci, h in enumerate(headers, start=1):
                fail_ws.cell(row=ri, column=ci).value = rec.get(h, "")

    # 保存
    output_path = Path(OUTPUT_FILE)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    if DELETE_TEMP_IMAGES_AFTER_SAVE:
        delete_used_images(used_image_paths)

    print("=" * 60)
    print(f"产品矩阵生成完成：{output_path}")
    if failed_records:
        print(f"失败记录共 {len(failed_records)} 条，已写入 sheet：图片下载失败记录")


if __name__ == "__main__":
    main()
