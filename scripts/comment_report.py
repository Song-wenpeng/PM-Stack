"""Reader-facing views for a completed comment analysis workbook.

The AI classification sheets remain the source of category totals.  This
module adds review-level counts and presentation without another AI call.
"""

from __future__ import annotations

import ast
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.views import Selection


NAVY = "263449"
BLUE = "3758A8"
RED = "DC5757"
GREEN = "219477"
AMBER = "E8A433"
LIGHT = "F3F6FA"
GREY = "667085"


def parse_tags(value):
    """Read Step 1 JSON/Python lists, allowing plain text in older exports."""
    if value is None or value == "":
        return []
    if isinstance(value, (list, tuple)):
        return [str(x).strip() for x in value if str(x).strip()]
    if not isinstance(value, str):
        return [str(value)]
    text = value.strip()
    if not text or text in ("[]", "nan", "None", "null"):
        return []
    if text.startswith("["):
        for parser in (json.loads, ast.literal_eval):
            try:
                parsed = parser(text)
                if isinstance(parsed, (list, tuple)):
                    return [str(x).strip() for x in parsed if str(x).strip()]
            except (ValueError, SyntaxError, TypeError, json.JSONDecodeError):
                pass
        if text.endswith("]"):
            inner = text[1:-1]
            return [part.strip().strip('"').strip("'") for part in inner.split(",")
                    if part.strip()]
    if any(separator in text for separator in ("；", ";", "、", "|")):
        return [part.strip() for part in re.split(r"[；;、|]", text) if part.strip()]
    return [text]


def _integer(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _internal_link(cell, sheet_name, address):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet_name}'!{address}")


def _rating(value):
    try:
        number = int(str(value).strip().split()[0])
        return number if 1 <= number <= 5 else None
    except (IndexError, TypeError, ValueError):
        return None


def _taxonomy(workbook, dimension):
    category_sheet = workbook[f"{dimension}_AI归类"]
    normalization_sheet = workbook[f"{dimension}_标签标准化"]
    normalization = {
        str(row[0]).strip(): str(row[1]).strip()
        for row in normalization_sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None and row[1] is not None
    }
    category_counts = {}
    label_to_category = {}
    representatives = {}
    mode = ""
    for row in category_sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 8 or not row[1]:
            continue
        mode = str(row[0] or mode)
        category = str(row[1]).strip()
        category_counts[category] = _integer(row[2])
        if row[5]:
            label_to_category[str(row[5]).strip()] = category
        representatives[category] = str(row[7] or "")
    return normalization, label_to_category, category_counts, representatives, mode


def collect_metrics(workbook):
    """Return category mentions and de-duplicated source-row counts."""
    raw = workbook["合并数据"]
    headers = {str(cell.value): cell.column for cell in raw[1] if cell.value}
    row_count = raw.max_row - 1
    ratings = Counter(_rating(raw.cell(row_idx, headers.get("星级", 1)).value)
                      for row_idx in range(2, raw.max_row + 1))
    metrics = {}
    dimensions = [sheet.title[:-len("_AI归类")] for sheet in workbook
                  if sheet.title.endswith("_AI归类")]
    for dim in dimensions:
        if dim not in headers or f"{dim}_AI归类" not in workbook:
            continue
        normalization, lookup, counts, reps, mode = _taxonomy(workbook, dim)
        per_category = defaultdict(set)
        per_rating = defaultdict(lambda: defaultdict(set))
        evidence = defaultdict(list)
        unmapped = set()
        for row_idx in range(2, raw.max_row + 1):
            raw_tags = parse_tags(raw.cell(row_idx, headers[dim]).value)
            rating = _rating(raw.cell(row_idx, headers.get("星级", 1)).value)
            matched = set()
            for tag in raw_tags:
                standard = normalization.get(tag, tag)
                category = lookup.get(standard)
                if category:
                    matched.add(category)
                else:
                    unmapped.add(standard)
            for category in matched:
                per_category[category].add(row_idx)
                if rating is not None:
                    per_rating[category][rating].add(row_idx)
                original = raw.cell(row_idx, headers.get("内容", 3)).value
                if original:
                    excerpt = " ".join(str(original).split())[:170]
                    evidence[category].append((row_idx, excerpt, rating))
        metrics[dim] = {
            "counts": counts,
            "total_mentions": sum(counts.values()),
            "review_rows": per_category,
            "rating_rows": per_rating,
            "representatives": reps,
            "evidence": evidence,
            "mode": mode,
            "unmapped": unmapped,
        }
    return row_count, ratings, metrics


def _role(metrics, keyword, field_roles=None):
    role = {"负向": "negative", "正向": "positive", "改进": "improvement"}.get(keyword)
    if role and field_roles:
        configured = next((dimension for dimension in metrics if field_roles.get(dimension) == role), None)
        if configured:
            return configured
    return next((dimension for dimension in metrics if keyword in dimension), None)


def _title(ws, text, subtitle, end_col=10):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", size=17, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 31
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=10, color=GREY)
    ws.row_dimensions[2].height = 24
    for col in range(1, end_col + 1):
        ws.cell(3, col).border = Border(bottom=Side(style="thin", color="D9E0EA"))


def _header(ws, row, columns):
    for col, value in enumerate(columns, 1):
        cell = ws.cell(row, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 25


def _style_grid(ws, start_row, end_row, end_col):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=NAVY)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = Border(bottom=Side(style="hair", color="E1E7EF"))
        ws.row_dimensions[row[0].row].height = 23


def _add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def _summary(workbook, row_count, metrics):
    ws = workbook["摘要汇总"]
    ws.delete_rows(1, ws.max_row)
    headers = ["维度", "一级类别", "提及次数", "维度内提及占比", "去重评论数", "评论触及率", "代表标签", "归类模式", "累计提及占比"]
    _header(ws, 1, headers)
    category_rows = defaultdict(dict)
    index = 2
    for dim in metrics:
        item = metrics[dim]
        total = item["total_mentions"]
        cumulative = 0
        for category, mentions in sorted(item["counts"].items(), key=lambda pair: (-pair[1], pair[0])):
            cumulative += mentions
            reviews = len(item["review_rows"].get(category, ()))
            values = [
                dim, category, mentions, mentions / total if total else None,
                reviews, reviews / row_count if row_count else None,
                item["representatives"].get(category, ""), item["mode"],
                cumulative / total if total else None,
            ]
            for col, value in enumerate(values, 1):
                ws.cell(index, col, value)
            category_rows[dim][category] = index
            index += 1
    _style_grid(ws, 2, index - 1, 9)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{index - 1}"
    if index > 2:
        _add_table(ws, f"A1:I{index - 1}", "CommentSummary")
    for col, width in {"A": 20, "B": 25, "C": 13, "D": 18, "E": 15,
                       "F": 17, "G": 46, "H": 30, "I": 18}.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["H"].hidden = True
    for row in range(2, index):
        for col in (4, 6, 9):
            ws.cell(row, col).number_format = "0.0%"
    for dim, rows in category_rows.items():
        if not rows:
            continue
        positions = list(rows.values())
        color = RED if "负向" in dim else GREEN if "正向" in dim else AMBER if "改进" in dim else BLUE
        ws.conditional_formatting.add(
            f"D{min(positions)}:D{max(positions)}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                        color=color, showValue=True),
        )
    return category_rows


def _cross(workbook, ratings, metrics, field_roles=None):
    ws = workbook.create_sheet("星级问题", 2)
    low_total = ratings[1] + ratings[2]
    _title(ws, "星级与负向问题", f"按评论去重；1–2 星样本 {low_total} 条；同一评论可涉及多个类别", 9)
    headers = ["负向一级类别", "1 星", "2 星", "3 星", "4 星", "5 星", "1–2 星去重评论数", "低星触及率", "全部星级去重评论数"]
    _header(ws, 5, headers)
    item = metrics.get(_role(metrics, "负向", field_roles), {})
    rows = []
    for category in sorted(item.get("counts", {}), key=lambda name: -item["counts"][name]):
        rating_rows = item["rating_rows"].get(category, {})
        values = [len(rating_rows.get(rating, ())) for rating in range(1, 6)]
        low_count = values[0] + values[1]
        all_count = len(item["review_rows"].get(category, ()))
        rows.append([category, *values, low_count, low_count / low_total if low_total else None, all_count])
    for row_idx, row in enumerate(rows, 6):
        for col, value in enumerate(row, 1):
            ws.cell(row_idx, col, value)
        ws.cell(row_idx, 8).number_format = "0.0%"
    if rows:
        _style_grid(ws, 6, 5 + len(rows), 9)
        ws.auto_filter.ref = f"A5:I{5 + len(rows)}"
        _add_table(ws, f"A5:I{5 + len(rows)}", "RatingIssues")
        ws.conditional_formatting.add(
            f"H6:H{5 + len(rows)}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                        color=RED, showValue=True),
        )
    ws.freeze_panes = "B6"
    for col, width in {"A": 27, "B": 10, "C": 10, "D": 10, "E": 10, "F": 10,
                       "G": 20, "H": 15, "I": 23}.items():
        ws.column_dimensions[col].width = width
    ws["A4"] = "按评论去重；筛选类别后，可在“合并数据”查看评论原文。"
    ws["A4"].font = Font(name="Arial", size=10, color=GREY)
    return ws


def _evidence(workbook, metrics, field_roles=None):
    """One row per review/category, with only that review's secondary topics."""
    ws = workbook.create_sheet("问题证据", 3)
    _title(ws, "负向问题证据", "同一评论在同一一级类别下只列一次；提取到的负向关注点保留该评论全部原始标签；点击评论内容可返回原始数据", 7)
    _header(ws, 5, ["负向一级类别", "负向二级类别", "提取到的负向关注点", "星级", "原始行号", "评论内容", "原文链接"])
    dimension = _role(metrics, "负向", field_roles)
    item = metrics.get(dimension, {})
    raw = workbook["合并数据"]
    headers = {str(cell.value).strip(): cell.column for cell in raw[1] if cell.value}
    rating_col = headers.get("星级")
    content_col = headers.get("内容")
    normalization, _, _, _, _ = _taxonomy(workbook, dimension) if dimension else ({}, {}, {}, {}, "")
    secondary_lookup = defaultdict(set)
    if dimension:
        for values in workbook[f"{dimension}_AI归类"].iter_rows(min_row=2, values_only=True):
            if len(values) >= 6 and values[1] and values[5]:
                category, label = str(values[1]).strip(), str(values[5]).strip()
                secondary_lookup[(category, label)].add(str(values[3]).strip() if values[3] else label)

    def rating_at(row):
        return _rating(raw.cell(row, rating_col).value) if rating_col else None

    first_row = {}
    index = 6
    for category in sorted(item.get("counts", {}), key=lambda name: -item["counts"][name]):
        rows = sorted(item["review_rows"].get(category, ()),
                      key=lambda row: (rating_at(row) or 6, row))
        for source_row in rows:
            first_row.setdefault(category, index)
            content = " ".join(str(raw.cell(source_row, content_col).value or "").split()) if content_col else ""
            secondary = set()
            extracted_tags = parse_tags(raw.cell(source_row, headers[dimension]).value)
            for tag in extracted_tags:
                secondary.update(secondary_lookup.get((category, normalization.get(tag, tag)), ()))
            ws.cell(index, 1, category)
            ws.cell(index, 2, "；".join(sorted(secondary)))
            ws.cell(index, 3, "；".join(extracted_tags))
            ws.cell(index, 4, rating_at(source_row))
            ws.cell(index, 5, source_row)
            quote = ws.cell(index, 6, content)
            link = ws.cell(index, 7, "查看原始评论")
            address = f"{get_column_letter(content_col or headers[dimension])}{source_row}"
            _internal_link(quote, "合并数据", address)
            _internal_link(link, "合并数据", address)
            index += 1
    if index > 6:
        _style_grid(ws, 6, index - 1, 7)
        for row_idx in range(6, index):
            ws.cell(row_idx, 7).font = Font(name="Arial", size=10, color=BLUE, underline="single")
        ws.auto_filter.ref = f"A5:G{index - 1}"
        _add_table(ws, f"A5:G{index - 1}", "IssueEvidence")
    ws.freeze_panes = "F6"
    for col, width in {"A": 28, "B": 40, "C": 48, "D": 9, "E": 13, "F": 92, "G": 18}.items():
        ws.column_dimensions[col].width = width
    ws.sheet_properties.tabColor = RED
    return first_row


def repair_evidence_report(path, output=None, field_roles=None):
    """Rebuild only the evidence sheet using stored labels, with no API calls."""
    path = Path(path)
    workbook = load_workbook(path)
    _, _, metrics = collect_metrics(workbook)
    if not _role(metrics, "负向", field_roles):
        raise ValueError("报告中没有负向归类数据，无法重建问题证据。")
    position = workbook.index(workbook["问题证据"]) if "问题证据" in workbook else 3
    if "问题证据" in workbook:
        del workbook["问题证据"]
    _evidence(workbook, metrics, field_roles)
    sheet = workbook["问题证据"]
    workbook.move_sheet(sheet, offset=position - workbook.index(sheet))
    destination = Path(output) if output else path
    workbook.save(destination)
    workbook.close()
    return destination


def _bar_chart(ws, data_sheet, first_row, last_row, title, color, anchor):
    if last_row < first_row:
        return
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.title = title
    chart.x_axis.title = "提及次数"
    chart.legend = None
    chart.width = 11.5
    chart.height = 7.2
    chart.add_data(Reference(data_sheet, min_col=3, min_row=first_row, max_row=min(last_row, first_row + 6)))
    chart.set_categories(Reference(data_sheet, min_col=2, min_row=first_row, max_row=min(last_row, first_row + 6)))
    chart.series[0].graphicalProperties.solidFill = color
    chart.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(chart, anchor)


def _dimension_color(dimension, role=None):
    if role == "negative" or "负向" in dimension:
        return RED, "FCE9EA"
    if role == "positive" or "正向" in dimension:
        return GREEN, "E6F5F0"
    if role == "improvement" or "改进" in dimension:
        return AMBER, "FFF1DA"
    return BLUE, "E8EEFC"


def _write_dimension_table(ws, dimension, item, title_row, link_to_detail=False, role=None):
    """Render every primary category in the reference table style."""
    color, track = _dimension_color(dimension, role)
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=6)
    heading = ws.cell(title_row, 1, dimension)
    heading.font = Font(name="Arial", size=12, bold=True, color=color)
    heading.fill = PatternFill("solid", fgColor=track)
    heading.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[title_row].height = 27
    for col, value in enumerate(("维度", "一级类别", "提及次数", "占比", "热度", "代表性标签"), 1):
        cell = ws.cell(title_row + 1, col, value)
        cell.fill = PatternFill("solid", fgColor="F3F4F6")
        cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
        cell.alignment = Alignment(vertical="center", horizontal="center" if col in (3, 4, 5) else "left",
                                   indent=0 if col in (3, 4, 5) else 1)
        cell.border = Border(bottom=Side(style="thin", color="D7DDE5"))
    ws.row_dimensions[title_row + 1].height = 27
    ranked = sorted(item["counts"].items(), key=lambda pair: (-pair[1], pair[0]))
    largest = ranked[0][1] if ranked else 0
    start = title_row + 2
    for offset, (category, mentions) in enumerate(ranked):
        row = start + offset
        ws.cell(row, 1, dimension)
        category_cell = ws.cell(row, 2, category)
        if link_to_detail:
            _internal_link(category_cell, dimension, f"B{6 + offset}")
        ws.cell(row, 3, mentions)
        ws.cell(row, 4, mentions / item["total_mentions"] if item["total_mentions"] else None)
        heat = ws.cell(row, 5, "█" * max(1, round(16 * mentions / largest)) if largest else "")
        heat.fill = PatternFill("solid", fgColor=track)
        ws.cell(row, 6, "；".join(item["representatives"].get(category, "").split("；")[:3]))
        for col in range(1, 7):
            cell = ws.cell(row, col)
            cell.font = Font(name="Arial", size=10, color=BLUE if col == 2 else NAVY)
            cell.alignment = Alignment(vertical="center", horizontal="right" if col in (3, 4) else "left",
                                       indent=1 if col in (1, 2, 6) else 0, shrink_to_fit=col == 6)
            cell.border = Border(bottom=Side(style="thin", color="E0E4E9"))
        heat.font = Font(name="Arial", size=17, bold=True, color=color)
        heat.alignment = Alignment(vertical="center", indent=1, shrink_to_fit=True)
        share = ws.cell(row, 4).value
        ws.cell(row, 4).number_format = "0.0%" if share is not None and share < 0.01 else "0%"
        ws.row_dimensions[row].height = 29
    if len(ranked) > 1:
        ws.merge_cells(start_row=start, start_column=1, end_row=start + len(ranked) - 1, end_column=1)
    return start + len(ranked)

def _dashboard(workbook, row_count, ratings, metrics, category_rows, product_name=None, field_roles=None):
    ws = workbook.create_sheet("VOC看板", 0)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.showGridLines = False
    for col, width in {"A": 17, "B": 26, "C": 13, "D": 11, "E": 29, "F": 43}.items():
        ws.column_dimensions[col].width = width
    title = f"VOC 摘要汇总 · {product_name or '评论分析'}（{row_count} 条评论）"
    ws.merge_cells("A1:E1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["F1"] = "可筛选 · 可追溯"
    ws["F1"].font = Font(name="Arial", size=9, color=GREY)
    ws["F1"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:F2")
    ws["A2"] = "占比＝该类别提及次数 ÷ 同维度全部一级类别提及次数；热度在每个维度内缩放。"
    ws["A2"].font = Font(name="Arial", size=9, color=GREY)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 9
    for col, label in enumerate(("维度", "一级类别", "提及次数", "占比", "热度", "代表性标签"), 1):
        cell = ws.cell(4, col, label)
        cell.fill = PatternFill("solid", fgColor="F3F4F6")
        cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
        cell.alignment = Alignment(vertical="center", horizontal="center" if col in (3, 4, 5) else "left", indent=0 if col in (3, 4, 5) else 1)
        cell.border = Border(top=Side(style="thin", color="D7DDE5"),
                             bottom=Side(style="thin", color="D7DDE5"))
    ws.row_dimensions[4].height = 29
    negative_dim = _role(metrics, "负向", field_roles)
    positive_dim = _role(metrics, "正向", field_roles)
    improvement_dim = _role(metrics, "改进", field_roles)
    role_specs = [
        (negative_dim, "负向关注点", RED, "FCE9EA"),
        (positive_dim, "正向关注点", GREEN, "E6F5F0"),
        (improvement_dim, "改进需求", AMBER, "FFF1DA"),
    ]
    role_specs = [entry for entry in role_specs if entry[0] and metrics[entry[0]]["counts"]]
    if not role_specs:
        role_specs = [(dim, dim, BLUE, "E8EEFC") for dim in list(metrics)[:3]]
    next_row = 5
    for dim, label, color, track in role_specs:
        item = metrics[dim]
        leaders = sorted(item["counts"].items(), key=lambda pair: (-pair[1], pair[0]))[:2]
        if not leaders:
            continue
        first_row = next_row
        for rank, (category, mentions) in enumerate(leaders):
            row_idx = next_row
            share = mentions / item["total_mentions"] if item["total_mentions"] else None
            ws.cell(row_idx, 2, category)
            ws.cell(row_idx, 3, mentions)
            ws.cell(row_idx, 4, share)
            bar_length = max(1, round(16 * mentions / leaders[0][1]))
            heat = ws.cell(row_idx, 5, "█" * bar_length)
            heat.fill = PatternFill("solid", fgColor=track)
            representative = item["representatives"].get(category, "")
            ws.cell(row_idx, 6, "；".join(representative.split("；")[:3]))
            _internal_link(ws.cell(row_idx, 2), dim, f"B{6 + rank}")
            for col in range(1, 7):
                cell = ws.cell(row_idx, col)
                cell.font = Font(name="Arial", size=10, color=BLUE if col == 2 else NAVY)
                cell.alignment = Alignment(vertical="center", horizontal="right" if col in (3, 4) else "left",
                                           indent=1 if col in (2, 6) else 0, shrink_to_fit=col == 6)
                cell.border = Border(bottom=Side(style="thin", color="E0E4E9"))
            heat.font = Font(name="Arial", size=17, bold=True, color=color)
            heat.alignment = Alignment(vertical="center", horizontal="left", indent=1,
                                       shrink_to_fit=True)
            ws.cell(row_idx, 4).number_format = "0%"
            ws.row_dimensions[row_idx].height = 30
            next_row += 1
        ws.cell(first_row, 1, label)
        ws.cell(first_row, 1).font = Font(name="Arial", size=10, bold=True, color=color)
        ws.cell(first_row, 1).alignment = Alignment(vertical="center", indent=1)
        if next_row - first_row > 1:
            ws.merge_cells(start_row=first_row, start_column=1, end_row=next_row - 1, end_column=1)
    footer_row = next_row
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=6)
    ws.cell(footer_row, 1, "下方列出全部目标字段；点击类别进入对应字段明细。")
    ws.cell(footer_row, 1).font = Font(name="Arial", size=9, color=GREY)
    ws.cell(footer_row, 1).fill = PatternFill("solid", fgColor=LIGHT)
    ws.cell(footer_row, 1).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[footer_row].height = 25
    negative_item = metrics.get(negative_dim, {})
    if negative_item.get("counts"):
        top_three = sum(sorted(negative_item["counts"].values(), reverse=True)[:3])
        total = negative_item["total_mentions"]
        ws.cell(footer_row + 1, 1, f"负向前三类合计 {top_three}/{total} 次提及（{top_three / total:.1%}）。")
        ws.cell(footer_row + 1, 1).font = Font(name="Arial", size=10, bold=True, color=RED)
    detail_row = footer_row + 4
    ws.cell(detail_row - 1, 1, "目标字段完整分布").font = Font(name="Arial", size=12, bold=True, color=NAVY)
    for dim, item in metrics.items():
        detail_row = _write_dimension_table(ws, dim, item, detail_row, link_to_detail=True,
                                            role=(field_roles or {}).get(dim)) + 2
    ws.sheet_properties.tabColor = BLUE
    return ws


def _source_sheet(workbook, metrics):
    ws = workbook["合并数据"]
    original = workbook.copy_worksheet(ws)
    original.title = "标签原值"
    original.sheet_state = "hidden"
    headers = {str(cell.value): cell.column for cell in ws[1] if cell.value}
    for dim in metrics:
        col_idx = headers.get(dim)
        if col_idx is None:
            continue
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, col_idx, "；".join(parse_tags(ws.cell(row_idx, col_idx).value)))
    _header(ws, 1, [ws.cell(1, col).value for col in range(1, ws.max_column + 1)])
    ws.auto_filter.ref = ws.dimensions
    # Start the source at a neutral viewport so WPS can scroll to link targets.
    ws.freeze_panes = None
    ws.sheet_view.topLeftCell = "A1"
    ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    ws.sheet_view.tabSelected = False
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24 if col >= 6 else 19
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 65
    ws.column_dimensions["E"].width = 65
    for row in range(2, ws.max_row + 1):
        rating = _rating(ws.cell(row, headers.get("星级", 1)).value)
        if rating is not None:
            ws.cell(row, headers["星级"]).font = Font(name="Arial", size=10, bold=True,
                                                       color=RED if rating <= 2 else GREEN if rating >= 4 else GREY)
    ws.sheet_properties.tabColor = "8494A7"


def _dimension_sheets(workbook, row_count, metrics, field_roles=None):
    """Create one useful, traceable reader view per target field."""
    raw = workbook["合并数据"]
    raw_headers = {str(cell.value): cell.column for cell in raw[1] if cell.value}
    for dim, item in metrics.items():
        ws = workbook.create_sheet(dim)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.freeze_panes = "C6"
        ws.sheet_properties.tabColor = _dimension_color(dim, (field_roles or {}).get(dim))[0]
        for col, width in {"A": 17, "B": 26, "C": 13, "D": 11, "E": 29, "F": 43}.items():
            ws.column_dimensions[col].width = width
        ws.merge_cells("A1:F1")
        ws["A1"] = f"{dim}（{row_count} 条评论）"
        ws["A1"].font = Font(name="Arial", size=15, bold=True, color=NAVY)
        ws["A1"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 33
        ws.merge_cells("A2:F2")
        ws["A2"] = "占比＝该类别提及次数 ÷ 本字段全部一级类别提及次数；下方列出二级主题和可回查的原文。"
        ws["A2"].font = Font(name="Arial", size=9, color=GREY)
        ws["A2"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 25
        end = _write_dimension_table(ws, dim, item, 4, role=(field_roles or {}).get(dim))
        start = end + 3
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
        ws.cell(start, 1, "二级主题与评论原文").font = Font(name="Arial", size=12, bold=True, color=NAVY)
        ws.row_dimensions[start].height = 28
        for col, label in enumerate(("一级类别", "二级主题", "提及次数", "类内占比", "标准标签", "评论原文"), 1):
            cell = ws.cell(start + 1, col, label)
            cell.fill = PatternFill("solid", fgColor="F3F4F6")
            cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
            cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[start + 1].height = 27
        taxonomy = workbook[f"{dim}_AI归类"]
        themes = {}
        tag_to_theme = {}
        normalization, _, _, _, _ = _taxonomy(workbook, dim)
        for values in taxonomy.iter_rows(min_row=2, values_only=True):
            if len(values) < 7 or not values[1] or not values[3]:
                continue
            category, theme = str(values[1]).strip(), str(values[3]).strip()
            key = category, theme
            entry = themes.setdefault(key, {"count": _integer(values[4]), "tags": []})
            tag = str(values[5]).strip() if values[5] else ""
            if tag:
                if tag not in entry["tags"]:
                    entry["tags"].append(tag)
                tag_to_theme[tag] = key
        evidence = {}
        col = raw_headers.get(dim)
        if col:
            for source_row in range(2, raw.max_row + 1):
                quote = " ".join(str(raw.cell(source_row, raw_headers.get("内容", 3)).value or "").split())
                if not quote:
                    continue
                for tag in parse_tags(raw.cell(source_row, col).value):
                    key = tag_to_theme.get(normalization.get(tag, tag))
                    if key and key not in evidence:
                        evidence[key] = source_row, quote
        ranked_categories = sorted(item["counts"], key=lambda category: (-item["counts"][category], category))
        theme_row = start + 2
        first_theme = {}
        for category in ranked_categories:
            related = sorted(((key, value) for key, value in themes.items() if key[0] == category),
                             key=lambda pair: (-pair[1]["count"], pair[0][1]))
            for (name, theme), details in related:
                first_theme.setdefault(category, theme_row)
                ws.cell(theme_row, 1, name)
                ws.cell(theme_row, 2, theme)
                ws.cell(theme_row, 3, details["count"])
                total = item["counts"][name]
                ws.cell(theme_row, 4, details["count"] / total if total else None)
                ws.cell(theme_row, 4).number_format = "0%"
                ws.cell(theme_row, 5, "；".join(details["tags"]))
                if (name, theme) in evidence:
                    source_row, quote = evidence[name, theme]
                    cell = ws.cell(theme_row, 6, f"第{source_row}行 · {quote[:65]}")
                    _internal_link(cell, "合并数据", f"C{source_row}")
                    cell.hyperlink.tooltip = f"合并数据 C{source_row}"
                for col_idx in range(1, 7):
                    cell = ws.cell(theme_row, col_idx)
                    cell.font = Font(name="Arial", size=9, color=BLUE if col_idx == 6 and cell.hyperlink else NAVY)
                    cell.alignment = Alignment(vertical="center", horizontal="right" if col_idx in (3, 4) else "left",
                                               indent=1 if col_idx in (1, 2, 5, 6) else 0,
                                               shrink_to_fit=col_idx == 5,
                                               wrap_text=col_idx == 6)
                    cell.border = Border(bottom=Side(style="thin", color="E0E4E9"))
                ws.row_dimensions[theme_row].height = 38
                theme_row += 1
        for offset, category in enumerate(ranked_categories):
            if category in first_theme:
                _internal_link(ws.cell(6 + offset, 2), dim, f"A{first_theme[category]}")
        if theme_row > start + 2:
            ws.auto_filter.ref = f"A{start + 1}:F{theme_row - 1}"

def _detail_sheets(workbook, metrics):
    for dim in metrics:
        name = f"{dim}_AI归类"
        if name not in workbook:
            continue
        ws = workbook[name]
        ws.sheet_state = "hidden"
    for ws in workbook:
        if ws.title.endswith("_原始统计") or ws.title.endswith("_标签标准化"):
            ws.sheet_state = "hidden"


def enhance_report(path, product_name=None, field_roles=None):
    """Add a dashboard and readable views to a Step 2 workbook in place."""
    path = Path(path)
    workbook = load_workbook(path)
    if "VOC看板" in workbook:
        return path  # avoid silently replacing an already reviewed presentation
    row_count, ratings, metrics = collect_metrics(workbook)
    if not metrics:
        raise ValueError("没有可用于生成报告的 AI 归类结果")
    category_rows = _summary(workbook, row_count, metrics)
    negative_dim = _role(metrics, "负向", field_roles)
    if negative_dim:
        _cross(workbook, ratings, metrics, field_roles)
        first_evidence = _evidence(workbook, metrics, field_roles)
        for category, evidence_row in first_evidence.items():
            summary_row = category_rows.get(negative_dim, {}).get(category)
            if summary_row:
                cell = workbook["摘要汇总"].cell(summary_row, 2)
                _internal_link(cell, "问题证据", f"A{evidence_row}")
                cell.font = Font(name="Arial", size=10, color=BLUE, underline="single")
        for row in workbook["星级问题"].iter_rows(min_row=6, max_col=1):
            cell = row[0]
            if cell.value in first_evidence:
                _internal_link(cell, "问题证据", f"A{first_evidence[cell.value]}")
                cell.font = Font(name="Arial", size=10, color=BLUE, underline="single")
    _dashboard(workbook, row_count, ratings, metrics, category_rows, product_name, field_roles)
    _dimension_sheets(workbook, row_count, metrics, field_roles)
    _source_sheet(workbook, metrics)
    _detail_sheets(workbook, metrics)
    workbook["摘要汇总"].sheet_state = "hidden"
    workbook["二级主题汇总"].sheet_state = "hidden"
    front = ["VOC看板", *metrics, "星级问题", "问题证据", "合并数据"]
    for index, name in enumerate(front):
        if name in workbook:
            workbook.move_sheet(workbook[name], offset=index - workbook.index(workbook[name]))
    workbook.active = 0
    workbook.save(path)
    return path


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="整理评论分析 Excel 报告")
    parser.add_argument("workbook", help="Step 2 输出的 .xlsx 文件")
    parser.add_argument("--product-name", default=None, help="看板显示的产品或品类名称")
    args = parser.parse_args()
    enhance_report(args.workbook, product_name=args.product_name)
