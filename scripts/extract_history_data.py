# -*- coding: utf-8 -*-
"""按可选 Excel ASIN 名单提取拆分文件，生成历史销量和销额。"""

import argparse
import os
import re
import sys
import tempfile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import IllegalCharacterError


MONTH_RE = re.compile(r"\d{4}-\d{2}$")
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def is_month(value):
    return bool(MONTH_RE.fullmatch(value)) and 1 <= int(value[-2:]) <= 12


def normalize_month(value):
    if isinstance(value, (pd.Timestamp,)):
        return value.strftime("%Y-%m")
    label = str(value).strip()
    match = re.search(r"(\d{4})[-./年](\d{1,2})", label)
    if match and 1 <= int(match.group(2)) <= 12:
        return f"{match.group(1)}-{int(match.group(2)):02d}"
    match = re.fullmatch(r"(\d{4})(\d{2})", label)
    if match and 1 <= int(match.group(2)) <= 12:
        return f"{match.group(1)}-{match.group(2)}"
    return label


def normalize_asins(frame):
    frame = frame.copy()
    frame["ASIN"] = frame["ASIN"].astype("string").str.strip().str.upper()
    return frame[frame["ASIN"].notna() & frame["ASIN"].ne("")]


def read_source_folder(folder, volume_field, revenue_field):
    if not os.path.isdir(folder):
        raise ValueError(f"源数据文件夹不存在：{folder}")
    files = []
    for name in os.listdir(folder):
        if not name.lower().endswith(".xlsx") or name.startswith("~$") or "_销量拆分" not in name:
            continue
        if any(tag in name.lower() for tag in ('.backup', '_backup', '_备份')):
            continue
        month = normalize_month(name)
        if not is_month(month):
            match = re.search(r"(\d{4}-\d{2}|\d{4}\.\d{2}|\d{6})(?=_销量拆分)", name)
            month = normalize_month(match.group(1)) if match else ""
        if is_month(month):
            files.append((os.path.join(folder, name), month))
    if not files:
        raise ValueError("源数据中未找到带年月的 *_销量拆分.xlsx 文件")
    months = [month for _, month in files]
    if len(months) != len(set(months)):
        duplicates = {}
        for path, month in files:
            duplicates.setdefault(month, []).append(os.path.basename(path))
        detail = '; '.join(f"{month}：{', '.join(names)}" for month, names in duplicates.items() if len(names) > 1)
        raise ValueError(f"源数据中同一个月份有多个拆分文件，请先确认应使用哪一个。{detail}")

    volumes, revenues = {}, {}
    for path, month in sorted(files, key=lambda item: item[1]):
        frame = pd.read_excel(path, dtype={"ASIN": str})
        required = {"ASIN", volume_field, revenue_field}
        missing = required.difference(frame.columns)
        if missing:
            raise ValueError(f"{os.path.basename(path)} 缺少列：{', '.join(sorted(missing))}")
        frame = normalize_asins(frame)
        for field, target in ((volume_field, volumes), (revenue_field, revenues)):
            frame[field] = pd.to_numeric(frame[field], errors="coerce")
            target[month] = frame.groupby("ASIN", sort=False)[field].sum(min_count=1)
    return pd.DataFrame(volumes), pd.DataFrame(revenues)


def read_asin_list(path, sheet_name=''):
    """读取首行 ASIN 表头，保持名单顺序，忽略空值并去重。"""
    if not os.path.isfile(path):
        raise ValueError(f"ASIN 名单文件不存在：{path}")
    if not path.lower().endswith(('.xlsx', '.xlsm')):
        raise ValueError('ASIN 名单文件须为 .xlsx 或 .xlsm 格式')
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"ASIN 名单工作表不存在：{sheet_name}")
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        columns = [i for i, value in enumerate(headers) if str(value).strip().upper() == 'ASIN']
        if len(columns) != 1:
            raise ValueError(f"名单工作表「{sheet.title}」首行须包含唯一的 ASIN 列")
        asins = []
        seen = set()
        for number, row in enumerate(rows, 2):
            value = row[columns[0]] if len(row) > columns[0] else None
            if value is None:
                continue
            asin = str(value).strip().upper()
            if asin and asin not in seen:
                seen.add(asin)
                asins.append(asin)
        if not asins:
            raise ValueError('ASIN 名单没有可读取的 ASIN，请检查工作表')
        return asins, sheet.title
    finally:
        workbook.close()


def validate_sheet_names(volume_name, revenue_name):
    for name in (volume_name, revenue_name):
        if not name or len(name) > 31 or INVALID_SHEET_CHARS.search(name) or name.startswith("'") or name.endswith("'"):
            raise ValueError(f"无效的 Sheet 名称：{name!r}（不能超过 31 字，或含 Excel 禁用字符）")
    if volume_name.casefold() == revenue_name.casefold():
        raise ValueError("销量和销额的生成 Sheet 名称必须不同")


def write_workbook(output, volume, revenue, volume_name, revenue_name):
    validate_sheet_names(volume_name, revenue_name)
    if not output.lower().endswith(".xlsx"):
        raise ValueError("输出文件必须是 .xlsx 格式")
    if os.path.exists(output):
        workbook = load_workbook(output)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    for name, frame in ((volume_name, volume), (revenue_name, revenue)):
        existing = next((item for item in workbook.sheetnames
                         if item.casefold() == name.casefold()), None)
        if existing is not None:
            position = workbook.sheetnames.index(existing)
            del workbook[existing]
            sheet = workbook.create_sheet(name, position)
        else:
            sheet = workbook.create_sheet(name)
        sheet.append(["ASIN", *frame.columns])
        for asin, row in frame.iterrows():
            sheet.append([asin, *[None if pd.isna(value) else value.item() if hasattr(value, "item") else value for value in row]])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 17

    directory = os.path.dirname(os.path.abspath(output))
    os.makedirs(directory, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".pm-stack-data-", suffix=".xlsx", dir=directory)
    os.close(fd)
    try:
        workbook.save(temporary)
        if os.path.exists(output):
            import shutil
            from datetime import datetime
            backup = output + '.' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '.backup'
            shutil.copy2(output, backup)
            print(f"已备份原输出文件：{backup}")
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)


def extract(folder, volume_field, revenue_field, output,
            output_volume_sheet='历史销量', output_revenue_sheet='历史销额',
            asin_file='', asin_sheet=''):
    if not folder:
        raise ValueError("请指定源数据文件夹")
    validate_sheet_names(output_volume_sheet, output_revenue_sheet)
    requested = None
    if asin_file:
        requested, actual_sheet = read_asin_list(asin_file, asin_sheet)
        if os.path.normcase(os.path.abspath(asin_file)) == os.path.normcase(os.path.abspath(output)):
            if actual_sheet.casefold() in (output_volume_sheet.casefold(), output_revenue_sheet.casefold()):
                raise ValueError('生成工作表不能覆盖同一文件中的 ASIN 名单工作表')
        print(f"ASIN 名单：{actual_sheet}，{len(requested)} 个唯一 ASIN（全部行）")
    elif asin_sheet:
        raise ValueError('请先指定 ASIN 名单文件')
    volume, revenue = read_source_folder(folder, volume_field, revenue_field)
    available = set(volume.index) | set(revenue.index)
    if not available:
        raise ValueError('源数据中没有可提取的 ASIN')
    asins = requested if requested is not None else sorted(available)
    if requested is not None:
        missing = [asin for asin in requested if asin not in available]
        found = len(requested) - len(missing)
        print(f"匹配结果：找到 {found} 个，未找到 {len(missing)} 个 ASIN")
        if not found:
            raise ValueError('名单中的 ASIN 均未在源数据中找到，未修改输出文件')
        if missing:
            print('[提示] 未找到的 ASIN（保留空白行）：' + ', '.join(missing))
    months = sorted(set(volume.columns) | set(revenue.columns))
    volume = volume.reindex(index=asins, columns=months)
    revenue = revenue.reindex(index=asins, columns=months)
    volume.index.name = revenue.index.name = 'ASIN'
    write_workbook(output, volume, revenue, output_volume_sheet, output_revenue_sheet)
    print(f"已写入：{output}")
    print(f"{output_volume_sheet}、{output_revenue_sheet}：{len(asins)} 个 ASIN，{len(months)} 个月")


def main():
    parser = argparse.ArgumentParser(description="按 Excel ASIN 名单提取历史销量和销额")
    parser.add_argument("--folder", required=True)
    parser.add_argument("--asin-file", default="")
    parser.add_argument("--asin-sheet", default="")
    parser.add_argument("--volume-field", default="子体销量_矫正")
    parser.add_argument("--revenue-field", default="子体销售额_矫正")
    parser.add_argument("--output", required=True)
    parser.add_argument("--output-volume-sheet", default="历史销量")
    parser.add_argument("--output-revenue-sheet", default="历史销额")
    args = parser.parse_args()
    try:
        extract(args.folder, args.volume_field, args.revenue_field, args.output,
                args.output_volume_sheet, args.output_revenue_sheet,
                args.asin_file, args.asin_sheet)
    except (OSError, ValueError, KeyError, IllegalCharacterError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
