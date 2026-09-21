"""按「商品详情页链接」批量抓取五点描述与主副图，回填到竞品表 Excel。

列位置规则：五点描述放在「商品标题」后面，image_1..N 放在「商品主图」后面；
若表中已有这些列则原地填充。用 openpyxl 读写以保留原表格式与缩略图。
断点文件为「输出文件.checkpoint.json」，重跑自动跳过已抓取的 ASIN。

用法：python scripts/fill_product_details.py <输入.xlsx> [输出.xlsx] [--limit N] [--delay 秒] [--headed]
"""

from __future__ import annotations

import json
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import argparse
from copy import copy

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from core.reviews.product_fetcher import (
    ProductFetchError,
    ProductFetcherSession,
    resolve_asin_and_url,
)

LINK_CANDIDATES = ["商品详情页链接", "详情页链接", "商品链接", "链接"]
TITLE_CANDIDATES = ["商品标题", "标题"]
MAIN_IMG_CANDIDATES = ["商品主图", "主图"]
BULLET_COL = "五点描述"
STATUS_COL = "抓取状态"
MAX_IMAGE_COLS = 9
IMAGE_COL_PATTERN = re.compile(r"image_(\d+)")


def log(message: str) -> None:
    print(message, flush=True)


def scan_headers(ws):
    return {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value is not None and str(cell.value).strip()
    }


def pick_header(headers, candidates):
    for name in candidates:
        if name in headers:
            return name, headers[name]
    return None, None


def find_sheet(wb):
    for ws in wb.worksheets:
        headers = scan_headers(ws)
        if pick_header(headers, LINK_CANDIDATES)[0]:
            return ws
    raise SystemExit("未找到包含「商品详情页链接」列的工作表")


def copy_header_style(src, dst):
    """复制表头单元格样式，使插入列的表头与原表头行外观一致。"""
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)
    dst.number_format = src.number_format


def shift_column_dims(ws, pos, amount):
    """openpyxl 的 insert_cols 不移动列宽，插入后手动把 pos 及之后的列宽右移。"""
    snapshot = {
        column_index_from_string(letter): dim
        for letter, dim in list(ws.column_dimensions.items())
    }
    ws.column_dimensions.clear()
    for idx, dim in snapshot.items():
        new_idx = idx + amount if idx >= pos else idx
        letter = get_column_letter(new_idx)
        dim.index = letter
        ws.column_dimensions[letter] = dim


def load_checkpoint(path):
    if os.path.exists(path):
        try:
            with open(path, encoding="utf-8") as fh:
                return json.load(fh)
        except Exception as exc:
            log(f"[警告] 断点文件读取失败，忽略：{exc}")
    return {}


def save_checkpoint(path, results):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(results, fh, ensure_ascii=False)


def main() -> None:
    parser = argparse.ArgumentParser(description="批量抓取五点描述与主副图并回填 Excel")
    parser.add_argument("input", nargs="?", default=os.environ.get("INPUT_FILE"))
    parser.add_argument("output", nargs="?", default=os.environ.get("OUTPUT_FILE"))
    parser.add_argument("--limit", type=int, default=0, help="本次最多新抓取多少个 ASIN（测试用）")
    parser.add_argument("--delay", type=float, default=2.0, help="两次抓取之间的间隔秒数")
    parser.add_argument("--headed", action="store_true", help="显示浏览器窗口（默认无头）")
    args = parser.parse_args()

    if not args.input:
        raise SystemExit("请提供输入 Excel 路径")
    input_path = os.path.abspath(args.input)
    if args.output:
        output_path = os.path.abspath(args.output)
    else:
        stem, _ = os.path.splitext(input_path)
        output_path = f"{stem}_已补五点与主副图.xlsx"
    checkpoint_path = output_path + ".checkpoint.json"

    log(f"[读取] {input_path}")
    wb = load_workbook(input_path)
    ws = find_sheet(wb)
    headers = scan_headers(ws)
    link_name, link_col = pick_header(headers, LINK_CANDIDATES)
    title_name, _ = pick_header(headers, TITLE_CANDIDATES)
    main_img_name, _ = pick_header(headers, MAIN_IMG_CANDIDATES)
    log(f"[读取] 工作表「{ws.title}」，链接列「{link_name}」，共 {ws.max_row - 1} 行数据")

    row_links = {}
    asin_order = []
    asin_of_row = {}
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, link_col).value
        text = str(value).strip() if value is not None else ""
        row_links[row] = text
        if not text:
            continue
        try:
            asin, _ = resolve_asin_and_url(text)
        except ProductFetchError:
            continue
        asin_of_row[row] = asin
        if asin not in asin_order:
            asin_order.append(asin)
    log(f"[读取] 去重后待抓取 ASIN {len(asin_order)} 个")

    results = load_checkpoint(checkpoint_path)
    pending = [
        asin for asin in asin_order
        if asin not in results or not results[asin].get("ok")
    ]
    if args.limit > 0:
        pending = pending[: args.limit]

    if pending:
        with ProductFetcherSession(headless=not args.headed, log_callback=log) as session:
            for i, asin in enumerate(pending, 1):
                log(f"[进度] {i}/{len(pending)} {asin}")
                try:
                    data = session.fetch(asin)
                    results[asin] = {"ok": True, **data}
                except ProductFetchError as exc:
                    log(f"[重试] {asin}：{exc}")
                    time.sleep(max(args.delay, 3))
                    try:
                        data = session.fetch(asin)
                        results[asin] = {"ok": True, **data}
                    except ProductFetchError as exc2:
                        log(f"[失败] {asin}：{exc2}")
                        results[asin] = {"ok": False, "error": str(exc2)}
                save_checkpoint(checkpoint_path, results)
                if i < len(pending) and args.delay > 0:
                    time.sleep(args.delay)
    else:
        log("[进度] 断点已覆盖全部 ASIN，无需新抓取")

    ok_results = {a: r for a, r in results.items() if r.get("ok")}
    n_images = min(
        (max((len(r["images"]) for r in ok_results.values()), default=0)),
        MAX_IMAGE_COLS,
    )

    headers = scan_headers(ws)
    if BULLET_COL not in headers:
        _, title_col = pick_header(headers, TITLE_CANDIDATES)
        pos = title_col + 1 if title_col else ws.max_column + 1
        ws.insert_cols(pos, 1)
        shift_column_dims(ws, pos, 1)
        ws.cell(1, pos, BULLET_COL)
        copy_header_style(ws.cell(1, pos - 1), ws.cell(1, pos))
        ws.column_dimensions[get_column_letter(pos)].width = 45
        log(f"[列] 在「{title_name or '表尾'}」后插入「{BULLET_COL}」列")
    if n_images:
        headers = scan_headers(ws)
        existing = sorted(
            (name for name in headers if IMAGE_COL_PATTERN.fullmatch(name)),
            key=lambda name: int(IMAGE_COL_PATTERN.fullmatch(name).group(1)),
        )
        if existing:
            base = headers[existing[0]]
            if len(existing) < n_images:
                insert_at = base + len(existing)
                ws.insert_cols(insert_at, n_images - len(existing))
                shift_column_dims(ws, insert_at, n_images - len(existing))
                for j in range(len(existing) + 1, n_images + 1):
                    cell = ws.cell(1, base + j - 1)
                    cell.value = f"image_{j}"
                    copy_header_style(ws.cell(1, base), cell)
                    ws.column_dimensions[get_column_letter(base + j - 1)].width = 22
                log(f"[列] 已有图片列扩展到 image_{n_images}")
        else:
            _, main_img_col = pick_header(headers, MAIN_IMG_CANDIDATES)
            _, bullet_col = pick_header(headers, [BULLET_COL])
            anchor = main_img_col or bullet_col
            pos = anchor + 1 if anchor else ws.max_column + 1
            ws.insert_cols(pos, n_images)
            shift_column_dims(ws, pos, n_images)
            ref = ws.cell(1, pos - 1) if anchor else ws.cell(1, pos + n_images)
            for j in range(1, n_images + 1):
                cell = ws.cell(1, pos + j - 1)
                cell.value = f"image_{j}"
                copy_header_style(ref, cell)
                ws.column_dimensions[get_column_letter(pos + j - 1)].width = 22
            log(f"[列] 在「{main_img_name or BULLET_COL}」后插入 image_1..image_{n_images} 列")

    headers = scan_headers(ws)
    bullet_col = headers[BULLET_COL]
    image_cols = {j: headers[f"image_{j}"] for j in range(1, n_images + 1)}
    status_col = ws.max_column + 1
    ws.cell(1, status_col, STATUS_COL)
    copy_header_style(ws.cell(1, status_col - 1), ws.cell(1, status_col))
    ws.column_dimensions[get_column_letter(status_col)].width = 14
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    n_ok = n_fail = n_nolink = n_skip = 0
    for row in range(2, ws.max_row + 1):
        asin = asin_of_row.get(row)
        if asin is None:
            ws.cell(row, status_col, "无有效链接" if row_links[row] else "无链接")
            n_nolink += 1
            continue
        result = results.get(asin)
        if result is None:
            ws.cell(row, status_col, "未抓取")
            n_skip += 1
            continue
        if not result.get("ok"):
            ws.cell(row, status_col, f"失败: {result.get('error', '未知错误')}")
            n_fail += 1
            continue
        bullets = result.get("bullets", [])
        ws.cell(row, bullet_col, "About this item\n" + "\n".join(bullets))
        for j, col in image_cols.items():
            images = result.get("images", [])
            ws.cell(row, col, images[j - 1] if j <= len(images) else None)
        ws.cell(row, status_col, "成功")
        n_ok += 1

    try:
        wb.save(output_path)
    except PermissionError:
        stem, ext = os.path.splitext(output_path)
        output_path = f"{stem}_{time.strftime('%H%M%S')}{ext}"
        log(f"[警告] 原输出文件被占用（可能正被 WPS/Excel 打开），改存为：{output_path}")
        wb.save(output_path)
    log(f"[完成] 成功 {n_ok} 行，失败 {n_fail} 行，未抓取 {n_skip} 行，无链接 {n_nolink} 行")
    log(f"[完成] 输出文件：{output_path}")


if __name__ == "__main__":
    main()
