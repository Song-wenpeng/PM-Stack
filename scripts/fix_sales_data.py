# -*- coding: utf-8 -*-
"""修正拆分数据 — 按 ASIN + 年月 修改子体销量/销额矫正值

用法:
  # 单个 ASIN 单月修正
  python fix_sales_data.py --folder ./拆分数据/ --asin B00XXX --month 2025-03 \
      --volume 12345 --revenue 67890

  # 只改销量，不改销额
  python fix_sales_data.py --folder ./拆分数据/ --asin B00XXX --month 2025-03 \
      --volume 12345

  # 批量修正（CSV 文件: ASIN, 年月, 销量, 销额）
  python fix_sales_data.py --folder ./拆分数据/ --batch fix_list.csv
"""

import os
import sys
import re
import argparse
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


def extract_year_month(filename):
    """从文件名提取年月。"""
    # YYYY-MM
    m = re.search(r'(\d{4})-(\d{2})', filename)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    # YYYY.MM
    m = re.search(r'(\d{4})\.(\d{2})', filename)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    # YYYYMM
    m = re.search(r'(\d{4})(\d{2})(?:\D|$)', filename)
    if m:
        year, month = m.group(1), m.group(2)
        if 1 <= int(month) <= 12:
            return f"{year}-{month}"
    return None


def scan_files(folder):
    """扫描文件夹，返回 {年月: 文件路径}。"""
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"文件夹不存在: {folder}")

    file_map = {}
    for fname in os.listdir(folder):
        if fname.startswith('~$') or not fname.endswith('.xlsx'):
            continue
        if '_销量拆分' not in fname:
            continue
        ym = extract_year_month(fname)
        if ym:
            file_map[ym] = os.path.join(folder, fname)
    return file_map


def fix_single(folder, asin, month, volume=None, revenue=None):
    """修正单个 ASIN 单月数据。"""
    file_map = scan_files(folder)

    if month not in file_map:
        available = sorted(file_map.keys())
        print(f"错误: 未找到 {month} 的拆分文件")
        print(f"  可用年月: {available[0] if available else '无'} ~ {available[-1] if available else '无'}")
        return False

    filepath = file_map[month]
    print(f"  文件: {os.path.basename(filepath)}")

    # 使用 openpyxl 原位修改单元格，保留其他 Sheet、公式与样式。
    workbook = load_workbook(filepath)
    worksheet = workbook.active
    headers = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    if 'ASIN' not in headers:
        print(f"错误: 文件中没有 ASIN 列。列名: {list(headers)}")
        workbook.close()
        return False

    asin_column = headers['ASIN']
    row_number = None
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=asin_column).value
        if value is not None and str(value).strip() == asin.strip():
            row_number = row
            break

    if row_number is None:
        print(f"错误: 文件中未找到 ASIN={asin}")
        examples = [
            worksheet.cell(row=row, column=asin_column).value
            for row in range(2, min(worksheet.max_row + 1, 7))
        ]
        print(f"  文件中前5个 ASIN: {examples}")
        workbook.close()
        return False

    volume_column = headers.get('子体销量_矫正')
    revenue_column = headers.get('子体销售额_矫正')
    old_volume = (
        worksheet.cell(row=row_number, column=volume_column).value
        if volume_column else 'N/A'
    )
    old_revenue = (
        worksheet.cell(row=row_number, column=revenue_column).value
        if revenue_column else 'N/A'
    )

    print(f"  ASIN: {asin}")
    print(f"  旧值: 销量={old_volume}, 销额={old_revenue}")

    if volume is not None:
        if not volume_column:
            print("  警告: 文件中没有「子体销量_矫正」列，跳过销量修改")
        else:
            worksheet.cell(row=row_number, column=volume_column).value = float(volume)
            print(f"  新值: 销量={volume}")

    if revenue is not None:
        if not revenue_column:
            print("  警告: 文件中没有「子体销售额_矫正」列，跳过销额修改")
        else:
            worksheet.cell(row=row_number, column=revenue_column).value = float(revenue)
            print(f"  新值: 销额={revenue}")

    # 同目录临时文件 + 原子替换；替换前保留带时间戳的原始备份。
    source_path = Path(filepath)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = source_path.with_name(
        f"{source_path.stem}.backup_{timestamp}{source_path.suffix}"
    )
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{source_path.stem}.",
            suffix=".tmp.xlsx",
            dir=source_path.parent,
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)
        workbook.save(temp_path)
        shutil.copy2(source_path, backup_path)
        os.replace(temp_path, source_path)
    finally:
        workbook.close()
        if temp_path and temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass

    print("  [OK] 已保存")
    print(f"  备份: {backup_path}")
    return True


def fix_batch(folder, batch_file):
    """批量修正，batch_file 为 CSV（含 ASIN, 年月, 销量, 销额 列）。"""
    if not os.path.exists(batch_file):
        print(f"错误: 批量文件不存在: {batch_file}")
        return 0, 0

    df_batch = pd.read_csv(batch_file, dtype=str)
    required = ['ASIN', '年月']
    for r in required:
        if r not in df_batch.columns:
            print(f"错误: 批量文件缺少「{r}」列。列名: {df_batch.columns.tolist()}")
            return 0, 0

    success, fail = 0, 0
    for _, row in df_batch.iterrows():
        asin = row['ASIN'].strip()
        month = row['年月'].strip()
        volume = float(row['销量']) if '销量' in df_batch.columns and pd.notna(row.get('销量')) else None
        revenue = float(row['销额']) if '销额' in df_batch.columns and pd.notna(row.get('销额')) else None

        print(f"\n--- {asin} @ {month} ---")
        try:
            if fix_single(folder, asin, month, volume, revenue):
                success += 1
            else:
                fail += 1
        except Exception as e:
            print(f"  错误: {e}")
            fail += 1

    print(f"\n完成: 成功 {success}, 失败 {fail}")
    return success, fail


def main():
    parser = argparse.ArgumentParser(description='修正拆分数据中的销量/销额矫正值')
    parser.add_argument('--folder', required=True, help='拆分文件所在文件夹')
    parser.add_argument('--asin', default='', help='ASIN（单个修正时使用）')
    parser.add_argument('--month', default='', help='年月，格式 YYYY-MM（单个修正时使用）')
    parser.add_argument('--volume', type=float, default=None, help='新的子体销量_矫正值')
    parser.add_argument('--revenue', type=float, default=None, help='新的子体销售额_矫正值')
    parser.add_argument('--batch', default='', help='批量修正 CSV 文件（列: ASIN, 年月, 销量, 销额）')
    args = parser.parse_args()

    print("=" * 55)
    print("  拆分数据修正")
    print(f"  文件夹: {args.folder}")
    print("=" * 55)

    if args.batch:
        fix_batch(args.folder, args.batch)
    elif args.asin and args.month:
        if args.volume is None and args.revenue is None:
            print("错误: 请至少指定 --volume 或 --revenue")
            sys.exit(1)
        ok = fix_single(args.folder, args.asin, args.month,
                        args.volume, args.revenue)
        if not ok:
            sys.exit(1)
    else:
        print("错误: 请指定 --asin --month（单个修正）或 --batch（批量修正）")
        print("示例:")
        print("  python fix_sales_data.py --folder ./拆分/ --asin B00XXX --month 2025-03 --volume 12345")
        print("  python fix_sales_data.py --folder ./拆分/ --batch fix_list.csv")
        sys.exit(1)


if __name__ == '__main__':
    main()
