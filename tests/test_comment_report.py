"""Checks for review-level denominators and the reader-facing report export."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.comment_report import collect_metrics, enhance_report, repair_evidence_report


class CommentReportTests(unittest.TestCase):
    def _source(self, path):
        workbook = Workbook()
        raw = workbook.active
        raw.title = "合并数据"
        raw.append(["星级", "标题", "内容", "用户负向关注点"])
        raw.append(["1 star", "Bad", "Failed immediately", '["短期失效", "质量差"]'])
        raw.append(["2 star", "Failed", "Stopped after a week", '["短期失效"]'])
        raw.append(["5 star", "Good", "Works", "[]"])

        stat = workbook.create_sheet("用户负向关注点_原始统计")
        stat.append(["原始标签", "出现次数"])
        stat.append(["短期失效", 2])
        stat.append(["质量差", 1])

        normalization = workbook.create_sheet("用户负向关注点_标签标准化")
        normalization.append(["原始标签", "标准标签"])
        normalization.append(["短期失效", "短期失效"])
        normalization.append(["质量差", "质量差"])

        category = workbook.create_sheet("用户负向关注点_AI归类")
        category.append(["归类模式", "一级类别", "一级类别总次数", "二级主题",
                         "二级主题总次数", "标准标签", "标准标签次数", "代表性标签"])
        category.append(["AI", "可靠性问题", 3, "早期失效", 3, "短期失效", 2, "短期失效；质量差"])
        category.append(["AI", "可靠性问题", 3, "早期失效", 3, "质量差", 1, "短期失效；质量差"])
        workbook.create_sheet("摘要汇总")
        workbook.create_sheet("二级主题汇总")
        workbook.save(path)

    def test_counts_reviews_once_per_category(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            self._source(path)
            workbook = load_workbook(path)
            count, ratings, metrics = collect_metrics(workbook)
            self.assertEqual(count, 3)
            self.assertEqual(ratings[1], 1)
            self.assertEqual(metrics["用户负向关注点"]["total_mentions"], 3)
            self.assertEqual(len(metrics["用户负向关注点"]["review_rows"]["可靠性问题"]), 2)

    def test_evidence_uses_headers_and_review_specific_secondary_categories(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reordered.xlsx"
            self._source(path)
            book = load_workbook(path)
            raw = book["合并数据"]
            raw.insert_cols(1, 3)
            for col, name in enumerate(("型号", "VP评论", "Vine评论"), 1):
                raw.cell(1, col, name)
            raw.cell(2, 1, 5)  # Must not be mistaken for this review's actual 1 star.
            raw.cell(2, 3, "不是评论内容")
            raw.cell(2, 7, '["短期失效", "质量差", "质量差", "连接不稳"]')
            taxonomy = book["用户负向关注点_AI归类"]
            taxonomy.cell(3, 4, "制造缺陷")
            taxonomy.append(["AI", "连接问题", 1, "接触不良", 1, "连接不稳", 1, "连接不稳"])
            normalization = book["用户负向关注点_标签标准化"]
            normalization.cell(2, 2, "提前损坏")
            taxonomy.cell(2, 6, "提前损坏")
            book.save(path)
            enhance_report(path)
            result = load_workbook(path)
            evidence = result["问题证据"]
            self.assertEqual([c.value for c in evidence[5]],
                             ["负向一级类别", "负向二级类别", "提取到的负向关注点", "星级", "原始行号", "评论内容", "原文链接"])
            self.assertEqual([c.value for c in evidence[6]][:6],
                             ["可靠性问题", "制造缺陷；早期失效", "短期失效；质量差；质量差；连接不稳", 1, 2, "Failed immediately"])
            self.assertEqual(evidence["B7"].value, "早期失效")
            self.assertEqual(evidence["D7"].value, 2)
            self.assertEqual(evidence["F7"].value, "Stopped after a week")
            self.assertEqual(evidence["A8"].value, "连接问题")
            self.assertEqual(evidence["B8"].value, "接触不良")
            self.assertEqual(evidence["C8"].value, "短期失效；质量差；质量差；连接不稳")
            self.assertEqual(evidence["C7"].value, "短期失效")
            self.assertEqual(evidence["G6"].hyperlink.location, "'合并数据'!F2")
            self.assertEqual(evidence["F6"].hyperlink.location, "'合并数据'!F2")
            self.assertEqual(evidence.tables["IssueEvidence"].ref, "A5:G8")
            result.close()

    def test_existing_report_repair_preserves_other_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            output = Path(directory) / "repaired.xlsx"
            self._source(path)
            enhance_report(path)
            original = load_workbook(path)
            original["问题证据"]["D6"] = None
            original["问题证据"]["F6"] = None
            original.save(path)
            before = {ws.title: list(ws.values) for ws in original if ws.title != "问题证据"}
            names = original.sheetnames
            original.close()
            repair_evidence_report(path, output)
            repaired = load_workbook(output)
            self.assertEqual(repaired.sheetnames, names)
            self.assertEqual({ws.title: list(ws.values) for ws in repaired if ws.title != "问题证据"}, before)
            self.assertEqual(repaired["问题证据"]["D6"].value, 1)
            self.assertEqual(repaired["问题证据"]["F6"].value, "Failed immediately")
            self.assertEqual(repaired["问题证据"]["B6"].value, "早期失效")
            repaired.close()

    def test_export_keeps_raw_tags_and_links_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            self._source(path)
            enhance_report(path, product_name="测试品类")
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames[:4],
                             ["VOC看板", "用户负向关注点", "星级问题", "问题证据"])
            self.assertEqual(workbook["摘要汇总"]["C2"].value, 3)
            self.assertEqual(workbook["摘要汇总"]["E2"].value, 2)
            self.assertAlmostEqual(workbook["摘要汇总"]["F2"].value, 2 / 3)
            self.assertEqual(workbook["星级问题"]["G6"].value, 2)
            self.assertTrue(workbook["星级问题"]["A6"].hyperlink)
            self.assertEqual(workbook["合并数据"]["D2"].value, "短期失效；质量差")
            self.assertEqual(workbook["标签原值"]["D2"].value, '["短期失效", "质量差"]')
            self.assertEqual(workbook["用户负向关注点_原始统计"].sheet_state, "hidden")
            self.assertEqual(len(workbook["VOC看板"]._charts), 0)
            self.assertEqual(len(workbook["星级问题"]._charts), 0)
            self.assertEqual(workbook["摘要汇总"].sheet_state, "hidden")
            self.assertEqual(workbook["用户负向关注点_AI归类"].sheet_state, "hidden")
            self.assertEqual(workbook["用户负向关注点"]["B6"].value, "可靠性问题")
            self.assertEqual(workbook["用户负向关注点"]["B12"].value, "早期失效")
            self.assertTrue(workbook["用户负向关注点"]["F12"].hyperlink)
            self.assertTrue(workbook["用户负向关注点"]["F12"].value.startswith("第2行 ·"))
            self.assertEqual(workbook["合并数据"].sheet_view.topLeftCell, "A1")
            self.assertIsNone(workbook["合并数据"].freeze_panes)
            self.assertIn("测试品类（3 条评论）", workbook["VOC看板"]["A1"].value)
            self.assertEqual(workbook["VOC看板"]["B5"].value, "可靠性问题")
            self.assertEqual(workbook["VOC看板"]["C5"].value, 3)
            self.assertEqual(workbook["VOC看板"]["D5"].value, 1)
            self.assertEqual(workbook["VOC看板"]["E5"].value, "█" * 16)

    def test_other_category_uses_configured_fields_without_eight_field_assumption(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "other_category.xlsx"
            workbook = Workbook()
            raw = workbook.active
            raw.title = "合并数据"
            raw.append(["星级", "标题", "内容", "安装位置", "外观风格"])
            raw.append(["5 star", "Works", "Fits on a desk", '["桌面"]', '["简约"]'])
            raw.append(["4 star", "Good", "Looks neat", '["桌面"]', '["简约"]'])
            for field, label, category in (
                ("安装位置", "桌面", "桌面安装"),
                ("外观风格", "简约", "简约设计"),
            ):
                normalization = workbook.create_sheet(f"{field}_标签标准化")
                normalization.append(["原始标签", "标准标签"])
                normalization.append([label, label])
                classified = workbook.create_sheet(f"{field}_AI归类")
                classified.append(["归类模式", "一级类别", "一级类别总次数", "二级主题",
                                   "二级主题总次数", "标准标签", "标准标签次数", "代表性标签"])
                classified.append(["AI", category, 2, category, 2, label, 2, label])
            workbook.create_sheet("摘要汇总")
            workbook.create_sheet("二级主题汇总")
            workbook.save(path)
            enhance_report(path, product_name="另一类目")
            result = load_workbook(path)
            self.assertEqual(result.sheetnames[:3], ["VOC看板", "安装位置", "外观风格"])
            self.assertNotIn("星级问题", result)
            self.assertNotIn("问题证据", result)
            self.assertEqual(result["安装位置"]["C6"].value, 2)
            self.assertEqual(result["外观风格"]["B6"].value, "简约设计")
            self.assertTrue(result["外观风格"]["F12"].hyperlink)
            self.assertIn("安装位置", [result["VOC看板"].cell(r, 1).value
                                  for r in range(1, result["VOC看板"].max_row + 1)])
            configured_path = Path(directory) / "configured_role.xlsx"
            workbook.save(configured_path)
            enhance_report(configured_path, field_roles={"安装位置": "negative"})
            configured = load_workbook(configured_path)
            self.assertIn("星级问题", configured)
            self.assertIn("问题证据", configured)
            self.assertEqual(configured["安装位置"].sheet_properties.tabColor.rgb[-6:], "DC5757")

    def test_module_preview_reads_report(self):
        from PyQt6.QtWidgets import QApplication, QLabel, QPushButton, QTableWidget, QTabWidget, QWidget
        from modules.comment_analysis import ModuleWidget

        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            self._source(path)
            enhance_report(path)

            class Harness:
                _load_report = ModuleWidget._load_report

                def __init__(self):
                    self.result_context = QLabel()
                    self.result_table = QTableWidget(0, 5)
                    self.open_report_btn = QPushButton()
                    self.tabs = QTabWidget()
                    self.result_tab = QWidget()
                    self.tabs.addTab(self.result_tab, "结果")
                    self._report_path = None

            harness = Harness()
            harness._load_report(str(path))
            self.assertEqual(harness.result_table.rowCount(), 1)
            self.assertEqual(harness.result_table.item(0, 2).text(), "3")
            self.assertTrue(harness.open_report_btn.isEnabled())
            self.assertEqual(harness.tabs.currentWidget(), harness.result_tab)
            self.assertIn("可靠性问题", harness.result_context.text())
            app.processEvents()


if __name__ == "__main__":
    unittest.main()
