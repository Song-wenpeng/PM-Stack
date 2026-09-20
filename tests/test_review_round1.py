from tests import support

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,patch
from openpyxl import load_workbook
from core.reviews.store import ReviewStore
from core.reviews.repository import ReviewRepository
from core.reviews.cloud import SupabaseStore,SupabaseAuthError
from core.reviews.metadata import KEY,content_hash
from core.reviews.detail_export import export_review_details
from core.reviews.sync_lock import uploader_lock
from core.secure_storage import unprotect_secret

class RoundOneTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.store=ReviewStore(str(Path(self.temp.name)/"r.db"))
    def tearDown(self):
        self.store.close();self.temp.cleanup()
    def queue(self,n):
        with self.store.conn:
            for i in range(n):
                p={"product":{"asin":"B0DDG2KQQ9","marketplace":"amazon.com"},"reviews":[{"review_id":str(i),"marketplace":"amazon.com","content":"test"}]}
                self.store.conn.execute("INSERT INTO outbox(idempotency_key,entity_type,payload_json,status,created_at,updated_at) VALUES(?,'review_batch',?,'pending','2026-01-01','2026-01-01')",(str(i),json.dumps(p)))
    def repo(self,send=None):
        r=ReviewRepository(self.store)
        r.cloud=SimpleNamespace(ingest_batch=send or Mock())
        return r
    def test_407_tasks_continue_past_200(self):
        self.queue(407);r=self.repo();progress=[]
        result=r.sync_pending(limit=200,progress=progress.append)
        self.assertEqual((result["synced"],result["remaining"]),(407,0))
        self.assertEqual(r.cloud.ingest_batch.call_count,407)
        self.assertTrue(progress)
    def test_stop_and_resume_without_resending(self):
        self.queue(407);sent=[]
        r=self.repo(lambda p,k:sent.append(k))
        first=r.sync_pending(stop=lambda:len(sent)>=205)
        self.assertTrue(first["cancelled"]);self.assertEqual(first["remaining"],202)
        last=r.sync_pending()
        self.assertEqual(last["synced"],202)
        self.assertEqual(len(sent),len(set(sent)))
    def test_network_failure_stops_and_retries_only_remaining(self):
        self.queue(5)
        r=self.repo(Mock(side_effect=[None,RuntimeError("offline")]))
        result=r.sync_pending()
        self.assertEqual((result["synced"],result["failed"],result["remaining"]),(1,1,4))
        r.cloud.ingest_batch=Mock()
        self.assertEqual(r.sync_pending(force=True)["synced"],4)
    def test_new_tasks_wait_for_next_run(self):
        self.queue(1)
        def send(p,k):
            self.store.enqueue_collection("B0DDG2KQQ9","amazon.com","",[{"review_id":"new","content":"new"}])
        r=self.repo(send)
        result=r.sync_pending()
        self.assertEqual((result["synced"],result["remaining"]),(1,1))
    def test_exclusive_upload_lock(self):
        self.queue(1);r=self.repo()
        with uploader_lock(self.store.db_path):
            with self.assertRaises(RuntimeError):r.sync_pending()
        self.assertEqual(self.store.outbox_stats()["pending"],1)
    def row(self):
        a={"status":"success","content_hash":content_hash("hello<br>world"),"created_at":"2026-09-15",
           "category":"power strip","model":"test","template_hash":"v1",
           "labels":{"positive":["sturdy"],"translation":"=HYPERLINK(1)"},"field_labels":{"positive":"优点","translation":"翻译"}}
        return {"marketplace":"amazon.com","asin":"B0DDG2KQQ9","review_id":"R1","content":"hello<br>world","rating":5,"title":"=1+1",
                "review_url":"https://www.amazon.com/gp/customer-reviews/R1","raw_payload":{KEY:{"analyses":{"a":a}}}}
    def test_single_sheet_exports_ai_safe_text_and_bounded_layout(self):
        row=self.row();path=Path(self.temp.name)/"details.xlsx"
        export_review_details([row],str(path))
        w=load_workbook(path);s=w.active
        self.assertEqual(w.sheetnames,["评论明细"])
        headers=[c.value for c in s[1]]
        values={h:s.cell(2,i+1) for i,h in enumerate(headers)}
        self.assertEqual(values["AI · 优点"].value,"sturdy")
        self.assertEqual(values["AI · 翻译"].data_type,"s")
        self.assertEqual(values["标题"].data_type,"s")
        self.assertEqual(values["内容"].value,"hello\nworld")
        self.assertEqual(values["评论链接"].hyperlink.target,row["review_url"])
        self.assertEqual(s.row_dimensions[2].height,72)
        self.assertEqual(s.freeze_panes,"A2");w.close()
    def test_empty_analysis_differs_from_unanalysed(self):
        row=self.row()
        row["raw_payload"][KEY]["analyses"]["a"]["labels"]={"positive":[]}
        other={**row,"review_id":"R2","raw_payload":{}}
        path=Path(self.temp.name)/"details.xlsx";export_review_details([row,other],str(path))
        w=load_workbook(path);s=w.active
        column=[c.value for c in s[1]].index("分析状态")+1
        self.assertNotEqual(s.cell(2,column).value,s.cell(3,column).value);w.close()
    def test_export_snapshot_includes_local_labels_and_new_rows(self):
        row=self.row();self.store.add_reviews(row["asin"],row["marketplace"],[row])
        remote={**row,"raw_payload":{}}
        r=self.repo();r.cloud.query_reviews=Mock(return_value=[remote])
        rows,source=r.export_snapshot()
        self.assertEqual(len(rows),1)
        self.assertTrue(rows[0]["raw_payload"][KEY]["analyses"])
    def test_export_snapshot_cloud_failure_restarts_local(self):
        row=self.row();self.store.add_reviews(row["asin"],row["marketplace"],[row])
        r=self.repo();r.cloud.query_reviews=Mock(side_effect=RuntimeError("offline"))
        rows,source=r.export_snapshot()
        self.assertEqual(len(rows),1);self.assertIn("本地",source)
    def fake_auth(self):
        c=object.__new__(SupabaseStore);c.project_url="https://example.supabase.co";c.publishable_key="public-test"
        c.workspace_id=None
        session=SimpleNamespace(access_token="test-access",refresh_token="test-refresh")
        c.client=Mock();c.client.auth.sign_in_with_password.return_value=SimpleNamespace(session=session)
        c.ensure_workspace=Mock(return_value="workspace-one")
        return c
    def test_login_atomically_persists_encrypted_session_and_scope(self):
        c=self.fake_auth();c.sign_in("test@example.com","unused",self.store)
        encrypted=self.store.get_setting("supabase_session")
        self.assertNotIn("test-access",encrypted)
        self.assertEqual(json.loads(unprotect_secret(encrypted))["access_token"],"test-access")
        self.assertEqual(json.loads(self.store.get_setting("review_data_scope"))["workspace"],"workspace-one")
    def test_wrong_workspace_does_not_replace_existing_session(self):
        c=self.fake_auth();self.store.set_setting("review_data_scope",json.dumps({"url":c.project_url,"workspace":"original"}))
        self.store.set_setting("supabase_session","original-protected")
        with self.assertRaises(SupabaseAuthError):c.sign_in("test@example.com","unused",self.store)
        self.assertEqual(self.store.get_setting("supabase_session"),"original-protected")
    def test_encryption_failure_does_not_persist_partial_login(self):
        c=self.fake_auth()
        with patch("core.reviews.cloud.protect_secret",side_effect=RuntimeError("storage unavailable")):
            with self.assertRaises(RuntimeError):c.sign_in("test@example.com","unused",self.store)
        self.assertFalse(self.store.get_setting("supabase_session"))
        self.assertFalse(self.store.get_setting("supabase_workspace_id"))
