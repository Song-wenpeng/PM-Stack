"""Exercise batch mapping and real workbook writeback without network or user data."""
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import threading
import time
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from core.field_normalization import (distinct_values,
                                      normalize_after_extraction, quantity,
                                      unsafe_change, validate_mapping)
from core.field_values import is_empty_extracted_field, is_empty_field
from core.field_workbook import FieldWorkbookWriter


CONVENTION = {'kind': 'quantity', 'unit': 'm', 'format': '数值紧接m', 'labels': []}
TEXT = {'kind': 'text', 'unit': '', 'format': '统一表达', 'labels': []}


class FakeModel:
    def __init__(self, convert=None, fail_field=None, after=None):
        self.calls = []
        self.convert = convert or (lambda field, value: value)
        self.fail_field = fail_field
        self.after = after
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

    def create(self, **kwargs):
        prompt = kwargs['messages'][1]['content']
        data = json.loads(prompt.split('数据：\n', 1)[1])
        self.calls.append(data)
        if data['field'] == self.fail_field:
            raise TimeoutError('fake model failure')
        convention = data['convention'] or (CONVENTION if data['field'] == '线长' else TEXT)
        payload = {'convention': convention}
        if '仅制定全字段共享表达约定' not in prompt:
            payload['mappings'] = [dict(source_id=v['source_id'],
                                        target=self.convert(data['field'], v['value']),
                                        status='normalized') for v in data['values']]
        if self.after:
            self.after(len(self.calls))
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(
            content=json.dumps(payload, ensure_ascii=False)), finish_reason='stop')],
            usage=SimpleNamespace(total_tokens=100))


def length_target(field, value):
    parsed = quantity(value)
    return f'{parsed[1].normalize():f}m' if field == '线长' and parsed else value


class NormalizationWorkbookFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'source.xlsx'
        env = patch.dict(os.environ, {'EXTRACT_NORMALIZE': '1', 'EXTRACT_RPM_LIMIT': '0',
            'EXTRACT_TPM_LIMIT': '0', 'EXTRACT_MAX_RETRIES': '1', 'NORMALIZE_BATCH_SIZE': '80',
            'NORMALIZE_CONCURRENCY': '2'})
        env.start()
        self.addCleanup(env.stop)

    def workbook(self, columns, selected=None):
        wb = Workbook()
        ws = wb.active
        ws.title = '商品数据'
        names = list(columns)
        ws.append(['ASIN', *names, '商品标题', '详细参数', 'SKU', '五点描述', '打标', '计算'])
        for i in range(len(next(iter(columns.values())))):
            ws.append(['B000000001', *[columns[k][i] for k in names], 'title', 'spec', 'sku',
                       'bullets', 1 if selected is None or i in selected else 0, f'=1+{i}'])
        wb.create_sheet('其他').append(['保留', "='商品数据'!A2"])
        wb.save(self.path)
        wb.close()
        self.fields = names
        self.config = {'name': 'test', 'target_fields': names,
                       'field_config': {name: '提取' + name for name in names}}

    def run_normalize(self, model, indices=None, output=None, stopped=lambda: False, limiter=None):
        df = pd.read_excel(self.path, dtype=object)
        indices = list(df.index) if indices is None else indices
        writer = FieldWorkbookWriter(self.path, output or self.path, self.fields)
        writer.save(df.loc[indices])
        ok = normalize_after_extraction(df, self.fields, indices, self.config, writer,
            lambda: model, 'test-model', stopped=stopped, limiter=limiter)
        return ok, df

    def frame(self, path=None):
        return pd.read_excel(path or self.path, dtype=object)

    def audit(self, path=None):
        return json.loads(Path(str(path or self.path) + '.normalization.json').read_text(encoding='utf-8'))


class NormalizationTests(NormalizationWorkbookFixture):
    def test_600_asins_50_values_one_model_request_and_600_correct_rows(self):
        variants = []
        for number in range(1, 11):
            variants += [f'{number}Metre', f'{number} Metre', f'{number} m',
                         f'{number * 100}cm', f'{number * 1000}mm']
        values = variants * 12
        self.workbook({'线长': values})
        model = FakeModel(length_target)
        ok, _ = self.run_normalize(model)
        self.assertTrue(ok)
        self.assertEqual(len(model.calls), 1)
        self.assertEqual(len(model.calls[0]['values']), 50)
        self.assertEqual({v['count'] for v in model.calls[0]['values']}, {12})
        self.assertEqual(list(self.frame()['线长']), [length_target('线长', v) for v in values])
        self.assertEqual(len(set(self.frame()['线长'])), 10)
        self.assertEqual(len(self.audit()['fields']['线长']['raw']), 600)

    def test_selected_rows_duplicate_asins_and_other_columns_survive(self):
        self.workbook({'线长': ['1 Metre', '1 Metre', '100cm']}, selected={0, 2})
        from PIL import Image as PILImage
        from openpyxl.drawing.image import Image
        image = Path(self.temp.name) / 'image.png'
        PILImage.new('RGB', (12, 12), 'red').save(image)
        wb = load_workbook(self.path)
        wb.active.add_image(Image(image), 'A2')
        wb.active['B2'].number_format = '@'
        wb.save(self.path)
        wb.close()
        ok, _ = self.run_normalize(FakeModel(length_target), [0, 2])
        self.assertTrue(ok)
        self.assertEqual(list(self.frame()['线长']), ['1m', '1 Metre', '1m'])
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(len(wb.active._images), 1)
        headers = {c.value: c.column for c in wb.active[1]}
        self.assertEqual(wb.active.cell(4, headers['计算']).value, '=1+2')
        self.assertEqual(wb['其他']['B1'].value, "='商品数据'!A2")
        self.assertEqual(wb.active['B2'].number_format, '@')

    def test_resume_reuses_frozen_map_then_manual_value_invalidates_it(self):
        self.workbook({'线长': ['1 Metre', '100cm']})
        self.assertTrue(self.run_normalize(FakeModel(length_target))[0])
        model = FakeModel(lambda *args: self.fail('frozen map must not query AI'))
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(model.calls, [])
        self.assertEqual(self.audit()['fields']['线长']['raw'][0][1], '1 Metre')
        wb = load_workbook(self.path)
        wb.active['B2'] = '9 Metre'
        wb.save(self.path)
        wb.close()
        model = FakeModel(length_target)
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(len(model.calls), 1)
        self.assertEqual(list(self.frame()['线长']), ['9m', '1m'])

    def test_failed_feature_keeps_raw_and_successful_feature_resumes_without_call(self):
        self.workbook({'线长': ['1 Metre'], '颜色': ['Black']})
        self.assertFalse(self.run_normalize(FakeModel(length_target, fail_field='颜色'))[0])
        self.assertEqual(self.frame().loc[0, '线长'], '1m')
        self.assertEqual(self.frame().loc[0, '颜色'], 'Black')
        model = FakeModel(lambda f, v: '黑色' if f == '颜色' else v)
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual([call['field'] for call in model.calls], ['颜色'])
        self.assertEqual(self.frame().loc[0, '颜色'], '黑色')

    def test_batches_freeze_convention_and_resume_partial_mapping(self):
        self.workbook({'线长': [f'{i} Metre' for i in range(1, 6)]})
        stop = threading.Event()
        os.environ['NORMALIZE_BATCH_SIZE'] = '2'
        model = FakeModel(length_target, after=lambda n: stop.set() if n == 3 else None)
        self.assertFalse(self.run_normalize(model, stopped=stop.is_set)[0])
        self.assertEqual(list(self.frame()['线长']), [f'{i} Metre' for i in range(1, 6)])
        cached = self.audit()['fields']['线长']
        self.assertEqual(len(cached['batches']), 1)  # planning + one persisted batch
        model = FakeModel(length_target)
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(len(model.calls), 2)  # no replanning or repeating the first batch
        self.assertTrue(all(call['convention'] == CONVENTION for call in model.calls))
        self.assertEqual(list(self.frame()['线长']), [f'{i}m' for i in range(1, 6)])

    def test_bad_numeric_mapping_is_kept_with_reason(self):
        self.workbook({'线长': ['3ft', '1m/2m可选', '3m/16ft']})
        self.assertTrue(self.run_normalize(FakeModel(lambda *args: '1m'))[0])
        self.assertEqual(list(self.frame()['线长']), ['3ft', '1m/2m可选', '3m/16ft'])
        self.assertTrue(all(v['status'] == 'keep' and v['reason'] for v in
                            self.audit()['fields']['线长']['mapping'].values()))

    def test_output_copy_preserves_source_and_resume_maps(self):
        self.workbook({'线长': ['1 Metre', '100cm']})
        before = self.path.read_bytes()
        output = self.path.parent / 'copy.xlsx'
        self.assertTrue(self.run_normalize(FakeModel(length_target), output=output)[0])
        self.assertEqual(self.path.read_bytes(), before)
        model = FakeModel()
        self.assertTrue(self.run_normalize(model, output=output)[0])
        self.assertEqual(model.calls, [])
        self.assertEqual(list(self.frame(output)['线长']), ['1m', '1m'])

    def test_mapping_does_not_cascade_and_formula_like_text_stays_text(self):
        self.workbook({'颜色': ['A', 'B', 'C']})
        model = FakeModel(lambda f, v: {'A': 'B', 'B': 'C', 'C': '=NOW()'}[v])
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(list(self.frame()['颜色']), ['B', 'C', '=NOW()'])
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['B4'].data_type, 's')

    def test_formula_in_target_field_is_not_submitted_or_changed(self):
        self.workbook({'线长': ['1 Metre', '=1+1']})
        model = FakeModel(length_target)
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(len(model.calls[0]['values']), 1)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['B3'].value, '=1+1')

    def test_disable_leaves_values_and_never_requests_model(self):
        self.workbook({'线长': ['1 Metre']})
        os.environ['EXTRACT_NORMALIZE'] = '0'
        model = FakeModel()
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(model.calls, [])
        self.assertEqual(self.frame().loc[0, '线长'], '1 Metre')
        self.assertFalse(Path(str(self.path) + '.normalization.json').exists())

    def test_waiting_rate_limit_can_be_stopped(self):
        self.workbook({'线长': ['1 Metre']})
        stop = threading.Event()
        class WaitingLimiter:
            def acquire(self, tokens, event):
                stop.set()
                event.wait(2)
                return None
        start = time.monotonic()
        model = FakeModel()
        self.assertFalse(self.run_normalize(model, stopped=stop.is_set, limiter=WaitingLimiter())[0])
        self.assertLess(time.monotonic() - start, 1.5)
        self.assertEqual(model.calls, [])

    def test_two_feature_requests_can_overlap_without_parallel_workbook_writes(self):
        self.workbook({'颜色': ['Black'], '镜头': ['Single lens']})
        barrier = threading.Barrier(2)
        model = FakeModel(lambda f, v: (barrier.wait(timeout=2),
                                        {'颜色': '黑色', '镜头': '单镜头'}[f])[1])
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(list(self.frame().loc[0, ['颜色', '镜头']]), ['黑色', '单镜头'])

    def test_changes_in_source_during_ai_block_normalized_write(self):
        self.workbook({'线长': ['1 Metre']})
        def edit(*args):
            wb = load_workbook(self.path)
            headers = {c.value: c.column for c in wb.active[1]}
            wb.active.cell(2, headers['商品标题'], 'external edit')
            wb.save(self.path)
            wb.close()
            return '1m'
        with self.assertRaisesRegex(RuntimeError, '其他程序修改'):
            self.run_normalize(FakeModel(edit))
        self.assertEqual(self.frame().loc[0, '线长'], '1 Metre')
        self.assertEqual(self.audit()['fields']['线长']['status'], 'ready')

    def test_changes_in_output_copy_during_ai_are_preserved(self):
        self.workbook({'线长': ['1 Metre']})
        output = self.path.parent / 'copy.xlsx'
        def edit(*args):
            wb = load_workbook(output)
            wb.active['B2'] = 'manual edit'
            wb.save(output)
            wb.close()
            return '1m'
        with self.assertRaisesRegex(RuntimeError, '输出文件.*其他程序修改'):
            self.run_normalize(FakeModel(edit), output=output)
        self.assertEqual(self.frame(output).loc[0, '线长'], 'manual edit')

    def test_replacement_plan_cannot_overwrite_unselected_cells_or_stale_values(self):
        self.workbook({'线长': ['1 Metre', '2 Metre']})
        df = pd.read_excel(self.path, dtype=object)
        writer = FieldWorkbookWriter(self.path, self.path, self.fields)
        writer.save(df)
        before = self.path.read_bytes()
        for index, field, original in [(1, '线长', '2 Metre'), (0, '商品标题', 'title'),
                                       (0, '线长', 'stale')]:
            with self.assertRaises((ValueError, RuntimeError)):
                writer.save(df.loc[[0]], replacements=[dict(index=index, field=field,
                                                           original=original, value='1m')])
            self.assertEqual(self.path.read_bytes(), before)

    def test_corrupt_complete_mapping_is_rebuilt_not_applied(self):
        self.workbook({'线长': ['1 Metre']})
        self.assertTrue(self.run_normalize(FakeModel(length_target))[0])
        audit = self.audit()
        audit['fields']['线长']['mapping'] = {}
        Path(str(self.path) + '.normalization.json').write_text(json.dumps(audit), encoding='utf-8')
        model = FakeModel(length_target)
        self.assertTrue(self.run_normalize(model)[0])
        self.assertEqual(len(model.calls), 1)


class MappingValidationTests(unittest.TestCase):
    def test_zero_false_absence_and_strings_are_distinct(self):
        values = distinct_values(list(enumerate([0, False, '0', '无', '', None, '未知'])))
        self.assertEqual(len(values), 4)
        self.assertEqual([v['value'] for v in values], [0, False, '0', '无'])
        with patch.dict(os.environ, EXTRACT_NORMALIZE='1'):
            self.assertFalse(is_empty_extracted_field('无'))
            self.assertTrue(is_empty_field('无'))

    def test_equivalence_guards(self):
        for source, target in [('100cm', '1m'), ('3ft', '.9144m'),
                               ('15.7inch', '0.39878m'), ('3m/9.84ft', '3m')]:
            with self.subTest(source=source):
                self.assertEqual(unsafe_change(source, target, CONVENTION), '')
        for source, target in [('3ft', '1m'), ('1kWh', '1000W'), ('14AWG', '14mm'),
                               ('3（2）', '3'), ('1m/2m', '1m'), ('3m/16ft', '3m'),
                               ('4USB', '4A'), ('1,000m', '1m')]:
            with self.subTest(source=source):
                self.assertTrue(unsafe_change(source, target, CONVENTION))

    def test_missing_duplicate_and_unknown_ids_are_rejected(self):
        items = distinct_values([[0, 'Black'], [1, 'White']])
        entries = [dict(source_id=v['source_id'], target=v['value'], status='normalized') for v in items]
        for invalid in (entries[:1], entries + entries[:1], entries + [dict(entries[0], source_id='unknown')]):
            with self.assertRaises(ValueError):
                validate_mapping({'convention': TEXT, 'mappings': invalid}, items)

    def test_batch_cannot_change_convention_or_invent_category(self):
        items = distinct_values([[0, 'Black']])
        convention = dict(TEXT, kind='category', labels=['黑色'])
        payload = {'convention': convention, 'mappings': [dict(source_id='v00001', target='白色', status='normalized')]}
        _, result = validate_mapping(payload, items)
        self.assertEqual(result['v00001']['target'], 'Black')
        with self.assertRaises(ValueError):
            validate_mapping(payload, items, TEXT)

    def test_scalar_targets_are_canonical_and_keep_preserves_boolean(self):
        items = distinct_values([[0, '1 Metre'], [1, '100cm']])
        _, mapping = validate_mapping({'convention': CONVENTION, 'mappings': [
            dict(source_id='v00001', target='1.00 m', status='normalized'),
            dict(source_id='v00002', target='1m', status='normalized')]}, items)
        self.assertEqual([v['target'] for v in mapping.values()], ['1m', '1m'])
        _, mapping = validate_mapping({'convention': TEXT, 'mappings': [
            dict(source_id='v00001', target=False, status='keep')]}, distinct_values([[0, False]]))
        self.assertIs(mapping['v00001']['target'], False)


class ScriptNormalizationTests(NormalizationWorkbookFixture):
    def test_both_scripts_normalize_existing_selected_values_and_preserve_absence(self):
        for name in ('exstract_from_description', 'exstract_from_description_fast'):
            with self.subTest(script=name):
                self.workbook({'线长': ['1 Metre', '100cm', '2 Metre'],
                               '颜色': ['无', 'Black', 'Black']}, selected={0, 2})
                audit = Path(str(self.path) + '.normalization.json')
                audit.unlink(missing_ok=True)
                spec = importlib.util.spec_from_file_location(name, Path(__file__).parents[1] / 'scripts' / f'{name}.py')
                module = importlib.util.module_from_spec(spec)
                with patch.dict(os.environ, {'API_KEY': 'test-only', 'CURRENT_PRODUCT': '', 'ENABLE_FETCH': '0',
                    'EXTRACT_CHECKPOINT_FILE': '', 'EXTRACT_FILTER_ENABLED': '1',
                    'EXTRACT_FILTER_COLUMN': '打标', 'EXTRACT_FILTER_VALUE': '1', 'FIELDS': ''}):
                    spec.loader.exec_module(module)
                    module.INPUT_FILE = module.OUTPUT_FILE = str(self.path)
                    module.TARGET_FIELDS = self.fields[:]
                    module.PRODUCT_CONFIG = self.config
                    module.FIELD_CONFIG = self.config['field_config']
                    module.ENABLE_VISION_DEFAULT = False
                    module.SAVE_EVERY = module.CONCURRENCY = 1
                    model = FakeModel(lambda f, v: length_target(f, v) if f == '线长' else
                                      ('黑色' if v == 'Black' else v))
                    if name.endswith('_fast'):
                        with patch.object(module, 'load_product_config', return_value=('test', self.config)), \
                             patch.object(module, 'get_text_client', return_value=model), \
                             patch.object(module, 'check_stop_file', return_value=False), \
                             patch.object(module, 'process_single_row') as extract:
                            module.main()
                    else:
                        module.client = model
                        with patch.object(module, 'ask_enable_vision', return_value=False), \
                             patch.object(module, 'extract_text_features') as extract:
                            module.main()
                    extract.assert_not_called()
                    self.assertEqual(module.filter_valid_fields({'颜色': '无'}, ['颜色']), {'颜色': '无'})
                frame = self.frame()
                self.assertEqual(list(frame['线长']), ['1m', '100cm', '2m'])
                self.assertEqual(list(frame['颜色']), ['无', 'Black', '黑色'])
                self.assertEqual(len(model.calls), 2)

    def load_fast(self):
        spec = importlib.util.spec_from_file_location('normalization_fast',
            Path(__file__).parents[1] / 'scripts/exstract_from_description_fast.py')
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.INPUT_FILE = module.OUTPUT_FILE = str(self.path)
        module.ENABLE_VISION_DEFAULT = False
        module.SAVE_EVERY = module.CONCURRENCY = 1
        return module

    def test_fast_new_answers_are_normalized_and_failed_normalization_can_resume(self):
        self.workbook({'线长': [None, None], '颜色': [None, None]})
        with patch.dict(os.environ, {'API_KEY': 'test-only', 'CURRENT_PRODUCT': '',
            'ENABLE_FETCH': '0', 'EXTRACT_FILTER_ENABLED': '0', 'FIELDS': '', 'EXTRACT_CHECKPOINT_FILE': ''}):
            module = self.load_fast()
            with patch.object(module, 'load_product_config', return_value=('test', self.config)), \
                 patch.object(module, 'extract_text_features_concurrent', return_value={'线长': '100cm', '颜色': 'Black'}) as extract, \
                 patch.object(module, 'get_text_client', return_value=FakeModel(length_target, fail_field='颜色')):
                with self.assertRaises(SystemExit) as error:
                    module.main()
                self.assertEqual(error.exception.code, 2)
                self.assertEqual(extract.call_count, 2)
            self.assertTrue(Path(str(self.path) + '.checkpoint.json').exists())
            self.assertEqual(list(self.frame()['线长']), ['1m', '1m'])
            model = FakeModel(lambda f, v: '黑色' if f == '颜色' else v)
            with patch.object(module, 'load_product_config', return_value=('test', self.config)), \
                 patch.object(module, 'extract_text_features_concurrent') as extract, \
                 patch.object(module, 'get_text_client', return_value=model):
                module.main()
            extract.assert_not_called()
            self.assertEqual([c['field'] for c in model.calls], ['颜色'])
            self.assertEqual(list(self.frame()['颜色']), ['黑色', '黑色'])
            self.assertFalse(Path(str(self.path) + '.checkpoint.json').exists())


class NormalizationUITests(unittest.TestCase):
    def test_default_on_and_explicit_off_are_passed_to_both_script_modes(self):
        from tests import support
        from modules.market_planning import ModuleWidget
        support.qt_app()
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / 'fields.json'
            config_path.write_text(json.dumps({'products': {'test': {'name': 'test',
                'target_fields': ['线长'], 'field_config': {'线长': '提取线长'}}}}), encoding='utf-8')
            config = SimpleNamespace(resolve_config=lambda name: str(config_path))
            widget = ModuleWidget(config, None)
            self.addCleanup(widget.deleteLater)
            self.assertTrue(widget.normalize_cb.isChecked())
            widget.tag_input.setText('not-executed.xlsx')
            for enabled, fast in ((True, True), (False, False)):
                widget.normalize_cb.setChecked(enabled)
                widget.fast_cb.setChecked(fast)
                widget.fetch_cb.setChecked(False)
                with patch.object(widget, '_run', return_value=True) as run:
                    widget._run_tagging()
                args = run.call_args.args
                self.assertEqual(args[0], 'exstract_from_description_fast.py' if fast else 'exstract_from_description.py')
                self.assertEqual(args[2]['EXTRACT_NORMALIZE'], '1' if enabled else '0')
