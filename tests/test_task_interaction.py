"""User-facing task flow, asynchronous field choices and result safety."""
from tests import support

import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from PyQt6.QtCore import QThreadPool
from PyQt6.QtWidgets import QTabWidget, QPushButton, QScrollArea

from core.config_manager import ConfigManager
from core.main_window import MainWindow
from core.theme import THEME
from core.task_ui import TaskBar


class TaskInteractionTests(unittest.TestCase):
    def setUp(self):
        self.app = support.qt_app()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(support.cleanup_qt_widgets)
        self.root = Path(self.tmp.name)
        config = ConfigManager(str(self.root / 'config.json'))
        config.update({'auto_check_app_update': False, 'auto_check_update': False})
        with patch('core.main_window.ConfigManager', return_value=config):
            self.window = MainWindow()
        self.window.resize(880, 620)
        self.window.show()
        self.app.processEvents()

    def tearDown(self):
        QThreadPool.globalInstance().waitForDone(5000)
        THEME.set_theme('light')

    def module(self, name):
        index, widget = next((i, w) for i, (info, w) in enumerate(self.window._modules)
                             if info['name'] == name)
        self.window._switch_module(index)
        self.app.processEvents()
        return widget

    def wait_for(self, condition):
        deadline = time.monotonic() + 5
        while not condition() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(.01)
        self.assertTrue(condition())

    def workbook(self, name, sheets):
        path = self.root / name
        book = Workbook()
        book.remove(book.active)
        for title, columns in sheets.items():
            book.create_sheet(title).append(columns)
        book.save(path)
        book.close()
        return path

    def test_primary_actions_stay_inside_small_window_in_both_themes(self):
        for theme in ('light', 'dark'):
            THEME.set_theme(theme)
            self.app.setStyleSheet(THEME.qss())
            for name in ('销量数据', '产品矩阵', '市场规划'):
                module = self.module(name)
                tabs = module.findChild(QTabWidget)
                for index in range(tabs.count()):
                    tabs.setCurrentIndex(index)
                    self.app.processEvents()
                    for button in module.findChildren(QPushButton):
                        if button.isVisible() and button.property('buttonType') == 'primary':
                            position = button.mapTo(self.window, button.rect().bottomRight())
                            self.assertTrue(self.window.rect().contains(position), (name, index, theme))

    def test_missing_fields_block_run_and_show_inline_hints(self):
        matrix = self.module('产品矩阵')
        with patch.object(matrix, '_run') as run:
            self.assertFalse(matrix._run_matrix2())
        run.assert_not_called()
        self.assertTrue(matrix.matrix2_sheet.property('validationError'))
        self.assertIn('Sheet', matrix.matrix2_sheet._validation_hint.text())
        matrix.matrix2_sheet.setText('数据')
        self.assertFalse(matrix.matrix2_sheet.property('validationError'))
        self.assertTrue(matrix.matrix2_sheet._validation_hint.isHidden())

    def test_field_filter_defaults_and_worker_parameters(self):
        module = self.module('市场规划')
        self.assertFalse(module.tag_filter_cb.isChecked())
        self.assertFalse(module.tag_filter_column.isEnabled())
        path = self.workbook('products.xlsx', {'商品': ['ASIN', '打标']})
        module.tag_input.setText(str(path))
        self.wait_for(lambda: bool(module.tag_filter_choices.metadata))
        self.assertIn('打标', [module.tag_filter_column.itemText(i)
                            for i in range(module.tag_filter_column.count())])
        with patch.object(module, '_run', return_value=False) as run:
            module._run_tagging()
            self.assertEqual(run.call_args.args[2]['EXTRACT_FILTER_ENABLED'], '0')
            module.tag_filter_cb.setChecked(True)
            self.assertTrue(module.tag_filter_column.isEnabled())
            module._run_tagging()
            env = run.call_args.args[2]
            self.assertEqual((env['EXTRACT_FILTER_ENABLED'], env['EXTRACT_FILTER_COLUMN'],
                              env['EXTRACT_FILTER_VALUE']), ('1', '打标', '1'))
            module.tag_filter_value.clear()
            run.reset_mock()
            module._run_tagging()
            run.assert_not_called()

    def test_sheet_and_column_choices_are_async_and_keep_manual_text(self):
        matrix = self.module('产品矩阵')
        path = self.workbook('source.xlsx', {'元数据': ['ASIN', '品牌', '颜色'], '月数据': ['ASIN', '2026-10']})
        matrix.matrix2_sheet.setText('元数据')
        matrix.matrix2_row_col.setText('手动列名')
        matrix.matrix2_file.setText(str(path))
        choices = matrix._excel_choices[1]
        self.wait_for(lambda: len(choices.metadata) == 2)
        self.assertGreaterEqual(matrix.matrix2_sheet.findText('月数据'), 0)
        self.assertGreaterEqual(matrix.matrix2_row_col.findText('颜色'), 0)
        self.assertEqual(matrix.matrix2_row_col.text(), '手动列名')
        self.assertEqual(matrix.matrix2_data_sheet.text(), '')
        matrix.matrix2_sheet.setText('月数据')
        self.assertGreaterEqual(matrix.matrix2_row_col.findText('2026-10'), 0)
        self.assertEqual(matrix.matrix2_row_col.findText('颜色'), -1)
        generation = choices.generation
        matrix.matrix2_file.setText(str(self.root / 'missing.xlsx'))
        choices._loaded(generation, {'旧文件': ['旧列']}, '')
        self.assertEqual(choices.metadata, {})
        self.assertEqual(matrix.matrix2_sheet.findText('旧文件'), -1)
        self.assertEqual(matrix.matrix2_row_col.text(), '手动列名')

    def test_corrupt_workbook_retains_manual_input(self):
        matrix = self.module('产品矩阵')
        path = self.root / 'broken.xlsx'
        path.write_bytes(b'broken zip')
        matrix.matrix2_sheet.setText('自填工作表')
        matrix.matrix2_file.setText(str(path))
        self.wait_for(lambda: '读取失败' in matrix.matrix2_sheet.toolTip())
        self.assertEqual(matrix.matrix2_sheet.text(), '自填工作表')

    def test_results_ignore_stale_files_and_failed_runs(self):
        output = self.root / 'result.xlsx'
        output.write_bytes(b'old')
        bar = TaskBar()
        bar.begin([str(output)])
        bar.finish(0)
        self.assertTrue(bar.result_btn.isHidden())
        bar.begin([str(output)])
        output.write_bytes(b'new contents')
        bar.finish(0)
        self.assertEqual(bar.result_path, str(output))
        self.assertFalse(bar.result_btn.isHidden())
        bar.begin([str(output)])
        bar.finish(1)
        self.assertTrue(bar.result_btn.isHidden())
        self.assertIn('失败', bar.summary.text())

    def test_log_results_prefer_final_output_and_ignore_backup(self):
        backup = self.root / 'backup.xlsx'
        result = self.root / 'final.xlsx'
        backup.write_bytes(b'backup')
        result.write_bytes(b'final')
        bar = TaskBar()
        bar.begin()
        bar.observe(f'备份已保存：{backup}')
        bar.observe(f'处理完成，结果已输出到：{result}')
        bar.finish(0)
        self.assertEqual(bar.result_path, str(result))

    def test_field_search_preserves_selection_and_shows_count(self):
        market = self.module('市场规划')
        self.assertGreater(len(market._field_cbs), 0)
        selected = market._get_selected_fields()
        market.field_search.setText('不存在的字段')
        self.assertEqual(market._get_selected_fields(), selected)
        self.assertTrue(all(cb.isHidden() for cb in market._field_cbs))
        market.field_search.clear()
        self.assertIn(str(len(selected)), market.field_count.text())
        self.assertTrue(all(not cb.isHidden() for cb in market._field_cbs))

    def test_empty_library_disables_actions_and_product_list_can_collapse(self):
        reviews = self.module('评论分析').review_collection
        reviews._reviews_loaded(([], '本地', ''))
        self.assertFalse(reviews.library_analyze_btn.isEnabled())
        self.assertFalse(reviews.library_export_btn.isEnabled())
        reviews._reviews_loaded(([{'content': 'test', 'rating': 5}], '本地', ''))
        self.assertTrue(reviews.library_analyze_btn.isEnabled())
        self.assertTrue(reviews.products_table.isHidden())
        reviews.products_toggle.click()
        self.assertFalse(reviews.products_table.isHidden())
        reviews.products_toggle.click()
        self.assertTrue(reviews.products_table.isHidden())

    def test_legacy_matrix_export_requires_a_preview_with_products(self):
        matrix = self.module('产品矩阵').free_matrix
        with patch('core.matrix_workspace.QFileDialog.getSaveFileName') as choose:
            matrix.export()
        choose.assert_not_called()
        path = self.workbook('products.xlsx', {'数据': ['ASIN', '品牌']})
        from openpyxl import load_workbook
        book = load_workbook(path)
        book.active.append(['B000000001', 'A'])
        book.save(path)
        book.close()
        matrix.path.setText(str(path))
        matrix.read_data()
        output = self.root / 'legacy-matrix.xlsx'
        with patch('core.matrix_workspace.QFileDialog.getSaveFileName',
                   return_value=(str(output), '')):
            matrix.export()
        self.assertTrue(output.exists())
        import pandas as pd
        self.assertEqual(pd.read_excel(output, sheet_name='产品明细')['ASIN'].tolist(), ['B000000001'])
        matrix.search.setText('不存在的商品')
        matrix.render()
        with patch('core.matrix_workspace.QFileDialog.getSaveFileName') as choose:
            matrix.export()
        choose.assert_not_called()

    def test_unified_split_runs_file_and_folder_with_result_link(self):
        import pandas as pd
        from scripts.batch_sales_split import output_path
        sales = self.module('销量数据')
        source = self.root / 'Competitor-2026.10.xlsx'
        pd.DataFrame({'ASIN': ['A', 'B'], '父ASIN': ['P', 'P'],
                      '月销量': [100, 100], '子体销量': [1, 3],
                      '价格(€)': [2.5, 7.25]}).to_excel(source, index=False)
        for path in (source, self.root):
            sales.split_input.setText(str(path))
            sales.run_btn1.click()
            bar = sales.clean_actions
            self.wait_for(lambda: not bar.running)
            result = output_path(source)
            self.assertEqual(bar.result_path, result)
            self.assertTrue(sales.run_btn1.isEnabled())
            self.assertEqual(pd.read_excel(result)['子体销售额_矫正'].tolist(), [62.5, 543.75])
        self.assertEqual(len(sales.clean_tasks.buttons()), 2)

    def test_extraction_finishes_with_the_generated_workbook_link(self):
        sales = self.module('销量数据')
        path = self.workbook('US_2026-10_销量拆分.xlsx',
                             {'数据': ['ASIN', '子体销量_矫正', '子体销售额_矫正']})
        from openpyxl import load_workbook
        book = load_workbook(path)
        book.active.append(['B000000001', 2, 20])
        book.save(path)
        book.close()
        output = self.root / '历史.xlsx'
        sales.trend_folder.setText(str(self.root))
        sales.trend_output.setText(str(output))
        sales._run_extract_data()
        bar = sales.run_btn2._task_bar
        self.wait_for(lambda: not bar.running)
        self.assertTrue(output.is_file())
        self.assertEqual(bar.result_path, str(output))
        self.assertTrue(sales.run_btn2.isEnabled())


if __name__ == '__main__':
    unittest.main()
