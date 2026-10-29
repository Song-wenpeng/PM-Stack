"""Matrix-2 CLI, safe-save and amount-statistics regressions."""
from tests import support

from contextlib import chdir, redirect_stdout
import io
import math
import os
from pathlib import Path
import runpy
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts/product_matrix2.py'


class AmountParsingTests(unittest.TestCase):
    """R03: mixed-format amounts must survive, and non-amounts must be reported."""

    CASES = [
        (1200, 1200.0, ''),
        (1200.5, 1200.5, ''),
        ('1,200', 1200.0, ''),
        ('€1,200.50', 1200.5, ''),
        ('$300', 300.0, ''),
        ('1.234,56', 1234.56, ''),
        ('1,234.56', 1234.56, ''),
        ('(200)', -200.0, ''),
        (' 50 ', 50.0, ''),
        ('-15', -15.0, ''),
        ('0', 0.0, ''),
        (None, None, '空值'),
        (float('nan'), None, '空值'),
        ('', None, '空值'),
        ('   ', None, '空值'),
        ('-', None, '空值'),
        ('待定', None, '不可解析'),
        ('约1千', None, '不可解析'),
        ('12ab', None, '不可解析'),
        (True, None, '不可解析'),
        (float('inf'), None, '不可解析'),
        ('inf', None, '不可解析'),
    ]

    def test_amount_formats_and_rejections(self):
        from scripts.product_matrix2 import parse_amount
        for value, expected, reason in self.CASES:
            with self.subTest(value=repr(value)):
                number, actual = parse_amount(value)
                self.assertEqual(actual, reason)
                if expected is None:
                    self.assertTrue(math.isnan(number))
                else:
                    self.assertAlmostEqual(number, expected, places=6)

    def test_percent_is_not_treated_as_an_amount(self):
        from scripts.product_matrix2 import parse_amount
        number, reason = parse_amount('12%')
        self.assertEqual(reason, '不可解析')
        self.assertTrue(math.isnan(number))


class ProductMatrix2Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'source.xlsx'
        with pd.ExcelWriter(self.source, engine='openpyxl') as writer:
            pd.DataFrame({'行': ['a'], '列': ['b'], '金额': [10]}).to_excel(writer, sheet_name='数据', index=False)
            pd.DataFrame({'原件': ['=SUM(1,2)']}).to_excel(writer, sheet_name='保留', index=False)
        self.original = self.source.read_bytes()

    def execute(self, output=None):
        args = [str(SCRIPT), '--file', str(self.source), '--sheet', '数据',
                '--row-col', '行', '--col-col', '列', '--value-col', '金额']
        if output is not None:
            args += ['--output', str(output)]
        self.log = io.StringIO()
        with patch('sys.argv', args), redirect_stdout(self.log):
            try:
                runpy.run_path(str(SCRIPT), run_name='__main__')
            except SystemExit as exc:
                return exc.code
        return 0

    def test_cli_rejects_same_source_before_write(self):
        code = self.execute(self.source)
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertNotEqual(code, 0)

    def test_failed_write_preserves_existing_target(self):
        from scripts.product_matrix2 import export_excel
        target = self.root / 'existing.xlsx'
        target.write_bytes(self.original)
        with patch.object(pd.DataFrame, 'to_excel', side_effect=OSError('synthetic failure')):
            with self.assertRaises(Exception):
                export_excel({'a': {'b': {'sales': 10, 'count': 1, 'avg': 10}}}, ['a'], ['b'], '行', '列', target)
        self.assertEqual(target.read_bytes(), self.original)

    def test_cli_case_relative_and_normalized_aliases_reject_before_read(self):
        for output in (str(self.source).upper(), './source.xlsx', 'missing/../source.xlsx'):
            with self.subTest(output=output), chdir(self.root), patch.object(pd, 'read_excel', side_effect=AssertionError('must reject first')):
                self.assertNotEqual(self.execute(output), 0)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_cli_default_output_cannot_equal_source(self):
        self.source = self.root / 'product_matrix.xlsx'
        self.source.write_bytes(self.original)
        with chdir(self.root), patch.object(pd, 'read_excel', side_effect=AssertionError('must reject first')):
            self.assertNotEqual(self.execute(), 0)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_cli_separate_export_preserves_source_and_sheet_contents(self):
        target = self.root / 'matrix.xlsx'
        self.assertEqual(self.execute(target), 0)
        self.assertEqual(self.source.read_bytes(), self.original)
        frames = pd.read_excel(target, sheet_name=None)
        self.assertEqual(list(frames), ['产品矩阵', '参数信息'])
        self.assertEqual(frames['产品矩阵'].iloc[0, 1], '10 (1个, 10.0均)')

    def test_cli_blank_output_uses_default_and_preserves_source(self):
        with chdir(self.root):
            self.assertEqual(self.execute(''), 0)
        self.assertTrue((self.root / 'product_matrix.xlsx').is_file())
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_cli_100_rows_multisheet_formula_style_source_is_unchanged(self):
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill
        pd.DataFrame({'行': ['a'] * 100, '列': ['b'] * 100, '金额': [10] * 100}).to_excel(
            self.source, sheet_name='数据', index=False)
        book = load_workbook(self.source)
        book.active['A1'].fill = PatternFill('solid', fgColor='00FF00')
        book.create_sheet('保留公式').append(['=SUM(1,2)'])
        book.save(self.source)
        book.close()
        original = self.source.read_bytes()
        self.assertEqual(self.execute(self.root / 'separate.xlsx'), 0)
        self.assertEqual(self.source.read_bytes(), original)

    def test_cli_hardlink_to_source_is_rejected(self):
        alias = self.root / 'alias.xlsx'
        os.link(self.source, alias)
        with patch.object(pd, 'read_excel', side_effect=AssertionError('must reject first')):
            self.assertNotEqual(self.execute(alias), 0)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_direct_cli_subprocess_has_project_imports_and_source_guard(self):
        command = [sys.executable, str(SCRIPT), '--file', str(self.source), '--sheet', '数据',
                   '--row-col', '行', '--col-col', '列', '--value-col', '金额', '--output']
        target = self.root / 'subprocess.xlsx'
        result = subprocess.run(command + [str(target)], cwd=self.root, env=dict(os.environ, PYTHONIOENCODING='utf-8'),
                                capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))
        result = subprocess.run(command + [str(self.source)], cwd=self.root, capture_output=True, timeout=30)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertTrue(target.exists())

    def test_cli_permanent_replace_failure_retains_target_and_recovery(self):
        target = self.root / 'matrix.xlsx'
        target.write_bytes(self.original)
        with patch('core.report_io.os.replace', side_effect=PermissionError('locked')), patch('core.report_io.time.sleep'):
            self.assertNotEqual(self.execute(target), 0)
        self.assertEqual(target.read_bytes(), self.original)
        candidates = list(self.root.glob('matrix.pending-*.xlsx'))
        self.assertEqual(len(candidates), 1)
        self.assertEqual(list(pd.read_excel(candidates[0], sheet_name=None)), ['产品矩阵', '参数信息'])
        self.assertIn(str(candidates[0]), self.log.getvalue())

    def test_cli_transient_replace_failure_recovers(self):
        target = self.root / 'matrix.xlsx'
        real_replace = os.replace
        calls = []
        def replace(source, destination):
            calls.append(source)
            if len(calls) < 3:
                raise PermissionError('scanner')
            real_replace(source, destination)
        with patch('core.report_io.os.replace', side_effect=replace), patch('core.report_io.time.sleep'):
            self.assertEqual(self.execute(target), 0)
        self.assertEqual(len(calls), 3)
        self.assertEqual(self.source.read_bytes(), self.original)


class Matrix2StatisticsTests(unittest.TestCase):
    """R03: amounts must not be silently dropped, zero-faked or double-counted."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'source.xlsx'
        self.log = io.StringIO()

    def write_source(self, metadata, data_sheet=None):
        with pd.ExcelWriter(self.source, engine='openpyxl') as writer:
            metadata.to_excel(writer, sheet_name='元数据', index=False)
            if data_sheet is not None:
                data_sheet.to_excel(writer, sheet_name='子体销额', index=False)
        self.original = self.source.read_bytes()

    def execute(self, extra_args=(), output=None):
        target = output if output is not None else self.root / 'matrix.xlsx'
        args = [str(SCRIPT), '--file', str(self.source), '--sheet', '元数据',
                '--row-col', '行', '--col-col', '列', *extra_args, '--output', str(target)]
        self.log = io.StringIO()
        with patch('sys.argv', args), redirect_stdout(self.log):
            try:
                runpy.run_path(str(SCRIPT), run_name='__main__')
            except SystemExit as exc:
                return exc.code or 0, target
        return 0, target

    def cell(self, target, row=0, column=1):
        return pd.read_excel(target, sheet_name='产品矩阵').iloc[row, column]

    def info(self, target):
        frame = pd.read_excel(target, sheet_name='参数信息')
        return dict(zip(frame['参数'].astype(str), frame['值'].astype(str)))

    def test_thousands_and_currency_amounts_are_not_silently_dropped(self):
        self.write_source(pd.DataFrame({'行': ['a', 'a'], '列': ['b', 'b'],
                                        '金额': ['1,200', '€300']}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '1,500 (2个, 750.0均)')
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_european_and_accounting_amount_formats_parse(self):
        self.write_source(pd.DataFrame({'行': ['a'] * 3, '列': ['b'] * 3,
                                        '金额': ['1.234,56', '(200)', ' 50 ']}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '1,085 (3个, 361.5均)')

    def test_all_missing_amounts_are_reported_as_missing_not_zero(self):
        self.write_source(pd.DataFrame({'行': ['a', 'a'], '列': ['b', 'b'],
                                        '金额': [None, '']}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        text = self.cell(target)
        self.assertNotIn('0 (2个', text)
        self.assertIn('无有效金额', text)
        self.assertIn('2', text)
        frame = pd.read_excel(target, sheet_name='产品矩阵')
        self.assertIn('无有效金额', str(frame.iloc[-1, 1]))
        self.assertNotIn('0 (2个', str(frame.iloc[-1, 1]))

    def test_unparseable_amounts_are_counted_with_samples_not_dropped(self):
        # 'N/A' 之类会被 pandas 读表阶段转成 NaN（属真实缺失），
        # 这里用读表后仍为文本、但确实不是金额的值来验证不被静默丢弃。
        self.write_source(pd.DataFrame({'行': ['a'] * 3, '列': ['b'] * 3,
                                        '金额': ['1,200', '待定', '约1千']}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '1,200 (1个, 1,200.0均, 另2行无有效金额)')
        self.assertIn('待定', self.log.getvalue())
        self.assertIn('无法解析', self.log.getvalue())
        self.assertEqual(self.info(target)['不可解析金额行数'], '2')
        self.assertEqual(self.info(target)['有效金额行数'], '1')

    def test_partial_missing_amounts_average_over_valid_rows_only(self):
        self.write_source(pd.DataFrame({'行': ['a'] * 3, '列': ['b'] * 3,
                                        '金额': [100, 200, None]}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '300 (2个, 150.0均, 另1行无有效金额)')

    def test_duplicate_asin_in_data_sheet_does_not_inflate_sales(self):
        self.write_source(
            pd.DataFrame({'ASIN': ['A1', 'A2'], '行': ['a', 'a'], '列': ['b', 'b']}),
            pd.DataFrame({'ASIN': ['A1', 'A1', 'A2'], '2025-12': [100, 100, 50]}))
        code, target = self.execute(['--data-sheet', '子体销额', '--month-col', '2025-12'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '150 (2个, 75.0均)')
        self.assertIn('重复', self.log.getvalue())
        self.assertEqual(self.info(target)['子表重复ASIN数'], '1')
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_data_sheet_unparseable_amounts_keep_samples_after_merge(self):
        self.write_source(
            pd.DataFrame({'ASIN': ['A1', 'A2', 'A3', 'A4'],
                          '行': ['a'] * 4, '列': ['b'] * 4}),
            pd.DataFrame({'ASIN': ['A1', 'A2', 'A3'],
                          '2025-12': ['€1,200', '待定', '约1千']}))
        code, target = self.execute(['--data-sheet', '子体销额', '--month-col', '2025-12'])
        self.assertEqual(code, 0)
        info = self.info(target)
        self.assertEqual((info['有效金额行数'], info['空金额行数'], info['不可解析金额行数']),
                         ('1', '1', '2'))
        self.assertIn('待定', info['不可解析金额样例'])
        self.assertIn('约1千', self.log.getvalue())
        self.assertEqual(self.cell(target), '1,200 (1个, 1,200.0均, 另3行无有效金额)')

    def test_duplicate_asin_in_metadata_is_reported_without_silent_removal(self):
        self.write_source(pd.DataFrame({'ASIN': ['A1', 'A1'], '行': ['a', 'a'],
                                        '列': ['b', 'b'], '金额': [100, 100]}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '200 (2个, 100.0均)')
        self.assertEqual(self.info(target)['元数据重复ASIN行数'], '1')
        self.assertIn('元数据存在重复 ASIN', self.log.getvalue())

    def test_pandas_null_text_is_counted_as_missing_not_unparseable(self):
        self.write_source(pd.DataFrame({'行': ['a', 'a'], '列': ['b', 'b'],
                                        '金额': ['N/A', '1,000']}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        info = self.info(target)
        self.assertEqual((info['有效金额行数'], info['空金额行数'], info['不可解析金额行数']),
                         ('1', '1', '0'))
        self.assertEqual(self.cell(target), '1,000 (1个, 1,000.0均, 另1行无有效金额)')

    def test_valid_data_keeps_existing_cell_format(self):
        self.write_source(pd.DataFrame({'行': ['a'], '列': ['b'], '金额': [10]}))
        code, target = self.execute(['--value-col', '金额'])
        self.assertEqual(code, 0)
        self.assertEqual(self.cell(target), '10 (1个, 10.0均)')


class Matrix2UIProtectionTests(unittest.TestCase):
    def setUp(self):
        from core.config_manager import ConfigManager
        from modules.product_matrix import ModuleWidget
        self.app = support.qt_app()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(support.cleanup_qt_widgets)
        self.root = Path(self.tmp.name)
        self.runner = MagicMock()
        self.widget = ModuleWidget(ConfigManager(str(self.root / 'config.json')), self.runner)
        self.widget.matrix2_file.setText(str(self.root / 'source.xlsx'))
        self.widget.matrix2_sheet.setText('数据')
        self.widget.matrix2_row_col.setText('行')
        self.widget.matrix2_col_col.setText('列')
        self.widget.matrix2_value_col.setText('金额')

    def tearDown(self):
        from PyQt6.QtCore import QThreadPool
        QThreadPool.globalInstance().waitForDone(5000)

    def test_ui_same_source_never_starts_runner(self):
        for output in (str(self.root / 'source.xlsx'), str(self.root / 'source.xlsx').upper(), './source.xlsx'):
            with self.subTest(output=output), chdir(self.root), patch.object(self.widget, '_run') as run:
                self.widget.matrix2_output.setText(output)
                self.assertFalse(self.widget._run_matrix2())
                run.assert_not_called()
        self.runner.run_script.assert_not_called()

    def test_ui_default_same_source_never_starts_runner(self):
        self.widget.matrix2_file.setText(str(self.root / 'product_matrix.xlsx'))
        with chdir(self.root), patch.object(self.widget, '_run') as run:
            self.assertFalse(self.widget._run_matrix2())
        run.assert_not_called()
        self.runner.run_script.assert_not_called()

    def test_ui_separate_output_passes_absolute_path(self):
        with chdir(self.root), patch.object(self.widget, '_run') as run:
            self.widget._run_matrix2()
        args = run.call_args.kwargs['args']
        self.assertEqual(args[args.index('--output') + 1], str(self.root / 'product_matrix.xlsx'))
