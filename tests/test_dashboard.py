"""Calculation, import/export and navigation regressions for the market dashboard."""
from tests import support
from pathlib import Path
from dataclasses import replace
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
import numpy as np
from openpyxl import load_workbook
from PyQt6.QtCore import Qt, QPoint

from core.dashboard_data import (normalize, from_history, Query, analyze, export_report,
                                 product_snapshot, month_text)
from core.config_manager import ConfigManager
from core.dashboard_widgets import ImportDialog
from modules.display_dashboard import ModuleWidget, LoadWorkbook


def fixture():
    return pd.DataFrame([
        {'ASIN': asin, '月份': month, '销量': units, '销售额': units * price,
         '品牌': brand, '站点': 'US', '币种': 'USD', '品类': '排插', '插位数': sockets}
        for month, factor in [('2024-12', .5), ('2025-01', .5), ('2025-12', 1), ('2026-01', 2)]
        for asin, brand, units, price, sockets in [('A1', 'A', 10 * factor, 10, '6'), ('B1', 'B', 30 * factor, 20, '8')]
    ])


def dataset(frame=None):
    return normalize(fixture() if frame is None else frame,
                     {c: c for c in ['ASIN', '月份', '销量', '销售额']}, source='synthetic.xlsx')


def query(**kwargs):
    return replace(Query('2026-01', '2026-01', 'US', 'USD'), **kwargs)


class DashboardDataTests(unittest.TestCase):
    def test_fixed_denominator_and_weighted_price(self):
        result = analyze(dataset(), query(brand='A'))
        self.assertEqual(result['kpi']['销售额'], 200)
        self.assertAlmostEqual(result['kpi']['样本份额'], 1 / 7)
        self.assertEqual(len(result['market']), 2)
        self.assertEqual(len(result['detail']), 1)
        all_result = analyze(dataset(), query())
        self.assertEqual(all_result['kpi']['加权均价'], 17.5)
        self.assertEqual(all_result['kpi']['同比'], 3)
        self.assertEqual(all_result['kpi']['上期比'], 1)

    def test_missing_is_unknown_zero_denominator_not_infinite(self):
        frame = fixture()
        frame.loc[frame['品牌'] == 'B', '销售额'] = np.nan
        result = analyze(dataset(frame), query(brand='A'))
        self.assertTrue(pd.isna(result['kpi']['样本份额']))
        self.assertTrue(result['ranking']['份额'].isna().all())
        self.assertTrue(pd.isna(result['trend'].iloc[0]['市场数值']))
        frame['销量'], frame['销售额'] = 0, 0
        result = analyze(dataset(frame), query())
        self.assertTrue(pd.isna(result['kpi']['加权均价']))
        self.assertEqual(result['kpi']['活跃ASIN'], 0)

    def test_period_comparison_requires_complete_months(self):
        result = analyze(dataset(), query(start='2025-11', end='2026-01'))
        self.assertEqual(result['missing_months'], ['2025-11'])
        self.assertTrue(pd.isna(result['kpi']['同比']))
        self.assertTrue(pd.isna(result['kpi']['上期比']))
        self.assertTrue(pd.isna(result['trend'].iloc[0]['市场数值']))

    def test_attribute_and_price_filters_define_market(self):
        result = analyze(dataset(), query(filters={'插位数': ['6']}, price_max=15))
        self.assertEqual(set(result['market']['品牌']), {'A'})
        self.assertEqual(result['kpi']['样本份额'], 1)
        with self.assertRaises(ValueError):
            analyze(dataset(), query(filters={'品牌': ['A']}))
        with self.assertRaises(ValueError):
            analyze(dataset(), query(price_min=20, price_max=10))

    def test_site_currency_isolation(self):
        frame = fixture()
        gb = frame.assign(站点='UK', 币种='GBP', 销售额=1e8)
        result = analyze(dataset(pd.concat([frame, gb])), query())
        self.assertEqual(result['kpi']['销售额'], 1400)

    def test_duplicate_conflict_and_missing_identity(self):
        frame = fixture()
        loaded = dataset(pd.concat([frame, frame.iloc[:1]]))
        self.assertEqual(len(loaded.frame), len(frame))
        self.assertIn('重复', loaded.notes[0])
        conflict = frame.iloc[:1].assign(销售额=999)
        with self.assertRaisesRegex(ValueError, '冲突'):
            dataset(pd.concat([frame, conflict]))
        with self.assertRaises(ValueError):
            dataset(frame.assign(ASIN=''))

    def test_invalid_numbers_and_months(self):
        for value in ['oops', -1, np.inf]:
            with self.assertRaises(ValueError):
                dataset(fixture().assign(销量=value))
        for value in ['2025-13', 'not a date', None, 45292]:
            with self.assertRaises(ValueError):
                month_text(value)
        self.assertEqual(month_text('2025年1月'), '2025-01')
        self.assertEqual(month_text(pd.Timestamp('2025-01-15')), '2025-01')

    def test_history_adapter_matches_long_format_and_preserves_gaps(self):
        frame = fixture()
        meta = frame[['ASIN', '品牌', '插位数']].drop_duplicates()
        units = frame.pivot(index='ASIN', columns='月份', values='销量').reset_index()
        revenue = frame.pivot(index='ASIN', columns='月份', values='销售额').reset_index()
        loaded = from_history(meta, units, revenue, {'ASIN': 'ASIN', '品牌': '品牌'}, 'US', 'USD')
        self.assertEqual(analyze(loaded, query())['kpi']['销售额'], 1400)
        missing = from_history(meta, units, revenue.drop(columns='2026-01'), {'ASIN': 'ASIN'}, 'US', 'USD')
        self.assertTrue(pd.isna(analyze(missing, query())['kpi']['销售额']))
        with self.assertRaises(ValueError):
            from_history(meta.iloc[:1], units, revenue, {'ASIN': 'ASIN'}, 'US', 'USD')
        with self.assertRaises(ValueError):
            from_history(meta, pd.concat([units, units.iloc[:1]]), revenue, {'ASIN': 'ASIN'}, 'US', 'USD')

    def test_top_and_other_preserve_total_and_handle_name_collision(self):
        frame = fixture().iloc[:1]
        frame = pd.concat([frame.assign(ASIN=f'A{i}', 品牌=name) for i, name in enumerate(['其余分组合计', 'A', 'B', 'C', 'D'])])
        result = analyze(dataset(frame), query(start='2024-12', end='2024-12', top=3))
        self.assertAlmostEqual(result['top']['份额'].sum(), 1)
        self.assertEqual(len(result['top']), 4)
        self.assertNotEqual(result['others_label'], '其余分组合计')

    def test_export_traceability_and_formula_safety(self):
        loaded = dataset(fixture().assign(备注='=1+1'))
        q = query()
        result = analyze(loaded, q)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'report.xlsx'
            export_report(path, loaded, q, result, detail_context={'明细搜索': 'test'})
            book = load_workbook(path)
            self.assertEqual(len(book.sheetnames), 7)
            self.assertEqual(len(book['看板图表']._charts), 2)
            detail = book['商品月度明细']
            header = [c.value for c in detail[1]]
            self.assertEqual(detail.cell(2, header.index('备注') + 1).data_type, 's')
            self.assertEqual(detail.max_row, 3)
            self.assertIn('来源', [row[0].value for row in book['统计口径']])
            book.close()

    def test_matrix_snapshot_aggregates_months(self):
        result = analyze(dataset(), query(start='2025-12', end='2026-01'))
        products = product_snapshot(result['detail'])
        self.assertEqual(len(products), 2)
        self.assertEqual(products.set_index('ASIN').loc['A1', '销售额'], 300)
        self.assertEqual(products.set_index('ASIN').loc['B1', '价格($)'], 20)


class DashboardUITests(unittest.TestCase):
    def setUp(self):
        self.app = support.qt_app()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.addCleanup(support.cleanup_qt_widgets)
        self.config = ConfigManager(str(Path(self.temp.name) / 'config.json'))
        self.widget = ModuleWidget(self.config, None)
        self.widget.resize(1100, 800)
        self.widget.set_dataset(dataset())
        self.widget.start.setCurrentText('2026-01')
        self.widget.show()
        self.app.processEvents()

    def test_drill_search_sort_and_empty_result(self):
        self.widget.drill_group('A')
        self.assertEqual(self.widget.proxy.rowCount(), 1)
        self.assertAlmostEqual(self.widget.result['kpi']['样本份额'], 1 / 7)
        self.widget.search.setText('missing')
        self.assertEqual(self.widget.proxy.rowCount(), 0)
        self.assertFalse(self.widget.matrix_btn.isEnabled())
        self.widget.reset_filters()
        self.widget.start.setCurrentText('2026-01')
        col = list(self.widget.model.frame.columns).index('销售额')
        self.widget.proxy.sort(col, Qt.SortOrder.DescendingOrder)
        self.assertEqual(self.widget.visible_frame().iloc[0]['ASIN'], 'B1')
        self.widget.drill_month('2026-01')
        self.assertEqual(self.widget.proxy.rowCount(), 2)
        self.widget.filters = {'插位数': ['does-not-exist']}
        self.widget.refresh()
        self.assertEqual(self.widget.proxy.rowCount(), 0)
        self.assertTrue(pd.isna(self.widget.result['kpi']['销售额']))

    def test_real_import_mapping_worker_and_export(self):
        path = Path(self.temp.name) / 'input.xlsx'
        fixture().to_excel(path, index=False)
        dialog = ImportDialog(str(path), self.widget)
        self.assertEqual(dialog.fields['月份'].currentData(), '月份')
        worker = LoadWorkbook(str(path), dialog.options())
        loaded, errors = [], []
        worker.signals.done.connect(loaded.append)
        worker.signals.failed.connect(errors.append)
        worker.run()
        self.assertFalse(errors)
        self.assertEqual(len(loaded[0].frame), 8)
        out = Path(self.temp.name) / 'out.xlsx'
        with patch('modules.display_dashboard.QFileDialog.getSaveFileName', return_value=(str(out), '')):
            self.widget.export()
        self.assertTrue(out.exists())

    def test_plan_round_trip(self):
        path = Path(self.temp.name) / 'plan.json'
        self.widget.brand.setCurrentText('A')
        self.widget.filters = {'插位数': ['6']}
        self.widget.refresh()
        with patch('modules.display_dashboard.QFileDialog.getSaveFileName', return_value=(str(path), '')):
            self.widget.save_plan()
        self.widget.reset_filters()
        with patch('modules.display_dashboard.QFileDialog.getOpenFileName', return_value=(str(path), '')):
            self.widget.load_plan()
        self.assertEqual(self.widget.brand.currentText(), 'A')
        self.assertEqual(self.widget.filters, {'插位数': ['6']})

    def test_resident_attributes_multi_select_reset_remove(self):
        from PyQt6.QtTest import QTest
        frame = fixture().assign(线长=lambda df: df['ASIN'].map({'A1': '1m', 'B1': '2m'}))
        self.widget.set_dataset(dataset(frame))
        self.widget.start.setCurrentText('2026-01')
        self.widget.add_attribute_fields(['插位数'])
        self.widget.add_attribute_fields(['线长', '插位数'])
        self.assertEqual(list(self.widget.attribute_controls), ['插位数', '线长'])
        combo = self.widget.attribute_controls['插位数']
        QTest.mouseClick(combo, Qt.MouseButton.LeftButton)
        self.app.processEvents()
        self.assertTrue(combo.popup.isVisible())
        combo.toggle_all(False)
        combo.options.item(0).setCheckState(Qt.CheckState.Checked)
        QTest.mouseClick(combo.apply_button, Qt.MouseButton.LeftButton)
        self.assertEqual(self.widget.query().filters['插位数'], ['6'])
        self.assertEqual(self.widget.result['kpi']['销售额'], 200)
        combo.showPopup()
        combo.search.setText('8')
        self.assertTrue(combo.options.item(0).isHidden())
        combo.options.item(1).setCheckState(Qt.CheckState.Checked)
        QTest.mouseClick(combo.apply_button, Qt.MouseButton.LeftButton)
        self.assertEqual(combo.selected, [])  # All values means unrestricted, including future values.
        self.assertEqual(self.widget.result['kpi']['销售额'], 1400)
        self.widget.attribute_controls['线长'].commit(['1m'])
        self.assertEqual(self.widget.result['kpi']['销售额'], 200)
        self.widget.reset_filters()
        self.assertEqual(len(self.widget.attribute_controls), 2)
        self.assertEqual(combo.currentText(), '全部')
        self.widget.start.setCurrentText('2026-01')
        combo.commit(['6'])
        self.widget.remove_attribute_field('插位数')
        self.assertNotIn('插位数', self.widget.filters)
        self.assertEqual(self.widget.result['kpi']['销售额'], 1400)
        self.assertEqual(self.config.get('dashboard_filter_fields'), ['线长'])

    def test_add_dialog_can_be_used_repeatedly(self):
        from PyQt6.QtWidgets import QListWidget, QDialog
        def accept_field(dialog):
            fields = dialog.findChild(QListWidget)
            self.assertGreater(fields.count(), 0)
            fields.item(0).setCheckState(Qt.CheckState.Checked)
            return QDialog.DialogCode.Accepted
        self.widget.set_dataset(dataset(fixture().assign(线长='1m')))
        with patch('modules.display_dashboard.QDialog.exec', accept_field):
            self.widget.add_filter()
            self.widget.add_filter()
        self.assertEqual(set(self.widget.attribute_controls), {'插位数', '线长'})

    def test_pinned_layout_survives_reimport_restart_and_plan(self):
        loaded = dataset(fixture().assign(线长='1m'))
        self.widget.set_dataset(loaded)
        self.widget.add_attribute_fields(['插位数', '线长'])
        self.widget.attribute_controls['插位数'].commit(['6'])
        path = Path(self.temp.name) / 'pinned.json'
        with patch('modules.display_dashboard.QFileDialog.getSaveFileName', return_value=(str(path), '')):
            self.widget.save_plan()
        self.widget.remove_attribute_field('线长')
        with patch('modules.display_dashboard.QFileDialog.getOpenFileName', return_value=(str(path), '')):
            self.widget.load_plan()
        self.assertEqual(list(self.widget.attribute_controls), ['插位数', '线长'])
        self.assertEqual(self.widget.attribute_controls['线长'].selected, [])
        self.assertEqual(self.widget.attribute_controls['插位数'].selected, ['6'])
        fresh = ModuleWidget(ConfigManager(self.config.get_path()), None)
        fresh.set_dataset(loaded)
        self.assertEqual(list(fresh.attribute_controls), ['插位数', '线长'])
        fresh.set_dataset(dataset())
        self.assertEqual(list(fresh.attribute_controls), ['插位数'])
        self.assertFalse(fresh.filters)

    def test_currency_hidden_market_isolated_and_scope_in_footer(self):
        frame = fixture()
        self.widget.set_dataset(dataset(pd.concat([frame, frame.assign(币种='EUR', 销售额=1e8)])))
        self.widget.start.setCurrentText('2026-01')
        self.assertTrue(self.widget.currency.isHidden())
        self.widget.site.setCurrentText('US (USD)')
        self.assertEqual(self.widget.query().currency, 'USD')
        self.assertEqual(self.widget.result['kpi']['销售额'], 1400)
        self.widget.site.setCurrentText('US (EUR)')
        self.assertEqual(self.widget.query().currency, 'EUR')
        self.assertEqual(self.widget.result['kpi']['销售额'], 2e8)
        self.app.processEvents()
        self.assertGreater(self.widget.scope.y(), self.widget.status.y())
        self.assertGreater(self.widget.status.y(), self.widget.tabs.y())

    def test_searchable_popups_are_bounded_and_all_action_stays_fixed(self):
        from PyQt6.QtTest import QTest
        from PyQt6.QtWidgets import QPushButton
        self.widget.fill(self.widget.brand, [f'Brand {i}' for i in range(150)], True)
        combo = self.widget.brand
        combo.showPopup()
        self.app.processEvents()
        self.assertLessEqual(combo.popup.height(), 380)
        before = combo.all_button.pos()
        combo.options.scrollToBottom()
        self.app.processEvents()
        self.assertEqual(combo.all_button.pos(), before)
        self.assertTrue(combo.all_button.isVisible())
        combo.search.setText('Brand 149')
        self.assertEqual(sum(not combo.options.item(i).isHidden() for i in range(combo.options.count())), 1)
        combo.hidePopup()
        self.widget.brand.setCurrentText('Brand 149')
        combo.showPopup()
        QTest.mouseClick(combo.all_button, Qt.MouseButton.LeftButton)
        self.assertEqual(combo.currentText(), '全部')
        self.widget.start.showPopup()
        QTest.mouseClick(self.widget.start.all_button, Qt.MouseButton.LeftButton)
        self.assertEqual(self.widget.start.currentText(), '2024-12')
        self.assertEqual(self.widget.end.currentText(), '2026-01')
        visible = [b.text() for b in self.widget.findChildren(QPushButton) if b.isVisible()]
        for old in ['保存分析方案', '加载方案', '调整数据来源', '数据说明', '收起筛选']:
            self.assertNotIn(old, visible)
        self.assertIn('重新加载数据源', visible)

    def test_attribute_all_toggle_does_not_scroll_and_does_not_apply_empty(self):
        from PyQt6.QtTest import QTest
        self.widget.add_attribute_fields(['插位数'])
        combo = self.widget.attribute_controls['插位数']
        combo.showPopup()
        self.app.processEvents()
        self.assertTrue(combo.select_all.isChecked())
        QTest.mouseClick(combo.select_all, Qt.MouseButton.LeftButton, pos=QPoint(8, combo.select_all.height() // 2))
        self.assertFalse(combo.apply_button.isEnabled())
        QTest.mouseClick(combo.select_all, Qt.MouseButton.LeftButton, pos=QPoint(8, combo.select_all.height() // 2))
        self.assertTrue(combo.apply_button.isEnabled())
        pos = combo.select_all.pos()
        combo.options.scrollToBottom()
        self.assertEqual(combo.select_all.pos(), pos)
        QTest.mouseClick(combo.apply_button, Qt.MouseButton.LeftButton)
        self.assertNotIn('插位数', self.widget.filters)

    def test_reload_reads_new_file_preserves_filters_and_tracks_latest(self):
        from core.dashboard_import import load_workbook as read_dashboard
        from PyQt6.QtCore import QThreadPool
        path = Path(self.temp.name) / 'source.xlsx'
        frame = fixture()
        frame.to_excel(path, index=False)
        self.widget.set_dataset(read_dashboard(str(path))[0])
        self.widget.start.setCurrentText('2025-12')
        self.widget.brand.setCurrentText('A')
        self.widget.add_attribute_fields(['插位数'])
        self.widget.attribute_controls['插位数'].commit(['6'])
        newer = frame[frame['月份'] == '2026-01'].assign(月份='2026-02', 销售额=lambda df: df['销售额'] * 3)
        pd.concat([frame, newer]).to_excel(path, index=False)
        self.widget.reload_source()
        self.assertFalse(self.widget.reload_btn.isEnabled())
        QThreadPool.globalInstance().waitForDone(10000)
        self.app.processEvents()
        self.assertEqual(self.widget.start.currentText(), '2025-12')
        self.assertEqual(self.widget.end.currentText(), '2026-02')
        self.assertEqual(self.widget.brand.currentText(), 'A')
        self.assertEqual(self.widget.filters, {'插位数': ['6']})
        self.assertEqual(self.widget.result['kpi']['销售额'], 900)
        self.assertIn('已重新加载', self.widget.status.text())
        previous = self.widget.dataset
        path.write_bytes(b'invalid workbook')
        with patch('modules.display_dashboard.QMessageBox.warning') as warning:
            self.widget.reload_source()
            QThreadPool.globalInstance().waitForDone(10000)
            self.app.processEvents()
            warning.assert_called_once()
        self.assertIs(self.widget.dataset, previous)
        self.assertEqual(self.widget.filters, {'插位数': ['6']})
        self.assertTrue(self.widget.reload_btn.isEnabled())

    def test_reload_prunes_missing_values_with_notice(self):
        self.widget.add_attribute_fields(['插位数'])
        self.widget.attribute_controls['插位数'].commit(['6', '8'])
        self.widget.brand.setCurrentText('B')
        self.widget._reload_state = {'query': self.widget.query(), 'latest': '2026-01',
                                     'fields': ['插位数'], 'search': ''}
        self.widget.set_dataset(dataset(fixture().query("ASIN == 'A1'")))
        self.assertEqual(self.widget.filters, {'插位数': ['6']})
        self.assertEqual(self.widget.brand.currentText(), '全部')
        self.assertIn('失效筛选', self.widget.status.text())

    def test_fast_detail_search_keeps_numeric_sort_and_refresh_consistent(self):
        self.assertEqual(self.widget.proxy.sortColumn(), -1)
        self.widget.search.setText('1,200')
        self.assertEqual(self.widget.visible_frame()['ASIN'].tolist(), ['B1'])
        self.widget.start.setCurrentText('2025-12')
        self.assertEqual(self.widget.visible_frame()['ASIN'].tolist(), ['B1'])
        self.widget.search.clear()
        col = list(self.widget.model.frame).index('销售额')
        self.widget.proxy.sort(col, Qt.SortOrder.DescendingOrder)
        self.assertEqual(self.widget.visible_frame()['销售额'].tolist(), [1200, 600, 200, 100])

    def test_matrix_signal_carries_current_asins(self):
        paths = []
        self.widget.matrix_requested.connect(paths.append)
        self.widget.search.setText('A1')
        self.widget.send_matrix()
        frame = pd.read_excel(paths[0])
        self.assertEqual(frame['ASIN'].tolist(), ['A1'])

    def test_navigation_integration_and_theme_render(self):
        from core.main_window import MainWindow
        from core.theme import THEME
        self.config.update({'auto_check_app_update': False, 'auto_check_update': False})
        with patch('core.main_window.ConfigManager', return_value=self.config):
            window = MainWindow()
        dashboard = next(w for info, w in window._modules if info['name'] == '展示看板')
        dashboard.set_dataset(dataset())
        dashboard.send_matrix()
        index = window.stack.currentIndex()
        self.assertEqual(window._modules[index][0]['name'], '产品矩阵')
        products = window._modules[index][1]
        self.assertIs(products.matrix_tabs.currentWidget(), products.free_matrix)
        self.assertEqual(len(products.free_matrix.frame), 2)
        for theme in ['dark', 'light']:
            THEME.set_theme(theme)
            self.app.processEvents()
            self.assertFalse(self.widget.grab().isNull())


if __name__ == '__main__':
    unittest.main()
