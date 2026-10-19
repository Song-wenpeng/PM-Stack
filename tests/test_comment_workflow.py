"""User workflows and durable resume behavior, without paid API calls."""
from tests import support

import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from core.config_manager import ConfigManager
from core.runner import ScriptRunner
from core.reviews.analysis_jobs import AnalysisJobs, inspect_workbook
from core.reviews.analysis_dialog import template_fingerprint
from modules.comment_analysis import ModuleWidget


def source_book(path, header="内容"):
    book = Workbook()
    book.active.title = "5 star"
    book.active.append([header, "评论链接"])
    book.active.append(["Works well", "https://example.test/review/1"])
    book.active.append([None, None])
    sheet = book.create_sheet("1 star")
    sheet.append([header, "评论链接"])
    sheet.append(["Broke quickly", "https://example.test/review/2"])
    book.create_sheet("说明").append(["这是说明页"])
    book.save(path)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.app = support.qt_app()
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.manager = ConfigManager(str(self.root / "config.json"))
        self.runner = ScriptRunner()
        self.widget = ModuleWidget(self.manager, self.runner)
        self.source = self.root / "reviews.xlsx"
        source_book(self.source)
        self.env_patch = mock.patch.object(self.manager, "get_env_dict", return_value={"API_KEY": "test-only"})
        self.env_patch.start()
        self.launch = mock.patch.object(self.runner, "run_script", return_value=True).start()
        self.addCleanup(mock.patch.stopall)

    def tearDown(self):
        self.wait_for(lambda: self.widget._worker is None)
        self.widget._queue.clear()
        self.widget._active = None
        self.widget._disconnect_runner()
        self.app.processEvents()
        support.cleanup_qt_widgets()
        self.temp.cleanup()

    def wait_for(self, predicate):
        deadline = time.monotonic() + 5
        while not predicate() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.005)
        self.app.processEvents()
        self.assertTrue(predicate(), "UI worker did not finish")

    def start(self, paths=None):
        self.widget.product_combo.setCurrentIndex(self.widget.product_combo.findData("general"))
        self.widget._add_inputs([str(p) for p in paths or [self.source]])
        self.wait_for(lambda: self.widget._worker is None)
        self.widget._start_analysis()
        self.wait_for(lambda: self.widget._worker is None and self.widget._active is not None)
        return self.widget._active

    def test_start_is_default_and_batch_per_file_categories(self):
        self.assertEqual([self.widget.tabs.tabText(i) for i in range(self.widget.tabs.count())], ["开始分析", "评论库", "分析记录"])
        self.assertIs(self.widget.tabs.currentWidget(), self.widget.start_tab)
        other = self.root / "other.xlsx"
        source_book(other)
        self.widget._add_inputs([str(self.source), str(other)])
        self.wait_for(lambda: self.widget._worker is None)
        self.widget.product_combo.setCurrentIndex(self.widget.product_combo.findData("general"))
        combo = self.widget.input_table.cellWidget(1, 2)
        combo.setCurrentIndex(combo.findData("adapter"))
        self.widget._start_analysis()
        self.wait_for(lambda: self.widget._active is not None)
        self.assertEqual([r["product"] for r in self.widget.jobs.records], ["general", "adapter"])
        self.assertEqual(len(self.widget._queue), 1)
        self.assertEqual(self.widget._active["count"], 2)
        env = self.launch.call_args.kwargs["env_extra"]
        self.assertTrue(self.launch.call_args.args[0].endswith("comments_step_1_fast.py"))
        self.assertEqual(env["STEP1_CONCURRENCY"], "8")
        self.assertEqual(env["STEP1_RPM_LIMIT"], "900")
        self.assertEqual(env["STEP1_TPM_LIMIT"], "90000")
        self.assertEqual(env["STEP1_MAX_FAILURE_RATE"], "0")

    def test_partial_failure_blocks_summary_and_reuses_same_checkpoint(self):
        job = self.start()
        checkpoint = Path(job["step1"] + ".checkpoint.json")
        checkpoint.write_text('{"saved":true}')
        self.runner.finished_signal.emit(2)
        self.app.processEvents()
        self.assertEqual(job["status"], "failed")
        self.assertEqual(self.launch.call_count, 1)
        self.widget._resume_selected()
        self.assertEqual(self.widget._active["id"], job["id"])
        self.assertEqual(self.launch.call_args.kwargs["env_extra"]["STEP1_OUTPUT_FILE"], job["step1"])
        self.assertTrue(checkpoint.exists())
        self.assertEqual(len(self.widget.jobs.records), 1)

    def test_summary_failure_retry_does_not_repeat_extraction(self):
        job = self.start()
        source_book(Path(job["step1"]))
        self.runner.finished_signal.emit(0)
        self.wait_for(lambda: self.launch.call_count == 2)
        self.assertEqual(job["stage"], "summary")
        self.runner.finished_signal.emit(1)
        self.app.processEvents()
        self.widget._resume_selected()
        self.assertTrue(self.launch.call_args.args[0].endswith("comments_step_2.py"))
        self.assertEqual(self.launch.call_count, 3)

    def test_stop_preserves_queue_and_restart_recovers_tasks(self):
        other = self.root / "other.xlsx"
        source_book(other)
        job = self.start([self.source, other])
        self.widget._stop()
        self.assertTrue(self.widget._stop_file.exists())
        self.assertEqual(self.widget.jobs.records[1]["status"], "stopped")
        self.runner.finished_signal.emit(0)
        self.app.processEvents()
        self.assertEqual(job["stage"], "summary")
        self.assertEqual(job["status"], "stopped")
        self.assertEqual(self.launch.call_count, 1)
        restored = AnalysisJobs(self.widget.jobs.root)
        self.assertEqual(len(restored.records), 2)
        self.assertEqual(restored.records[0]["stage"], "summary")

    def test_success_keeps_report_for_each_file_and_persists_history(self):
        other = self.root / "other.xlsx"
        source_book(other)
        self.widget.mode_combo.setCurrentIndex(1)
        first = self.start([self.source, other])
        source_book(Path(first["report"]))
        self.runner.finished_signal.emit(0)
        self.wait_for(lambda: self.widget._active is not None and self.widget._active["id"] != first["id"])
        second = self.widget._active
        source_book(Path(second["report"]))
        self.runner.finished_signal.emit(0)
        self.wait_for(lambda: self.widget._active is None)
        self.assertTrue(all(job["status"] == "done" for job in self.widget.jobs.records))
        self.widget._refresh_history(first["id"])
        self.assertEqual(self.widget._report_path, first["report"])
        self.assertNotEqual(first["report"], second["report"])
        self.assertEqual(len(AnalysisJobs(self.widget.jobs.root).records), 2)

    def test_library_confirmed_scope_starts_without_save_dialog(self):
        key = "general"
        template = self.widget.templates[key]
        payload = {"path": str(self.source), "product_key": key, "fingerprint": template_fingerprint(template), "asins": ["B000TEST01"]}
        with mock.patch("core.reviews.workflow.QFileDialog.getSaveFileName") as save:
            self.widget._use_collected_review_file(payload)
            self.wait_for(lambda: self.widget._active is not None)
        save.assert_not_called()
        self.assertEqual(self.widget._active["product"], key)
        self.assertEqual(self.widget._active["title"], "B000TEST01")
        self.assertIs(self.widget.tabs.currentWidget(), self.widget.result_tab)

    def test_deleted_or_changed_template_blocks_library_start(self):
        for key, fingerprint in (("deleted", ""), ("general", "old-fingerprint")):
            self.widget._use_collected_review_file({"path": str(self.source), "product_key": key, "fingerprint": fingerprint})
            self.assertIsNone(self.widget._worker)
            self.assertIn("无法开始", self.widget.start_hint.text())
        self.launch.assert_not_called()

    def test_invalid_data_and_empty_range_never_launch(self):
        bad = self.root / "bad.xlsx"
        Workbook().save(bad)
        self.widget._add_inputs([str(bad)])
        self.wait_for(lambda: self.widget._worker is None)
        self.assertEqual(self.widget.inputs, [])
        self.assertIn("没有找到", self.widget.start_hint.text())
        self.widget._add_inputs([str(self.source)])
        self.wait_for(lambda: self.widget._worker is None)
        self.widget.inputs[0]["selected"] = []
        self.widget._start_analysis()
        self.assertIn("至少选择", self.widget.start_hint.text())
        self.launch.assert_not_called()

    def test_summary_only_validates_labels_and_uses_snapshot(self):
        template = self.widget.templates["general"]
        with self.assertRaisesRegex(ValueError, "标签列"):
            self.widget._validate_labels(str(self.source), ["5 star"], template)
        source = self.root / "labels.xlsx"
        book = Workbook()
        book.active.title = "reviews"
        fields = template["stats_fields"]
        book.active.append(["内容", *fields, "_处理状态"])
        book.active.append(["review", *["[]" for _ in fields], "AI提取失败"])
        book.save(source)
        with self.assertRaisesRegex(ValueError, "补跑"):
            self.widget._validate_labels(str(source), ["reviews"], template)
        book.active.cell(2, len(fields) + 2, "成功")
        book.save(source)
        self.widget.mode_combo.setCurrentIndex(2)
        job = self.start([source])
        self.assertEqual(job["step1"], job["input"])
        self.assertTrue(self.launch.call_args.args[0].endswith("comments_step_2.py"))


class SnapshotTests(unittest.TestCase):
    def test_report_names_are_readable_and_reserved_before_batch_is_saved(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "B09N7T1TN5-US-Reviews-261009-166558.xlsx"
            source_book(path)
            jobs = AnalysisJobs(root / "jobs")
            with mock.patch("core.reviews.analysis_jobs.datetime") as clock:
                clock.now.return_value.strftime.return_value = "20261009"
                first = jobs.prepare(path, ["1 star"], "general", {})
                self.assertEqual(Path(first["report"]).name, "B09N7T1TN5-US_评论分析_20261009.xlsx")
                second = jobs.prepare(path, ["1 star"], "general", {})
                self.assertEqual(Path(second["report"]).name, "B09N7T1TN5-US_评论分析_20261009_2.xlsx")
                Path(second["report"]).touch()
                restored = AnalysisJobs(root / "jobs")
                restored.records.append(first)
                third = restored.prepare(path, ["1 star"], "general", {})
                self.assertEqual(Path(third["report"]).name, "B09N7T1TN5-US_评论分析_20261009_3.xlsx")
                generic = root / "用户自定义评论.xlsx"
                source_book(generic)
                labels = restored.prepare(generic, ["1 star"], "general", {}, mode="extract")
                self.assertEqual(Path(labels["report"]).name, "用户自定义评论_评论标签_20261009.xlsx")

    def test_selected_sheets_aliases_ratings_and_originals_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "reviews.xlsx"
            source_book(path, "评论内容")
            original = path.read_bytes()
            sheets = inspect_workbook(path)
            self.assertEqual(sheets, [{"name": "5 star", "count": 1}, {"name": "1 star", "count": 1}])
            jobs = AnalysisJobs(root / "jobs")
            job = jobs.prepare(path, ["1 star"], "general", {"display_name": "通用"})
            self.assertEqual(job["count"], 1)
            book = load_workbook(job["input"])
            rows = list(book.active.values)
            values = dict(zip(rows[0], rows[1]))
            book.close()
            self.assertEqual(values["内容"], "Broke quickly")
            self.assertEqual(values["星级"], 1)
            self.assertEqual(values["来源行号"], 2)
            self.assertEqual(values["来源工作表"], "1 star")
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(json.loads(Path(job["config"]).read_text(encoding="utf-8")), {"general": {"display_name": "通用"}})

    def test_running_task_is_resumable_after_restart_and_bad_history_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            jobs = AnalysisJobs(directory)
            jobs.records = [{"id": "task", "status": "running", "stage": "summary"}]
            jobs.save()
            self.assertEqual(AnalysisJobs(directory).records[0]["status"], "stopped")
            jobs.path.write_text("invalid", encoding="utf-8")
            broken = AnalysisJobs(directory)
            with self.assertRaises(OSError):
                broken.save()
            self.assertEqual(jobs.path.read_text(), "invalid")


if __name__ == "__main__":
    unittest.main()
