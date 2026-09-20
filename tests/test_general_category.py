from tests import support
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
from openpyxl import load_workbook
from core.config_manager import ConfigManager
from core.runner import ScriptRunner
from core.reviews.detail_export import export_review_details
from core.reviews.excel_import import extract_labels
from core.reviews.metadata import latest_analysis
from core.reviews.store import ReviewStore
from modules.comment_analysis import ModuleWidget
from modules.comment_analysis_test import ModuleWidget as FastModule
from scripts.comments_step_1_fast import build_prompt

class GeneralCategoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.manager = ConfigManager(str(self.root / 'config.json'))
        self.configs = json.loads(Path(self.manager.resolve_config('product_configs.json')).read_text(encoding='utf-8'))
        self.config = self.configs['general']

    def test_every_analysis_entry_offers_general(self):
        support.qt_app()
        self.addCleanup(support.cleanup_qt_widgets)
        widget = ModuleWidget(self.manager, ScriptRunner())
        fast = FastModule(self.manager, ScriptRunner())
        for combo in (widget.s1_product, widget.s2_product, widget.s12_product,
                      widget.batch_product, widget.review_collection.excel_options.category,
                      fast.product_combo):
            index = combo.findData('general')
            self.assertGreaterEqual(index, 0)
            combo.setCurrentIndex(index)
            self.assertIn('通用', combo.currentText())
            self.assertEqual(combo.currentData(), 'general')

    def test_prompt_fields_and_stats_are_consistent(self):
        columns = [f['excel_col'] for f in self.config['fields']]
        self.assertEqual(len(columns), len(set(columns)))
        self.assertEqual(set(self.config['stats_fields']), set(columns) - {'翻译后评论'})
        prompt = build_prompt(self.config, 'Bought for camping. Easy to wash, but too small.')
        for field in self.config['fields']:
            self.assertIn(field['key'], prompt)
        self.assertIn('不把所有缺点自动转换成用户建议', prompt)
        self.assertTrue({'adapter', 'power strip', 'endoscope', 'vacuum auto switch',
                         'saw vacuum scenario'}.issubset(self.configs))

    def test_general_results_persist_cache_and_export(self):
        store = ReviewStore(str(self.root / 'reviews.db'))
        self.addCleanup(store.close)
        row = {'review_id': 'GENERAL1', 'content': 'Bought for camping. Easy to wash, but too small.', 'rating': 3}
        store.add_reviews('B0DDG2KQQ9', 'amazon.com', [row])
        data = {'asin': 'B0DDG2KQQ9', 'marketplace': 'amazon.com', 'rows': [row]}
        labels = {f['key']: [] if f['type'] == 'list' else '' for f in self.config['fields']}
        labels.update(translated_comment='为了露营购买。易清洗，但太小。',
                      usage_scenarios=['露营'], purchase_motivations=['露营需要'],
                      positive_focus=['易清洗'], negative_focus=['尺寸太小'])
        gateway = Mock(return_value=json.dumps(labels, ensure_ascii=False))
        self.assertEqual(extract_labels(store, data, 'general', self.config, {}, gateway=gateway)['success'], 1)
        self.assertEqual(extract_labels(store, data, 'general', self.config, {}, gateway=gateway)['cached'], 1)
        gateway.assert_called_once()
        rows = store.query_reviews()
        self.assertEqual(latest_analysis(rows[0])['category'], 'general')
        self.assertEqual(latest_analysis(rows[0])['labels'], labels)
        self.assertGreater(store.outbox_stats()['pending'], 0)
        target = self.root / 'general.xlsx'
        export_review_details(rows, str(target))
        book = load_workbook(target)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ['评论明细'])
        values = dict(zip((c.value for c in book.active[1]), (c.value for c in book.active[2])))
        self.assertEqual(values['AI · 正向关注点'], '易清洗')
        self.assertEqual(values['AI · 负向关注点'], '尺寸太小')
