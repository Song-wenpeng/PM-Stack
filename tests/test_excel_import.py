from tests import support

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
from types import SimpleNamespace
from openpyxl import Workbook
from core.reviews.excel_import import read_import, import_reviews, extract_labels
from core.reviews.metadata import latest_analysis, KEY, merge_payload
from core.reviews.store import ReviewStore
from core.reviews.repository import ReviewRepository
from core.reviews.cloud import SupabaseStore

ASIN="B0DDG2KQQ9"
CONFIG={"product_desc":"Power strips","tasks":["提取优缺点"],
        "fields":[{"key":"positive","type":"list","excel_col":"优点"},
                  {"key":"negative","type":"list","excel_col":"缺点"}]}

class ExcelImportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.store=ReviewStore(str(self.root/"reviews.db"))
        self.path=self.root/"reviews.xlsx"
        self.write()

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def write(self, second=False, wrong_site=False):
        w=Workbook(); s=w.active; s.title="Reviews"
        s.append(["型号","VP评论","标题","内容","星级","链接","评论时间"])
        host="amazon.de" if wrong_site else "amazon.com"
        s.append(["6 Outlet","Y","Good","Works well",5,
            f"https://www.{host}/gp/customer-reviews/RTEST0001/?ASIN=B0CM6L5WTY","2026-09-05"])
        if second:
            s.append(["6 Outlet","Y","Good","Works well",5,
                f"https://www.{host}/gp/customer-reviews/RTEST0001/?ASIN=B0CM6L5WTY","2026-09-05"])
        w.create_sheet("Notes").append(["SellerSprite notes"])
        w.save(self.path); w.close()

    def data(self):
        return read_import(self.path,ASIN,"amazon.com")

    def test_mapping_variant_notes_and_duplicate_preview(self):
        before=self.path.read_bytes()
        self.write(second=True)
        before=self.path.read_bytes()
        data=self.data()
        self.assertEqual((data["row_count"],data["unique_count"],data["file_duplicates"]),(2,1,1))
        self.assertEqual(data["skipped_sheets"],["Notes"])
        row=data["rows"][0]
        self.assertEqual(row["review_id"],"RTEST0001")
        self.assertTrue(row["verified_purchase"])
        source=list(row["raw_payload"][KEY]["imports"].values())[0]
        self.assertEqual(source["target_asin"],ASIN)
        self.assertEqual(source["linked_asin"],"B0CM6L5WTY")
        self.assertEqual(self.path.read_bytes(),before)

    def test_wrong_site_rejected_without_database_write(self):
        self.write(wrong_site=True)
        with self.assertRaises(ValueError): self.data()
        self.assertEqual(self.store.total_reviews(),0)

    def test_import_reimport_labels_and_queue_keep_annotations(self):
        data=self.data()
        self.assertEqual(import_reviews(self.store,data),(1,0))
        gateway=Mock(return_value='{"positive":["好用"],"negative":[]}')
        first=extract_labels(self.store,data,"power strip",CONFIG,{},gateway=gateway)
        self.assertEqual(first["success"],1)
        second=extract_labels(self.store,data,"power strip",CONFIG,{},gateway=gateway)
        self.assertEqual(second["cached"],1)
        gateway.assert_called_once()
        self.assertEqual(import_reviews(self.store,data),(0,1))
        row=self.store.query_reviews()[0]
        self.assertEqual(latest_analysis(row)["labels"]["positive"],["好用"])
        self.assertGreaterEqual(self.store.outbox_stats()["pending"],2)
        fake=Mock()
        repo=ReviewRepository(self.store); repo.cloud=fake
        result=repo.sync_pending()
        self.assertFalse(result["error"])
        for call in fake.ingest_batch.call_args_list:
            reviews=call.args[0]["reviews"]
            self.assertTrue(reviews[0]["raw_payload"][KEY]["analyses"])

    def test_failed_or_cancelled_model_never_marks_success(self):
        data=self.data(); import_reviews(self.store,data)
        result=extract_labels(self.store,data,"power strip",CONFIG,{},
                             gateway=lambda _: '{"positive":"wrong"}')
        self.assertEqual(result["failed"],1)
        self.assertIsNone(latest_analysis(self.store.query_reviews()[0]))
        gateway=Mock()
        result=extract_labels(self.store,data,"power strip",CONFIG,{},
                             gateway=gateway,stop=lambda:True)
        self.assertTrue(result["cancelled"]); gateway.assert_not_called()

    def test_template_change_reextracts_and_preserves_versions(self):
        data=self.data(); import_reviews(self.store,data)
        gateway=lambda _: '{"positive":[],"negative":[]}'
        extract_labels(self.store,data,"power strip",CONFIG,{},gateway=gateway)
        changed={**CONFIG,"tasks":["另一版任务"]}
        result=extract_labels(self.store,data,"power strip",changed,{},gateway=gateway)
        self.assertEqual(result["success"],1)
        raw=json.loads(self.store.query_reviews()[0]["raw_json"])
        self.assertEqual(len(raw[KEY]["analyses"]),2)

    def test_cloud_ingest_preserves_remote_metadata(self):
        cloud=object.__new__(SupabaseStore)
        cloud.workspace_id="workspace"
        cloud.review_metadata=lambda ids:[{"marketplace":"amazon.com","review_id":"RTEST0001",
            "raw_payload":{KEY:{"analyses":{"old":{"status":"success"}}}}}]
        cloud.client=Mock()
        cloud.client.rpc.return_value.execute.return_value=SimpleNamespace(data={"reviews":1})
        data=self.data()
        cloud.ingest_batch({"product":{"asin":ASIN,"marketplace":"amazon.com"},
                            "reviews":data["rows"]},"receipt")
        sent=cloud.client.rpc.call_args.args[1]["p_reviews"][0]["raw_payload"]
        self.assertIn("old",sent[KEY]["analyses"])
        self.assertTrue(sent[KEY]["imports"])

    def test_same_id_conflicting_body_rejected(self):
        w=Workbook();s=w.active;s.append(["内容","星级","评论ID"])
        s.append(["first",5,"RTEST0001"]);s.append(["different",5,"RTEST0001"])
        w.save(self.path);w.close()
        with self.assertRaises(ValueError): self.data()

    def test_three_failures_pause_remaining_reviews(self):
        data=self.data()
        data["rows"]=[{**data["rows"][0],"review_id":"RPAUSE00"+str(i)} for i in range(5)]
        import_reviews(self.store,data)
        gateway=Mock(side_effect=ValueError("invalid JSON"))
        result=extract_labels(self.store,data,"power strip",CONFIG,{},gateway=gateway)
        self.assertEqual((result["failed"],result["pending"]),(3,2))
        self.assertEqual(gateway.call_count,3)

    def test_explicit_different_product_rejected(self):
        w=Workbook();s=w.active;s.append(["ASIN","内容","星级"])
        s.append(["B0CM6L5WTY","Body",5])
        w.save(self.path);w.close()
        with self.assertRaises(ValueError): self.data()
