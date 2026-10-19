import tempfile
from copy import copy
import importlib.util
import os
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from core.field_workbook import FieldWorkbookWriter, source_fingerprint, LOG_FIELDS, LOG_SHEET


class FieldWorkbookTests(unittest.TestCase):
    def test_scrolled_panes_preserve_freeze_boundary(self):
        from openpyxl.worksheet.views import Pane
        original = self.path.read_bytes()
        for state, x, y, top_left, expected_x in [
            ('frozen', None, 1, 'A189', None),
            ('frozen', 1, 1, 'C189', 3),
            ('frozenSplit', 1, 2, 'C189', 3),
            ('split', 2400, 1200, 'C189', 2400),
        ]:
            with self.subTest(state=state, x=x):
                self.path.write_bytes(original)
                wb = load_workbook(self.path)
                # Reset to original layout for each case.
                ws = wb.active
                ws.sheet_view.pane = Pane(state=state, xSplit=x, ySplit=y,
                                          topLeftCell=top_left, activePane='bottomRight')
                wb.save(self.path)
                wb.close()
                writer = FieldWorkbookWriter(self.path, None, ['包装数', 'USB接口'])
                df = self.results()
                writer.save(df)
                writer.save(df)
                wb = load_workbook(self.path)
                pane = wb.active.sheet_view.pane
                self.assertEqual(pane.ySplit, y)
                self.assertEqual(pane.xSplit, expected_x)
                self.assertEqual(pane.topLeftCell, 'A189' if top_left == 'A189' else 'E189')
                self.assertEqual(pane.state, state)
                self.assertEqual(pane.activePane, 'bottomRight')
                wb.close()

    def test_carriage_returns_survive_repeated_saves(self):
        from zipfile import ZipFile
        wb = load_workbook(self.path)
        wb.active['B2'] = 'line one\r\nline two'
        wb.save(self.path)
        wb.close()
        # Construct the XML-escaped CR used in real WPS source files.
        with ZipFile(self.path) as archive:
            entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        with ZipFile(self.path, 'w') as archive:
            for info, data in entries:
                from openpyxl import LXML
                if os.name == 'nt' and not LXML and info.filename.startswith('xl/worksheets/'):
                    data = data.replace(b'\r\n', b'\n')
                archive.writestr(info, data.replace(b'\r', b'&#13;') if info.filename.endswith('.xml') else data)
        writer = FieldWorkbookWriter(self.path, None, ['包装数', 'USB接口'])
        df = self.results()
        writer.save(df)
        writer.save(df)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['D2'].value, 'line one\r\nline two')

    def test_images_follow_columns_and_survive_repeated_saves(self):
        from PIL import Image as PILImage
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
        image_path = Path(self.temp.name) / 'test.png'
        PILImage.new('RGB', (16, 16), 'red').save(image_path)
        wb = load_workbook(self.path)
        pic = Image(image_path)
        pic.anchor = TwoCellAnchor(_from=AnchorMarker(col=1, row=1, colOff=123),
                                   to=AnchorMarker(col=1, row=1, colOff=9000))
        wb.active.add_image(pic)
        wb['汇总'].add_image(Image(image_path), 'B2')
        wb.save(self.path)
        wb.close()
        writer = FieldWorkbookWriter(self.path, None, ['包装数', 'USB接口'])
        df = self.results()
        writer.save(df)
        writer.save(df)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(len(wb.active._images), 1)
        anchor = wb.active._images[0].anchor
        self.assertEqual((anchor._from.col, anchor.to.col, anchor._from.colOff), (3, 3, 123))
        self.assertEqual(wb['汇总']._images[0].anchor._from.col, 1)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'source.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '商品数据'
        ws.append(['ASIN', '标题', '销量', '金额'])
        ws.append(['B001', '电源', 3, '=C2*10'])
        ws.append(['B001', '同ASIN第二行', 4, '=C3*10'])
        for row in ws:
            for c in row:
                c.font = Font(name='微软雅黑', size=11)
                c.alignment = Alignment(horizontal='left')
        ws.auto_filter.ref = 'A1:D3'
        ws.freeze_panes = 'B2'
        wb.create_sheet('汇总').append(["='商品数据'!D2", '保留'])
        wb.save(self.path)
        wb.close()

    def results(self):
        df = pd.read_excel(self.path)
        df['包装数'] = ['1', '2']
        df['USB接口'] = ['3A1C', '2A2C']
        df['文本提取字段'] = '包装数；USB接口'
        df['_处理状态'] = ['', '已停止']
        return df

    def test_writeback_preserves_formulas_sheets_styles_and_log_rows(self):
        original = self.path.read_bytes()
        writer = FieldWorkbookWriter(self.path, None, ['包装数', 'USB接口'])
        writer.save(self.results())
        self.assertEqual(writer.backup.read_bytes(), original)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        ws = wb.worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ['ASIN', '包装数', 'USB接口', '标题', '销量', '金额'])
        self.assertEqual(ws['F2'].value, '=E2*10')
        self.assertEqual(wb['汇总']['A1'].value, "='商品数据'!F2")
        self.assertEqual(ws['B2'].value, 1)
        self.assertEqual(copy(ws['B2'].font), copy(ws['D2'].font))
        self.assertTrue(ws['B2'].alignment.wrap_text)
        self.assertEqual(ws.freeze_panes, 'D2')
        self.assertEqual(ws.auto_filter.ref, 'A1:F3')
        log = wb[LOG_SHEET]
        self.assertEqual(log.max_row, 3)
        self.assertEqual(log['B3'].value, 3)
        self.assertEqual(log['C3'].value, 'B001')
        self.assertEqual(log['I3'].value, '已停止')

    def test_rerun_reuses_columns_logs_and_fingerprint(self):
        fields = ['包装数', 'USB接口']
        before = source_fingerprint(self.path, fields)
        writer = FieldWorkbookWriter(self.path, None, fields)
        writer.save(self.results())
        self.assertEqual(before, source_fingerprint(self.path, fields))
        frame = pd.read_excel(self.path)
        writer.restore_logs(frame)
        self.assertEqual(frame.at[1, '_处理状态'], '已停止')
        frame['包装数'] = '99'
        second = FieldWorkbookWriter(self.path, None, fields)
        second.save(frame)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.worksheets[0]['B2'].value, 1)
        self.assertEqual(wb.worksheets[0].max_column, 6)
        self.assertEqual(wb[LOG_SHEET].max_row, 3)

    def test_locked_output_retains_results_and_source(self):
        writer = FieldWorkbookWriter(self.path, None, ['包装数', 'USB接口'])
        original = self.path.read_bytes()
        with patch('core.field_workbook.os.replace', side_effect=PermissionError('locked')):
            with self.assertRaisesRegex(RuntimeError, '关闭 Excel/WPS'):
                writer.save(self.results())
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(len(list(self.path.parent.glob('*.pending-*.xlsx'))), 1)

    def test_external_edit_blocks_write(self):
        writer = FieldWorkbookWriter(self.path, None, ['包装数'])
        results = self.results()
        wb = load_workbook(self.path)
        wb.active['B2'] = '人工改动'
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(RuntimeError, '其他程序修改'):
            writer.save(results)

    def test_legacy_columns_move_to_log_and_new_location(self):
        wb = load_workbook(self.path)
        ws = wb.active
        ws['E1'], ws['E2'] = '包装数', '3'
        ws['F1'], ws['F2'] = '文本提取字段', '包装数'
        wb.save(self.path)
        wb.close()
        df = pd.read_excel(self.path)
        FieldWorkbookWriter(self.path, None, ['包装数']).save(df)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['B2'].value, 3)
        self.assertNotIn('文本提取字段', [c.value for c in wb.active[1]])
        self.assertEqual(wb[LOG_SHEET]['D2'].value, '包装数')

    def test_alternate_output_preserves_source(self):
        original = self.path.read_bytes()
        output = self.path.parent / 'copy.xlsx'
        FieldWorkbookWriter(self.path, output, ['包装数']).save(self.results())
        self.assertEqual(self.path.read_bytes(), original)
        self.assertTrue(output.exists())

    def test_unsafe_range_blocked_before_model_calls(self):
        wb = load_workbook(self.path)
        wb['汇总']['A1'] = "=SUM('商品数据'!A2:D2)"
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, '公式范围'):
            FieldWorkbookWriter(self.path, None, ['包装数'])

    def test_source_changes_invalidate_resume(self):
        first = source_fingerprint(self.path, ['包装数'])
        wb = load_workbook(self.path)
        wb.active['B2'] = 'changed'
        wb.save(self.path)
        wb.close()
        self.assertNotEqual(first, source_fingerprint(self.path, ['包装数']))

    def load_script(self, name):
        location = Path(__file__).resolve().parents[1] / 'scripts' / (name + '.py')
        spec = importlib.util.spec_from_file_location(name, location)
        module = importlib.util.module_from_spec(spec)
        with patch.dict(os.environ, {'API_KEY': 'test-only', 'CURRENT_PRODUCT': '', 'ENABLE_FETCH': '0'}):
            spec.loader.exec_module(module)
        return module

    def extraction_source(self):
        wb = load_workbook(self.path)
        for i, field in enumerate(['商品标题', '详细参数', 'SKU', '五点描述'], 5):
            wb.active.cell(1, i, field)
            wb.active.cell(2, i, 'source')
            wb.active.cell(3, i, 'source')
        wb.save(self.path)
        wb.close()

    def test_fast_main_stops_and_resumes_after_periodic_source_save(self):
        self.extraction_source()
        fast = self.load_script('exstract_from_description_fast')
        fast.INPUT_FILE, fast.OUTPUT_FILE = str(self.path), ''
        fast.SAVE_EVERY, fast.CONCURRENCY = 1, 1
        fast.ENABLE_VISION_DEFAULT = False
        config = {'name': 'test', 'target_fields': ['包装数'], 'field_config': {'包装数': '数量'}}
        calls = []
        def process(index, *args):
            calls.append(index)
            return index, {'_extracted': {'包装数': '2'}, '_text_fields': ['包装数']}, ''
        with patch.object(fast, 'load_product_config', return_value=('test', config)), patch.object(fast, 'process_single_row', side_effect=process), patch.object(fast, 'check_stop_file', side_effect=[False, True]):
            with self.assertRaises(SystemExit):
                fast.main()
        self.assertEqual(calls, [0])
        with patch.object(fast, 'load_product_config', return_value=('test', config)), patch.object(fast, 'process_single_row', side_effect=process), patch.object(fast, 'check_stop_file', return_value=False):
            fast.main()
        self.assertEqual(calls, [0, 1])
        self.assertFalse(Path(str(self.path) + '.checkpoint.json').exists())
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['B3'].value, 2)
        self.assertIsNone(wb[LOG_SHEET]['I3'].value)

    def test_serial_main_uses_same_workbook_layout(self):
        self.extraction_source()
        serial = self.load_script('exstract_from_description')
        serial.INPUT_FILE = serial.OUTPUT_FILE = str(self.path)
        serial.TARGET_FIELDS = ['包装数']
        serial.SAVE_EVERY = 1
        with patch.object(serial, 'ask_enable_vision', return_value=False), patch.object(serial, 'extract_text_features', return_value={'包装数': '1'}), patch.object(serial.time, 'sleep'):
            serial.main()
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['B1'].value, '包装数')
        self.assertEqual(wb.active['B2'].value, 1)
        self.assertIn(LOG_SHEET, wb.sheetnames)

    def add_labels(self):
        self.extraction_source()
        wb = load_workbook(self.path)
        wb.active.cell(1, 9, '打标')
        wb.active.cell(2, 9, 0)
        wb.active.cell(3, 9, 1)
        wb.save(self.path)
        wb.close()

    def test_both_modes_filter_before_model_calls_and_preserve_row_positions(self):
        self.add_labels()
        original = self.path.read_bytes()
        for name in ('exstract_from_description', 'exstract_from_description_fast'):
            with self.subTest(mode=name):
                self.path.write_bytes(original)
                module = self.load_script(name)
                module.INPUT_FILE = module.OUTPUT_FILE = str(self.path)
                module.SAVE_EVERY = module.CONCURRENCY = 1
                module.TARGET_FIELDS = ['包装数']
                module.ENABLE_VISION_DEFAULT = False
                calls = []
                def serial_extract(row, **kwargs):
                    calls.append(row.name)
                    return {'包装数': '2'}
                def fast_extract(index, *args):
                    calls.append(index)
                    return index, {'_extracted': {'包装数': '2'}, '_text_fields': ['包装数']}, ''
                config = {'name': 'test', 'target_fields': ['包装数'], 'field_config': {'包装数': '数量'}}
                with patch.dict(os.environ, {'EXTRACT_FILTER_ENABLED': '1', 'EXTRACT_FILTER_COLUMN': '打标', 'EXTRACT_FILTER_VALUE': '1'}):
                    if name.endswith('_fast'):
                        with patch.object(module, 'load_product_config', return_value=('test', config)), patch.object(module, 'process_single_row', side_effect=fast_extract):
                            module.main()
                    else:
                        with patch.object(module, 'ask_enable_vision', return_value=False), patch.object(module, 'extract_text_features', side_effect=serial_extract), patch.object(module.time, 'sleep'):
                            module.main()
                self.assertEqual(calls, [1])
                wb = load_workbook(self.path)
                self.assertIsNone(wb.active['B2'].value)
                self.assertEqual(wb.active['B3'].value, 2)
                self.assertEqual(wb[LOG_SHEET].max_row, 2)
                self.assertEqual(wb[LOG_SHEET]['B2'].value, 3)
                wb.close()

    def test_no_matches_leave_source_and_output_untouched(self):
        self.add_labels()
        original = self.path.read_bytes()
        for name in ('exstract_from_description', 'exstract_from_description_fast'):
            module = self.load_script(name)
            module.INPUT_FILE = str(self.path)
            module.OUTPUT_FILE = str(self.path.parent / 'no_matches.xlsx')
            with patch.dict(os.environ, {'EXTRACT_FILTER_ENABLED': '1', 'EXTRACT_FILTER_COLUMN': '打标', 'EXTRACT_FILTER_VALUE': '9'}):
                config = {'name': 'test', 'target_fields': ['包装数'], 'field_config': {'包装数': '数量'}}
                with patch.object(module, 'load_product_config', return_value=('test', config), create=True), patch.object(module, 'ask_enable_vision', return_value=False, create=True), patch.object(module, 'extract_text_features', create=True) as extract:
                    module.main()
                extract.assert_not_called()
            self.assertEqual(self.path.read_bytes(), original)
            self.assertFalse(Path(module.OUTPUT_FILE).exists())

    def test_fast_filter_excludes_unmatched_rows_from_fetch_and_resume(self):
        from concurrent.futures import Future
        self.add_labels()
        fast = self.load_script('exstract_from_description_fast')
        fast.INPUT_FILE = fast.OUTPUT_FILE = str(self.path)
        fast.ENABLE_FETCH = True
        fast.ENABLE_VISION_DEFAULT = False
        fast.CONCURRENCY = 1
        wb = load_workbook(self.path)
        wb.active.cell(1, 10, fast.FETCH_LINK_CANDIDATES[0])
        for r in (2, 3):
            wb.active.cell(r, 10, 'https://www.amazon.com/dp/B000000001')
            wb.active.cell(r, 8, '')  # missing bullets
        wb.save(self.path)
        wb.close()
        config = {'name': 'test', 'target_fields': ['包装数'], 'field_config': {'包装数': '数量'}}
        fetched = Future()
        fetched.set_result({'ok': True, 'bullets': ['fetched'], 'images': []})
        calls = []
        def process(index, *args):
            calls.append(index)
            return index, {'_extracted': {'包装数': '2'}}, ''
        with patch.dict(os.environ, {'EXTRACT_FILTER_ENABLED': '1', 'EXTRACT_FILTER_COLUMN': '打标', 'EXTRACT_FILTER_VALUE': '1'}), patch.object(fast, 'load_product_config', return_value=('test', config)), patch.object(fast, 'process_single_row', side_effect=process), patch.object(fast, 'FetchPool') as pool, patch('core.reviews.product_fetcher.resolve_asin_and_url', return_value=('B000000001', 'url')) as resolve, patch.object(fast.CheckpointStore, 'remove'):
            pool.return_value.request.return_value = fetched
            fast.main()
            self.assertEqual(calls, [1])
            self.assertEqual(resolve.call_count, 1)
            # A different condition must not reuse successful rows from the first run.
            os.environ['EXTRACT_FILTER_VALUE'] = '0'
            fast.main()
            self.assertEqual(calls, [1, 0])
