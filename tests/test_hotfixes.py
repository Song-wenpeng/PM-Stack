# -*- coding: utf-8 -*-
"""Regression tests for PM Stack security, UI and self-update fixes."""

import hashlib
import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PyQt6.QtWidgets import QApplication, QComboBox

from core.config_manager import ConfigManager
from core.app_updater import download_verified_file, load_update_manifest
from core.main_window import MainWindow, SettingsDialog
from core.runner import ScriptRunner
from core.updater import _safe_extract_zip
from scripts.fix_sales_data import fix_single
from updater_main import install_update


class ConfigSecurityTests(unittest.TestCase):
    @unittest.skipUnless(os.name == "nt", "DPAPI is only available on Windows")
    def test_secret_round_trip_has_no_plaintext_at_rest(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "config.json"
            secret = "pm-stack-regression-secret"

            manager = ConfigManager(str(config_path))
            manager.update({"api_key": secret, "theme": "dark"})

            raw = config_path.read_text(encoding="utf-8")
            payload = json.loads(raw)
            self.assertNotIn(secret, raw)
            self.assertNotIn("api_key", payload)
            self.assertIn("api_key_protected", payload)

            restored = ConfigManager(str(config_path))
            self.assertEqual(restored.get("api_key"), secret)
            self.assertEqual(restored.get("theme"), "dark")

    @unittest.skipUnless(os.name == "nt", "DPAPI is only available on Windows")
    def test_legacy_plaintext_is_migrated_in_place(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "config.json"
            secret = "legacy-regression-secret"
            config_path.write_text(
                json.dumps({"api_key": secret, "theme": "light"}),
                encoding="utf-8",
            )

            manager = ConfigManager(str(config_path))
            raw = config_path.read_text(encoding="utf-8")
            self.assertEqual(manager.get("api_key"), secret)
            self.assertNotIn(secret, raw)
            self.assertIn("api_key_protected", raw)


class UpdaterSecurityTests(unittest.TestCase):
    def test_zip_traversal_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_data = io.BytesIO()
            with zipfile.ZipFile(archive_data, "w") as archive:
                archive.writestr("../plugins_evil/payload.py", "bad")
            archive_data.seek(0)

            target = Path(temp_dir) / "plugins"
            with zipfile.ZipFile(archive_data, "r") as archive:
                with self.assertRaises(ValueError):
                    _safe_extract_zip(archive, target)

            self.assertFalse((Path(temp_dir) / "plugins_evil").exists())

    def test_normal_plugin_archive_is_extracted(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_data = io.BytesIO()
            with zipfile.ZipFile(archive_data, "w") as archive:
                archive.writestr("modules/sample.py", "VALUE = 1\n")
            archive_data.seek(0)

            target = Path(temp_dir) / "plugins"
            with zipfile.ZipFile(archive_data, "r") as archive:
                _safe_extract_zip(archive, target)

            self.assertEqual(
                (target / "modules" / "sample.py").read_text(encoding="utf-8"),
                "VALUE = 1\n",
            )


class ApplicationUpdaterTests(unittest.TestCase):
    def test_local_manifest_resolves_relative_executable(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            executable = root / "PM Stack.exe"
            executable.write_bytes(b"new-pm-stack")
            digest = hashlib.sha256(executable.read_bytes()).hexdigest()
            manifest_path = root / "update.json"
            manifest_path.write_text(
                json.dumps({
                    "version": "v1.10.0",
                    "url": "PM Stack.exe",
                    "sha256": digest,
                    "size": executable.stat().st_size,
                    "mandatory": False,
                    "notes": "test release",
                }),
                encoding="utf-8",
            )

            manifest = load_update_manifest(str(manifest_path))
            self.assertEqual(manifest.version, "v1.10.0")
            self.assertEqual(Path(manifest.url), executable.resolve())
            self.assertEqual(manifest.sha256, digest)

    def test_http_manifest_is_rejected_before_network_access(self):
        with self.assertRaisesRegex(ValueError, "不允许使用 HTTP"):
            load_update_manifest("http://example.test/update.json")

    def test_download_checks_size_and_hash_before_publishing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.exe"
            source.write_bytes(b"verified-update")
            digest = hashlib.sha256(source.read_bytes()).hexdigest()
            destination = root / "downloaded.exe"

            result = download_verified_file(
                str(source), str(destination), digest, source.stat().st_size,
            )
            self.assertEqual(Path(result).read_bytes(), b"verified-update")

            bad_destination = root / "bad.exe"
            with self.assertRaisesRegex(ValueError, "SHA-256"):
                download_verified_file(
                    str(source), str(bad_destination), "0" * 64,
                    source.stat().st_size,
                )
            self.assertFalse(bad_destination.exists())

    def test_independent_updater_replaces_and_keeps_rollback_copy(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "download.exe"
            target = root / "PM Stack.exe"
            source.write_bytes(b"new-version")
            target.write_bytes(b"old-version")
            digest = hashlib.sha256(source.read_bytes()).hexdigest()

            backup = Path(install_update(
                str(source), str(target), digest, restart=False,
            ))
            self.assertEqual(target.read_bytes(), b"new-version")
            self.assertEqual(backup.read_bytes(), b"old-version")
            self.assertFalse(source.exists())

    def test_external_update_channel_is_used_when_user_setting_is_blank(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            config_path = root / "config.json"
            channel_path = root / "update-channel.json"
            channel_path.write_text(
                json.dumps({"manifest_source": r"\\server\share\update.json"}),
                encoding="utf-8",
            )
            manager = ConfigManager(str(config_path))
            with mock.patch.object(manager, "get_data_dir", return_value=str(root)):
                self.assertEqual(
                    manager.get_app_update_manifest_source(),
                    r"\\server\share\update.json",
                )


class RunnerIsolationTests(unittest.TestCase):
    def test_runpy_tasks_are_globally_serialized(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            slow_script = Path(temp_dir) / "slow.py"
            fast_script = Path(temp_dir) / "fast.py"
            slow_script.write_text(
                "import time\ntime.sleep(0.25)\nprint('slow done')\n",
                encoding="utf-8",
            )
            fast_script.write_text("print('fast done')\n", encoding="utf-8")

            first = ScriptRunner()
            second = ScriptRunner()
            self.assertTrue(first.run_script(str(slow_script)))
            self.assertFalse(second.run_script(str(fast_script)))
            first._thread.join(timeout=3)
            self.assertFalse(first.is_running())

            self.assertTrue(second.run_script(str(fast_script)))
            second._thread.join(timeout=3)
            self.assertFalse(second.is_running())

    def test_main_window_uses_one_runner_per_module(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = str(Path(temp_dir) / "config.json")
            with mock.patch(
                "core.main_window.ConfigManager",
                side_effect=lambda: ConfigManager(config_path),
            ):
                window = MainWindow()
                try:
                    runners = [widget.runner for _, widget in window._modules]
                    self.assertGreaterEqual(len(runners), 5)
                    self.assertEqual(len(runners), len({id(item) for item in runners}))
                finally:
                    window.close()
                    app.processEvents()


class UIInteractionTests(unittest.TestCase):
    def test_dropdowns_and_dependent_controls_keep_visible_state_consistent(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = str(Path(temp_dir) / "config.json")
            with mock.patch(
                "core.main_window.ConfigManager",
                side_effect=lambda: ConfigManager(config_path),
            ):
                window = MainWindow()

            try:
                modules = {
                    info["name"]: widget for info, widget in window._modules
                }
                comments = modules["评论分析"]
                comments_test = modules["评论分析TEST"]
                sales = modules["销量数据"]
                market = modules["市场规划"]
                matrix = modules["产品矩阵"]

                combos = window.findChildren(QComboBox)
                # 评论分析 TEST 模块也包含一个品类下拉框。
                self.assertEqual(len(combos), 10)
                self.assertTrue(all(combo.count() > 0 for combo in combos))
                self.assertEqual(comments_test.concurrency_spin.value(), 8)
                self.assertEqual(comments_test.rpm_spin.value(), 900)
                self.assertEqual(comments_test.tpm_spin.value(), 90000)

                first = str(Path(temp_dir) / "first.xlsx")
                second = str(Path(temp_dir) / "second.xlsx")
                comments.s1_input.setText(first)
                comments.s1_input.setText(second)
                self.assertEqual(
                    comments.s1_output.text(),
                    str(Path(temp_dir) / "second_AI处理.xlsx"),
                )
                comments.s1_output.setText(str(Path(temp_dir) / "manual.xlsx"))
                comments.s1_input.setText(str(Path(temp_dir) / "third.xlsx"))
                self.assertEqual(
                    comments.s1_output.text(),
                    str(Path(temp_dir) / "manual.xlsx"),
                )

                comments.s1_product.setCurrentIndex(1)
                comments.s2_product.setCurrentIndex(0)
                comments.s2_input.setText(first)
                with mock.patch.object(comments, "_run") as run_comments:
                    comments._run_step2()
                run_comments.assert_not_called()
                self.assertIn("选择一个分析品类", comments.log_s2.text.toPlainText())

                sales.cross_filter_rows[0][0].setText("品牌")
                sales.cross_filter_rows[0][1].setText("A")
                sales.cross_filter_count.setValue(3)
                self.assertEqual(sales.cross_filter_rows[0][0].text(), "品牌")
                self.assertEqual(sales.cross_filter_rows[0][1].text(), "A")
                sales.cross_filter_count.setValue(1)
                self.assertEqual(sales.cross_filter_rows[0][0].text(), "品牌")
                self.assertEqual(sales.cross_filter_rows[0][1].text(), "A")

                sales.trend_folder.setText(temp_dir)
                sales.trend_config.setText(first)
                sales.trend_field.setEditText("   ")
                with mock.patch.object(sales, "_run") as run_sales:
                    sales._run_trend()
                run_sales.assert_not_called()
                self.assertIn("提取字段", sales.log1.text.toPlainText())

                sales.cross_trend_mode.setCurrentIndex(1)
                self.assertEqual(sales._cross_data_label.text(), "数据Sheet:")
                self.assertEqual(
                    sales._cross_data2_label.text(), "第二数据Sheet *:")
                self.assertIn("销额趋势", sales.cross_mode_hint.text())
                sales.cross_trend_mode.setCurrentIndex(2)
                self.assertIn("均价", sales.cross_mode_hint.text())

                self.assertFalse(matrix.m3_sort_order.isEnabled())
                matrix.m3_sort_col.setText("月销量")
                self.assertTrue(matrix.m3_sort_order.isEnabled())
                matrix.m3_sort_col.clear()
                self.assertFalse(matrix.m3_sort_order.isEnabled())

                market.tag_input.setText(first)
                market._stop_file = str(Path(temp_dir) / "stop.signal")
                with mock.patch.object(market, "_run", return_value=False):
                    market._run_tagging()
                self.assertFalse(market.stop_btn1.isEnabled())

                dialog = SettingsDialog(window.config_mgr, window.updater, window)
                try:
                    self.assertTrue(dialog.auto_app_update_cb.isChecked())
                    dialog.app_manifest_edit.setText(
                        str(Path(temp_dir) / "missing-update.json"))
                    with mock.patch.object(
                        dialog.app_update_checker, "check_update", return_value=True,
                    ) as check_app_update:
                        dialog._check_app_update()
                    check_app_update.assert_called_once_with(
                        str(Path(temp_dir) / "missing-update.json"), "v1.10.2")
                    self.assertFalse(dialog.check_app_update_btn.isEnabled())
                    dialog._on_app_check_done(False, "验证完成")
                    self.assertTrue(dialog.check_app_update_btn.isEnabled())

                    dialog.github_repo_edit.setText("owner/repo")
                    with mock.patch.object(
                        dialog.updater, "check_update"
                    ) as check_update:
                        dialog._check_update()
                        dialog._check_update()
                    check_update.assert_called_once()
                    self.assertFalse(dialog.check_update_btn.isEnabled())
                    dialog._on_check_done(False, "检查完成")
                    self.assertTrue(dialog.check_update_btn.isEnabled())
                finally:
                    dialog.close()
            finally:
                window.close()
                app.processEvents()


class ExcelSafetyTests(unittest.TestCase):
    def test_fix_preserves_workbook_and_creates_backup(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "market_2025-03_销量拆分.xlsx"
            workbook = Workbook()
            data_sheet = workbook.active
            data_sheet.title = "数据"
            data_sheet.append(["ASIN", "说明", "子体销量_矫正", "子体销售额_矫正"])
            data_sheet.append(["B00TEST", "保留", 10, 20.5])
            data_sheet["B1"].font = Font(bold=True)
            formula_sheet = workbook.create_sheet("公式")
            formula_sheet["A1"] = "=1+1"
            workbook.save(workbook_path)
            workbook.close()

            self.assertTrue(
                fix_single(temp_dir, "B00TEST", "2025-03", 99, 123.5)
            )

            updated = load_workbook(workbook_path, data_only=False)
            try:
                self.assertEqual(updated["数据"]["C2"].value, 99)
                self.assertEqual(updated["数据"]["D2"].value, 123.5)
                self.assertTrue(updated["数据"]["B1"].font.bold)
                self.assertEqual(updated["公式"]["A1"].value, "=1+1")
            finally:
                updated.close()

            backups = list(Path(temp_dir).glob("*.backup_*.xlsx"))
            self.assertEqual(len(backups), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
