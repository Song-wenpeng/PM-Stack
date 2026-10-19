"""Exercise real resume/writeback scheduling with model/browser calls replaced."""
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from core.field_workbook import FieldWorkbookWriter, LOG_SHEET


class FieldResumeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / 'source.xlsx'
        self.output = Path(self.temp.name) / 'result.xlsx'
        self.checkpoint = Path(str(self.output) + '.checkpoint.json')
        wb = Workbook()
        ws = wb.active
        ws.append(['ASIN', '商品标题', '详细参数', 'SKU', '五点描述',
                   '商品详情页链接', '包装数', 'USB接口', '打标'])
        ws.append(['B000000001', 'title', '', '', '', 'url-one', '', '', 1])
        ws.append(['B000000002', 'title', '', '', '', 'url-two', 3, 'manual', 1])
        ws.append(['B000000003', 'title', '', '', '', 'url-three', '', '', 0])
        wb.save(self.source)
        wb.close()
        location = Path(__file__).resolve().parents[1] / 'scripts/exstract_from_description_fast.py'
        spec = importlib.util.spec_from_file_location('field_resume_fast', location)
        self.fast = importlib.util.module_from_spec(spec)
        self.env = patch.dict(os.environ, {
            'API_KEY': 'test-only', 'ENABLE_FETCH': '0', 'FIELDS': '',
            'CURRENT_PRODUCT': '', 'EXTRACT_CHECKPOINT_FILE': '',
            'EXTRACT_FILTER_ENABLED': '1', 'EXTRACT_FILTER_COLUMN': '打标',
            'EXTRACT_FILTER_VALUE': '1',
        })
        self.env.start()
        self.addCleanup(self.env.stop)
        spec.loader.exec_module(self.fast)
        self.fast.INPUT_FILE = str(self.source)
        self.fast.OUTPUT_FILE = str(self.output)
        self.fast.ENABLE_VISION_DEFAULT = False
        self.fast.CONCURRENCY = self.fast.SAVE_EVERY = 1
        self.config = {'name': 'test', 'target_fields': ['包装数', 'USB接口'],
                       'field_config': {'包装数': '数量', 'USB接口': '接口'}}

    def run_main(self, extract):
        with patch.object(self.fast, 'load_product_config', return_value=('test', self.config)), \
             patch.object(self.fast, 'extract_text_features_concurrent', side_effect=extract) as model, \
             patch.object(self.fast, 'check_stop_file', return_value=False):
            self.fast.main()
        return model

    def frame(self):
        return pd.read_excel(self.output, dtype=object)

    def log(self):
        return pd.read_excel(self.output, sheet_name=LOG_SHEET).iloc[0]

    def test_incomplete_run_keeps_cache_and_next_run_only_requests_missing_field(self):
        first = self.run_main(lambda row, missing, *args: {'包装数': '2'})
        self.assertEqual(first.call_count, 1)  # no same-run infinite retries
        record = json.loads(self.checkpoint.read_text(encoding='utf-8'))['rows']['0']
        self.assertEqual(record['status'], '未完整')
        self.assertFalse(record['result']['_complete'])
        self.assertEqual(record['result']['_missing_fields'], ['USB接口'])
        self.assertEqual(self.log()['_处理状态'], '未完整')
        self.assertEqual(self.log()['提取完整状态'], '否')

        def fill(row, missing, *args):
            self.assertEqual(row['包装数'], '2')
            self.assertEqual(missing, ['USB接口'])
            return {'USB接口': '2A1C'}

        second = self.run_main(fill)
        self.assertEqual(second.call_count, 1)
        self.assertFalse(self.checkpoint.exists())
        result = self.frame()
        self.assertEqual(result.loc[0, '包装数'], 2)
        self.assertEqual(result.loc[0, 'USB接口'], '2A1C')
        self.assertEqual(result.loc[1, 'USB接口'], 'manual')
        self.assertTrue(pd.isna(result.loc[2, 'USB接口']))
        self.assertEqual(self.log()['提取完整状态'], '是')
        self.assertEqual(set(self.log()['文本提取字段'].split('；')), {'包装数', 'USB接口'})

    def test_legacy_success_stopped_and_failed_records_recover_partial_values(self):
        self.run_main(lambda row, missing, *args: {'包装数': '2'})
        original = json.loads(self.checkpoint.read_text(encoding='utf-8'))
        for status in ('', '成功', '已停止', 'AI提取失败'):
            with self.subTest(status=status):
                record = original['rows']['0']
                record['status'] = status
                record['result'].pop('_complete', None)
                record['result'].pop('_missing_fields', None)
                self.checkpoint.write_text(json.dumps(original, ensure_ascii=False), encoding='utf-8')

                def fill(row, missing, *args):
                    self.assertEqual(row['包装数'], '2')
                    self.assertEqual(missing, ['USB接口'])
                    return {'USB接口': 'restored'}

                self.assertEqual(self.run_main(fill).call_count, 1)
                self.assertEqual(self.frame().loc[0, '包装数'], 2)

    def test_manual_values_win_over_cache_and_incomplete_still_fetches(self):
        from concurrent.futures import Future
        self.fast.ENABLE_FETCH = True
        cached = Future()
        cached.set_result({'ok': True, 'bullets': ['spec'], 'images': ['image-url']})
        with patch.object(self.fast, 'FetchPool') as pool, \
             patch('core.reviews.product_fetcher.resolve_asin_and_url', return_value=('B000000001', 'url')) as resolve:
            pool.return_value.request.return_value = cached
            self.run_main(lambda row, missing, *args: {'包装数': '2'})
            self.assertEqual(resolve.call_count, 1)  # complete/unselected rows never fetched

        wb = load_workbook(self.source)
        wb.active['G2'] = 9  # selected attribute changes keep a valid checkpoint
        wb.save(self.source)
        wb.close()
        with patch.object(self.fast, 'FetchPool') as pool, \
             patch('core.reviews.product_fetcher.resolve_asin_and_url', return_value=('B000000001', 'url')):
            pool.return_value.request.return_value = cached

            def fill(row, missing, *args):
                self.assertEqual(row['包装数'], 9)
                self.assertEqual(missing, ['USB接口'])
                return {'USB接口': 'filled'}

            self.run_main(fill)
        self.assertEqual(self.frame().loc[0, '包装数'], 9)

    def test_restored_complete_rows_skip_models_and_browser_even_without_link_column(self):
        self.fast.ENABLE_FETCH = True
        self.fast.FETCH_LINK_CANDIDATES = ['absent-link-column']
        # Initially save a complete checkpoint without needing a browser.
        self.fast.ENABLE_FETCH = False
        with patch.object(self.fast.CheckpointStore, 'remove'):
            self.run_main(lambda row, missing, *args: {'包装数': '2', 'USB接口': 'cached'})
        data = json.loads(self.checkpoint.read_text(encoding='utf-8'))
        data['signature']['enable_fetch'] = True
        self.checkpoint.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        self.fast.ENABLE_FETCH = True
        with patch.object(self.fast, 'FetchPool') as pool:
            model = self.run_main(lambda *args: self.fail('complete row must not call model'))
        model.assert_not_called()
        pool.assert_not_called()
        self.assertFalse(self.checkpoint.exists())
        self.assertEqual(self.frame().loc[0, 'USB接口'], 'cached')

    def test_placeholders_are_missing_but_real_values_and_zero_are_preserved(self):
        wb = load_workbook(self.source)
        wb.active['G2'], wb.active['H2'] = 0, '未知'
        wb.save(self.source)
        wb.close()

        def fill(row, missing, *args):
            self.assertEqual(missing, ['USB接口'])
            return {'包装数': '99', 'USB接口': '2A'}

        self.run_main(fill)
        self.assertEqual(self.frame().loc[0, '包装数'], 0)
        self.assertEqual(self.frame().loc[0, 'USB接口'], '2A')

    def test_stopping_again_does_not_discard_recovered_partial_results(self):
        self.run_main(lambda row, missing, *args: {'包装数': '2'})
        with patch.object(self.fast, 'load_product_config', return_value=('test', self.config)), \
             patch.object(self.fast, 'check_stop_file', return_value=True), \
             patch.object(self.fast, 'extract_text_features_concurrent') as model:
            with self.assertRaises(SystemExit):
                self.fast.main()
        model.assert_not_called()
        record = json.loads(self.checkpoint.read_text(encoding='utf-8'))['rows']['0']
        self.assertEqual(record['status'], '已停止')
        self.assertEqual(record['result']['_extracted']['包装数'], '2')
        self.assertEqual(self.frame().loc[0, '包装数'], 2)

    def test_existing_workbook_evidence_survives_skipping_and_filling_without_cache(self):
        evidence = [{'批次': 1, '通过过滤后的结果': {'包装数': '2'}}]
        for complete in (True, False):
            with self.subTest(complete=complete):
                frame = pd.read_excel(self.source, dtype=object)
                frame.loc[0, '包装数'] = 2
                frame.loc[0, 'USB接口'] = 'known' if complete else None
                frame['图片识别批次日志'] = [json.dumps(evidence, ensure_ascii=False), '', '']
                frame['图片识别提取字段'] = ['包装数', '', '']
                FieldWorkbookWriter(self.source, self.source, self.config['target_fields']).save(frame)
                # The prior subcase wrote 'known'; clear it directly to simulate
                # an incomplete workbook without relying on any checkpoint.
                if not complete:
                    wb = load_workbook(self.source)
                    columns = {c.value: c.column for c in wb.active[1]}
                    wb.active.cell(2, columns['USB接口']).value = None
                    wb.save(self.source)
                    wb.close()
                model = self.run_main(lambda row, missing, *args: {'USB接口': 'filled'})
                self.assertEqual(model.call_count, 0 if complete else 1)
                self.assertEqual(json.loads(self.log()['图片识别批次日志']), evidence)
                self.assertEqual(self.log()['图片识别提取字段'], '包装数')


if __name__ == '__main__':
    unittest.main()
