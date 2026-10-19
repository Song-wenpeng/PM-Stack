"""ASIN-list filtering, output safety and extraction UI regression checks."""
from tests import support
from pathlib import Path
import tempfile
import unittest
from unittest import mock

import pandas as pd
from openpyxl import Workbook, load_workbook
from scripts.extract_history_data import extract, read_asin_list


class ExtractHistoryDataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root/'source'
        self.source.mkdir()
        for month, quantities in (('2025-01', [2, 4, 3]), ('2025-02', [6, 8, 1])):
            pd.DataFrame({'ASIN': ['b001', 'B002', 'b001'],
                          '子体销量_矫正': quantities,
                          '子体销售额_矫正': [v*10 for v in quantities]}).to_excel(
                self.source/f'US_{month}_销量拆分.xlsx', index=False)
        self.output = self.root/'history.xlsx'

    def run_extract(self, **kwargs):
        extract(str(self.source), '子体销量_矫正', '子体销售额_矫正', str(self.output), **kwargs)

    def list_file(self, values=(' b002 ', 'B001', 'b002', None, 'MISSING')):
        path = self.root/'products.xlsx'
        book = Workbook()
        book.active.title = '说明'
        book.active.append(['说明'])
        sheet = book.create_sheet('商品属性')
        sheet.append(['图片', ' ASIN ', '品牌'])
        for value in values:
            sheet.append([None, value, 'test'])
        sheet.row_dimensions[3].hidden = True
        book.save(path)
        book.close()
        return path

    def test_existing_workbook_picker_skips_replace_prompt(self):
        from PyQt6.QtWidgets import QFileDialog, QLineEdit
        from core.widgets import file_row
        support.qt_app()
        line = QLineEdit()
        row = file_row(line, 'save_no_confirm')
        with mock.patch('core.widgets.QFileDialog.getSaveFileName', return_value=('existing.xlsx', '')) as choose:
            row.itemAt(1).widget().click()
        self.assertEqual(choose.call_args.kwargs['options'], QFileDialog.Option.DontConfirmOverwrite)

    def test_tab_removed_supplement_and_passes_asin_list_from_real_button(self):
        from modules.sales_data import ModuleWidget
        support.qt_app()
        self.addCleanup(support.cleanup_qt_widgets)
        widget = ModuleWidget(None, None)
        self.assertFalse(hasattr(widget, 'trend_second_file'))
        self.assertFalse(hasattr(widget, 'trend_visible_only'))
        widget.trend_folder.setText(str(self.source))
        widget.trend_output.setText(str(self.output))
        widget.trend_asin_file.setText(str(self.list_file()))
        widget.trend_asin_sheet.setText('商品属性')
        with mock.patch.object(widget, '_run') as run:
            widget.run_btn2.click()
        args = run.call_args.args[2]
        self.assertEqual(args[args.index('--asin-sheet')+1], '商品属性')
        self.assertIn('--asin-file', args)
        self.assertNotIn('--second-file', args)
        widget.trend_asin_file.clear()
        self.assertFalse(widget.trend_asin_sheet.isEnabled())
        widget.close()

    def test_no_list_extracts_all_asins_and_sums_duplicate_rows(self):
        self.run_extract()
        result = pd.read_excel(self.output, sheet_name='历史销量').set_index('ASIN')
        self.assertEqual(result.index.tolist(), ['B001', 'B002'])
        self.assertEqual(result.loc['B001'].tolist(), [5, 7])

    def test_list_reads_column_b_all_rows_deduplicates_and_retains_order(self):
        path = self.list_file()
        asins, sheet = read_asin_list(str(path), '商品属性')
        self.assertEqual(asins, ['B002', 'B001', 'MISSING'])
        self.assertEqual(sheet, '商品属性')
        self.run_extract(asin_file=str(path), asin_sheet=sheet)
        volume = pd.read_excel(self.output, sheet_name='历史销量').set_index('ASIN')
        revenue = pd.read_excel(self.output, sheet_name='历史销额').set_index('ASIN')
        self.assertEqual(volume.index.tolist(), asins)
        self.assertEqual(revenue.index.tolist(), asins)
        self.assertEqual(revenue.loc['B002'].tolist(), [40, 80])
        self.assertTrue(volume.loc['MISSING'].isna().all())

    def test_only_requested_asins_are_exported(self):
        path = self.list_file(['B002'])
        self.run_extract(asin_file=str(path), asin_sheet='商品属性')
        self.assertEqual(pd.read_excel(self.output, sheet_name='历史销量')['ASIN'].tolist(), ['B002'])

    def test_missing_header_empty_list_and_wrong_sheet_report_errors(self):
        path = self.list_file()
        with self.assertRaisesRegex(ValueError, 'ASIN 列'):
            read_asin_list(str(path))
        with self.assertRaisesRegex(ValueError, '不存在'):
            read_asin_list(str(path), '不存在')
        path = self.list_file([None, '   '])
        with self.assertRaisesRegex(ValueError, '没有可读取'):
            read_asin_list(str(path), '商品属性')

    def test_no_matches_does_not_modify_existing_output(self):
        self.run_extract()
        original = self.output.read_bytes()
        path = self.list_file(['MISSING'])
        with self.assertRaisesRegex(ValueError, '均未'):
            self.run_extract(asin_file=str(path), asin_sheet='商品属性')
        self.assertEqual(self.output.read_bytes(), original)

    def test_existing_list_workbook_can_receive_history_and_is_backed_up(self):
        self.output = self.list_file(['B002'])
        original = self.output.read_bytes()
        self.run_extract(asin_file=str(self.output), asin_sheet='商品属性')
        book = load_workbook(self.output, data_only=True)
        self.assertEqual(book['商品属性']['B2'].value, 'B002')
        self.assertIn('说明', book.sheetnames)
        book.close()
        backup = list(self.root.glob('*.backup'))
        self.assertEqual(len(backup), 1)
        self.assertEqual(backup[0].read_bytes(), original)
        with self.assertRaisesRegex(ValueError, '不能覆盖'):
            self.run_extract(asin_file=str(self.output), asin_sheet='商品属性', output_volume_sheet='商品属性')

    def test_duplicate_month_error_lists_conflicting_files_and_ignores_backups(self):
        name = self.source/'Competitor-2025.01-001_2025-01_销量拆分.xlsx'
        (self.source/'US_2025-01_销量拆分.backup_20261009.xlsx').write_bytes(
            (self.source/'US_2025-01_销量拆分.xlsx').read_bytes())
        self.run_extract()
        name.write_bytes((self.source/'US_2025-01_销量拆分.xlsx').read_bytes())
        with self.assertRaisesRegex(ValueError, 'Competitor-2025'):
            self.run_extract()


if __name__ == '__main__':
    unittest.main()
