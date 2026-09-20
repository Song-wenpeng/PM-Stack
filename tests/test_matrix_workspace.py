from tests import support

import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, Mock
import pandas as pd
from PyQt6.QtCore import Qt, QEvent
from PyQt6.QtWidgets import QApplication, QListWidgetItem
from core.matrix_workspace import MatrixWorkspace, prepare_data


class MatrixWorkspaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = support.qt_app()

    def setUp(self):
        self.addCleanup(support.cleanup_qt_widgets)
        self.widget = MatrixWorkspace()
        self.widget.frame = prepare_data(pd.DataFrame({
            'ASIN': ['A', 'B'], '品牌': ['甲', '乙'],
            '三级分类': ['普通', '工具'], '上架时间': ['2024-01-01', '2025-01-01']}))
        for field in self.widget.frame.columns:
            item = QListWidgetItem(field)
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(Qt.CheckState.Checked if field == 'ASIN' else Qt.CheckState.Unchecked)
            self.widget.fields.addItem(item)
        for combo in (self.widget.horizontal, self.widget.vertical, self.widget.sort):
            combo.addItems(['不分组'] + list(self.widget.frame.columns))
        self.widget.render()

    def tearDown(self):
        self.widget.refresh_timer.stop()
        self.widget.close()
        self.widget.deleteLater()

    def test_drop_zones_and_canvas(self):
        w = self.widget
        w.apply_field_drop('horizontal', '品牌')
        w.apply_field_drop('vertical', '上架年份')
        self.assertTrue(w.refresh_timer.isActive())
        w.render()
        self.assertEqual((w.table.rowCount(), w.table.columnCount()), (3, 3))
        w.fields.setCurrentRow(2)
        event = Mock()
        event.type.return_value = QEvent.Type.Drop
        event.source.return_value = w.fields
        self.assertTrue(w.eventFilter(w.table.viewport(), event))
        self.assertIn('三级分类', w.selected_fields())
        w.render()
        self.assertTrue(any('三级分类: 普通' in cards.item(0).text() for cards in w.cells.values()))
        self.assertEqual(w.table.viewport().property('matrixDropRole'), 'cards')
        self.assertEqual(next(iter(w.cells.values())).viewport().property('matrixDropRole'), 'cards')

    def test_compact_resizable_panels(self):
        w = self.widget
        w.resize(1300, 800)
        w.show()
        self.app.processEvents()
        self.assertLessEqual(w.workspace_split.sizes()[0], 210)
        self.assertLessEqual(w.canvas_split.sizes()[0], 200)
        self.assertGreater(w.matrix_split.sizes()[1], 700)
        w.canvas_split.setSizes([230, 450])
        w.matrix_split.setSizes([240, 850])
        self.app.processEvents()
        self.assertGreater(w.canvas_split.sizes()[0], 200)
        self.assertGreater(w.matrix_split.sizes()[0], 200)
        self.assertEqual(w.card_summary.text(), '卡片字段')
        self.assertNotIn('勾选', w.vertical.value_title.text())

    def test_sparse_export_and_manual_order(self):
        w = self.widget
        cards = w.cells[(0, 0)]
        cards.insertItem(0, cards.takeItem(1))
        with tempfile.TemporaryDirectory() as folder:
            target = str(Path(folder) / 'matrix.xlsx')
            with patch('core.matrix_workspace.QFileDialog.getSaveFileName', return_value=(target, '')):
                w.export()
            detail = pd.read_excel(target, sheet_name='产品明细')
            self.assertEqual(detail['ASIN'].tolist(), ['B', 'A'])
            w.apply_field_drop('horizontal', '品牌')
            w.apply_field_drop('vertical', '上架年份')
            w.render()
            with patch('core.matrix_workspace.QFileDialog.getSaveFileName', return_value=(target, '')):
                w.export()
            self.assertEqual(len(pd.read_excel(target, sheet_name='产品明细')), 2)

    def test_large_sparse_matrix_not_blank(self):
        w = self.widget
        w.frame = prepare_data(pd.DataFrame({'ASIN': [str(i) for i in range(30)], '品牌': [str(i) for i in range(30)],
                                            '上架年份': [str(i) for i in range(30)]}))
        w.horizontal.setCurrentText('品牌')
        w.vertical.setCurrentText('上架年份')
        w.render()
        self.assertEqual(len(w.row_keys) * len(w.col_keys), 900)
        self.assertEqual(len(w.cells), 30)

    def test_nested_visibility_order_template_and_export(self):
        w = self.widget
        w.frame = prepare_data(pd.DataFrame({
            'ASIN': ['A', 'B', 'C', 'D'], '品牌': ['甲'] * 4,
            '三级分类': ['普通', '普通', '工具', '工具'],
            '排插外观': ['条形', '三角', '条形', '三角']}))
        w.vertical.addItem('排插外观')
        w.apply_field_drop('vertical', '三级分类')
        w.apply_field_drop('vertical', '排插外观')
        w.apply_field_drop('vertical', '三级分类')
        w.render()
        self.assertEqual(w.vertical.fields(), ['三级分类', '排插外观'])
        self.assertEqual(len(w.cells), 4)
        w.vertical.options['三级分类'] = [['普通', True], ['工具', True]]
        w.vertical.options['排插外观'] = [['条形', True], ['三角', True]]
        w.render()
        self.assertEqual(w.row_names[0], '三级分类：普通\n  排插外观：条形')
        w.vertical.options['三级分类'][1][1] = False
        w.render()
        self.assertEqual(len(w.row_names), 2)
        with tempfile.TemporaryDirectory() as folder:
            template = str(Path(folder) / 'layout.json')
            output = str(Path(folder) / 'matrix.xlsx')
            with patch('core.matrix_workspace.QFileDialog.getSaveFileName', return_value=(template, '')):
                w.save_template()
            w.vertical.restore('不分组')
            with patch('core.matrix_workspace.QFileDialog.getOpenFileName', return_value=(template, '')):
                w.load_template()
            self.assertEqual(w.vertical.fields(), ['三级分类', '排插外观'])
            self.assertEqual(len(w.row_names), 2)
            with patch('core.matrix_workspace.QFileDialog.getSaveFileName', return_value=(output, '')):
                w.export()
            detail = pd.read_excel(output, sheet_name='产品明细')
            self.assertEqual(detail['ASIN'].tolist(), ['A', 'B'])
            self.assertEqual(detail['纵向分组'].tolist(), w.row_names)
        w.vertical.options['三级分类'] = [['普通', False], ['工具', False]]
        w.render()
        self.assertEqual(w.table.rowCount(), w.header_rows)
        self.assertFalse(w.cells)

    def test_reorder_levels_values_and_swap(self):
        w = self.widget
        w.apply_field_drop('vertical', '三级分类')
        w.apply_field_drop('vertical', '品牌')
        w.render()
        axis = w.vertical
        axis.levels.insertItem(0, axis.levels.takeItem(1))
        w.render()
        self.assertEqual(axis.fields(), ['品牌', '三级分类'])
        axis.levels.setCurrentRow(0)
        axis.values.insertItem(0, axis.values.takeItem(1))
        axis.save_values()
        w.render()
        saved = axis.payload()
        w.swap_axes()
        self.assertEqual(w.horizontal.payload(), saved)
        self.assertEqual(w.vertical.fields(), [])

    def test_merged_headers_match_export_and_product_offsets(self):
        from openpyxl import load_workbook
        w = self.widget
        w.frame = prepare_data(pd.DataFrame({
            'ASIN': ['A', 'B', 'C', 'D', 'E'],
            '品牌': ['甲', '甲', '甲', '乙', '乙'],
            '三级分类': ['普通', '普通', '普通', '工具', '工具'],
            '上架时间': ['2024-01-01', '2024-01-01', '2025-01-01', '2024-01-01', '2025-01-01']}))
        for axis, fields in [(w.vertical, ['三级分类', '上架年份']),
                             (w.horizontal, ['品牌', '上架年份'])]:
            for field in fields:
                axis.add_field(field)
        w.render()
        self.assertEqual((w.header_rows, w.header_cols), (2, 2))
        for start in (0, 2):
            self.assertEqual(w.table.rowSpan(2 + start, 0), 2)
            self.assertEqual(w.table.columnSpan(0, 2 + start), 2)
            self.assertEqual(w.table.item(2 + start, 1).text(), '2024')
            self.assertEqual(w.table.rowSpan(2 + start, 1), 1)
        for (ri, ci), cards in w.cells.items():
            self.assertIs(w.table.cellWidget(2 + ri, 2 + ci), cards)
        with tempfile.TemporaryDirectory() as folder:
            output = str(Path(folder) / 'merged.xlsx')
            with patch('core.matrix_workspace.QFileDialog.getSaveFileName', return_value=(output, '')):
                w.export()
            book = load_workbook(output)
            sheet = book['自由组合矩阵']
            self.assertEqual(sheet.freeze_panes, 'C3')
            merged = {str(r) for r in sheet.merged_cells.ranges}
            self.assertTrue({'A1:A2', 'B1:B2', 'C1:D1', 'E1:F1'}.issubset(merged))
            # Export can occupy several rows for a leaf with several product cards.
            offset = 3
            expected_parent_ranges = []
            for ri, key in enumerate(w.row_keys):
                height = max([1] + [cards.count() for (row, _), cards in w.cells.items() if row == ri])
                if ri % 2 == 0:
                    parent_start = offset
                self.assertEqual(sheet.cell(offset, 2).value, key[1])
                for (row, ci), cards in w.cells.items():
                    if row == ri:
                        for i in range(cards.count()):
                            self.assertEqual(sheet.cell(offset + i, 3 + ci).value, cards.item(i).text())
                offset += height
                if ri % 2:
                    expected_parent_ranges.append(f'A{parent_start}:A{offset - 1}')
            self.assertTrue(set(expected_parent_ranges).issubset(merged))
            book.close()
        # Re-render must clear previous merged spans when dropping a level.
        w.vertical.setCurrentText('三级分类')
        w.render()
        self.assertEqual(w.header_cols, 1)
        self.assertEqual(w.table.rowSpan(2, 0), 1)


if __name__ == '__main__':
    unittest.main()
