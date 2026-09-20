from tests import support

import hashlib
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.reviews.exporter import export_reviews_for_analysis
from core.reviews.migration import import_legacy_database
from core.reviews.scraper import AmazonReviewScraper, CollectionCancelled
from core.reviews.store import ReviewStore


class ReviewAnalysisExportTests(unittest.TestCase):
    def test_exports_chinese_content_column_and_star_sheets(self):
        rows = [
            {
                "asin": "B0TEST0001",
                "review_id": "R1",
                "rating": 5,
                "content": "great product",
                "title": "Great",
            },
            {
                "asin": "B0TEST0001",
                "review_id": "R2",
                "rating": 1,
                "content": "bad product",
                "title": "Bad",
            },
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "analysis.xlsx")
            export_reviews_for_analysis(rows, path)
            workbook = load_workbook(path, read_only=True)
            self.assertEqual(workbook.sheetnames, ["5 star", "1 star"])
            headers = [cell.value for cell in next(workbook["5 star"].iter_rows())]
            self.assertIn("内容", headers)
            content_index = headers.index("内容") + 1
            self.assertEqual(workbook["5 star"].cell(2, content_index).value, "great product")
            workbook.close()


class ReviewCancellationTests(unittest.TestCase):
    def test_wait_can_be_cancelled_cooperatively(self):
        scraper = AmazonReviewScraper(should_stop=lambda: True)
        with self.assertRaises(CollectionCancelled):
            scraper._random_delay(0.01, 0.01)


class ReviewLegacyImportTests(unittest.TestCase):
    def test_import_is_non_destructive_and_idempotent(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = os.path.join(temp_dir, "source.db")
            destination_path = os.path.join(temp_dir, "destination.db")
            source = ReviewStore(source_path)
            source.upsert_product("B0TEST0001", "amazon.com", title="Test")
            source.add_reviews(
                "B0TEST0001",
                "amazon.com",
                [{"review_id": "R1", "rating": 5, "content": "great"}],
            )
            source.close()
            before = hashlib.sha256(Path(source_path).read_bytes()).hexdigest()

            destination = ReviewStore(destination_path)
            result = import_legacy_database(destination, source_path)
            self.assertEqual(result["products"], 1)
            self.assertEqual(result["reviews"], 1)
            self.assertEqual(destination.total_reviews(), 1)
            with self.assertRaises(RuntimeError):
                import_legacy_database(destination, source_path)
            destination.close()

            after = hashlib.sha256(Path(source_path).read_bytes()).hexdigest()
            self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
