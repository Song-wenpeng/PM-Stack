"""Category screening regressions: synthetic workbooks and fake model only."""
from tests import support

from contextlib import chdir, redirect_stdout
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import runpy
import tempfile
import types
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts/欧洲市场规划/类目初筛.py'
COLUMNS = ['Node root', '是否考虑', '网址链接', 'Node ID', 'Node Path', '二级类目', '类目路径_中文']


def source_book(path, count=100):
    book = Workbook()
    book.active.title = '类目'
    book.active.append(COLUMNS)
    for i in range(count):
        book.active.append(['root', 1, 'https://example.invalid', i, f'path{i}', '二级', f'类目{i}'])
    book.active['A1'].fill = PatternFill('solid', fgColor='FF0000')
    book.create_sheet('保留公式').append(['=SUM(1,2)', '不得更改'])
    book.save(path)
    book.close()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_script():
    spec = importlib.util.spec_from_file_location('category_screening_test', SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class CategoryReproductionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.fake = MagicMock()
        self.fake.chat.completions.create.return_value = types.SimpleNamespace(
            choices=[types.SimpleNamespace(message=types.SimpleNamespace(content='{"results": []}'))])
        self.factory = MagicMock(return_value=self.fake)
        self.addCleanup(patch.stopall)
        patch.dict('sys.modules', {'openai': types.SimpleNamespace(OpenAI=self.factory)}).start()
        patch('time.sleep').start()

    def execute(self, source, output=''):
        env = {'API_KEY': 'FAKE-NOT-A-KEY', 'SILICONFLOW_API_KEY': '',
               'CATEGORY_INPUT_FILE': str(source), 'CATEGORY_OUTPUT_FILE': str(output),
               'CATEGORY_STATE_DIR': str(self.root / 'state'), 'CATEGORY_STOP_FILE': ''}
        with patch.dict(os.environ, env):
            os.environ.pop('CATEGORY_CHECKPOINT_FILE', None)
            try:
                runpy.run_path(str(SCRIPT), run_name='__main__')
            except SystemExit as exc:
                return exc.code
        return 0

    def test_import_needs_no_key_or_client(self):
        with patch.dict(os.environ, {'API_KEY': '', 'SILICONFLOW_API_KEY': ''}):
            load_script()
        self.factory.assert_not_called()

    def test_empty_output_preserves_100_row_multisheet_original(self):
        source = self.root / 'source.xlsx'
        digest = source_book(source)
        try:
            code = self.execute(source)
        finally:
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), digest)
        self.assertTrue(source.with_name('source_类目筛选结果.xlsx').is_file())
        self.assertNotEqual(code, 0)

    def test_failed_result_as_input_is_retried(self):
        source = self.root / 'source.xlsx'
        source_book(source, 2)
        result = self.root / 'result.xlsx'
        self.execute(source, result)
        self.fake.reset_mock()
        self.execute(result, self.root / 'second.xlsx')
        self.assertGreater(self.fake.chat.completions.create.call_count, 0)
        frame = pd.read_excel(result)
        self.assertTrue(frame['匹配度等级'].isna().all())
        self.assertEqual(frame['处理状态'].tolist(), ['可重试失败'] * 2)


def valid_result(module, idx, score=4, confidence='高'):
    result = dict.fromkeys(module.OUTPUT_FIELDS, '合成判断')
    result.update(row_index=idx, 电工归属='核心电工', 匹配度等级='高匹配',
                  是否建议进入下一轮='建议重点验证', 判断信心=confidence)
    for field in module.SCORE_FIELDS:
        result[field] = score
    result['总分'] = score * 6
    result['百分制得分'] = score * 20
    return result


class CategoryStateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'source.xlsx'
        self.original = source_book(self.source, 4)
        self.output = self.root / 'output.xlsx'
        self.checkpoint = self.root / 'state' / 'checkpoint.json'
        self.env = patch.dict(os.environ, {'API_KEY': '', 'SILICONFLOW_API_KEY': '',
                                          'CATEGORY_STATE_DIR': str(self.root / 'state')})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.module = load_script()
        self.sleep = patch.object(self.module.time, 'sleep')
        self.sleep.start()
        self.addCleanup(self.sleep.stop)
        self.fake = MagicMock()
        self.fake.chat.completions.create.side_effect = AssertionError('real API path not allowed')
        self.log = io.StringIO()

    def execute(self, evaluator=None, **kwargs):
        source = kwargs.pop('input_file', self.source)
        output = kwargs.pop('output_file', self.output)
        checkpoint = kwargs.pop('checkpoint_file', self.checkpoint)
        if evaluator is None:
            evaluator = lambda items, *args: {'results': [valid_result(self.module, x['row_index']) for x in items]}
        with patch.object(self.module, 'evaluate_categories_batch', side_effect=evaluator) as evaluate, redirect_stdout(self.log):
            summary = self.module.run_screening(source, output, checkpoint, client=self.fake, **kwargs)
        return summary, evaluate

    def assert_source_unchanged(self):
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.original)

    def test_success_100_rows_keeps_source_and_saves_each_batch(self):
        self.original = source_book(self.source, 100)
        original_write = self.module.write_json
        with patch.object(self.module, 'write_json', wraps=original_write) as save:
            summary, evaluate = self.execute()
        self.assertEqual(summary['成功'], 100)
        self.assertEqual(summary['exit_code'], 0)
        self.assertEqual(evaluate.call_count, 5)
        self.assertGreaterEqual(save.call_count, 5)
        self.assert_source_unchanged()
        self.assertEqual(len(json.loads(self.checkpoint.read_text(encoding='utf-8'))['rows']), 100)

    def test_partial_batch_keeps_valid_rows_and_resume_only_retries_failures(self):
        summary, _ = self.execute(lambda items, *args: {'results': [valid_result(self.module, 3, 5), valid_result(self.module, 1, 2)]})
        self.assertEqual((summary['成功'], summary['失败'], summary['exit_code']), (2, 2, 2))
        saved = pd.read_excel(self.output)
        self.assertEqual(saved['Node ID'].tolist()[:2], [3, 1])
        state = json.loads(self.checkpoint.read_text(encoding='utf-8'))
        self.assertEqual(state['rows']['3']['总分'], 30)
        summary, evaluate = self.execute()
        self.assertEqual([r['row_index'] for r in evaluate.call_args.args[0]], [0, 2])
        self.assertEqual((summary['成功'], summary['跳过']), (2, 2))
        self.assert_source_unchanged()

    def test_sorted_output_as_source_resumes_by_row_status(self):
        self.execute(lambda items, *args: {'results': [valid_result(self.module, 3, 5), valid_result(self.module, 1, 2)]})
        summary, evaluate = self.execute(input_file=self.output, output_file=self.root / 'again.xlsx')
        items = evaluate.call_args.args[0]
        self.assertEqual([r['Node ID'] for r in items], ['0', '2'])
        self.assertEqual((summary['成功'], summary['跳过']), (2, 2))

    def test_no_pending_needs_no_client_or_key(self):
        self.execute()
        with patch.dict('sys.modules', {'openai': types.SimpleNamespace(OpenAI=MagicMock(side_effect=AssertionError('unneeded')))}), redirect_stdout(self.log):
            summary = self.module.run_screening(self.source, self.output, self.checkpoint)
        self.assertEqual((summary['跳过'], summary['成功'], summary['exit_code']), (4, 0, 0))

    def test_all_rows_excluded_needs_no_client_or_key(self):
        frame = pd.read_excel(self.source)
        frame['是否考虑'] = 0
        frame.to_excel(self.source, index=False)
        original = self.source.read_bytes()
        factory = MagicMock(side_effect=AssertionError('must not initialize'))
        with patch.dict('sys.modules', {'openai': types.SimpleNamespace(OpenAI=factory)}), redirect_stdout(self.log):
            summary = self.module.run_screening(self.source, self.output, self.checkpoint)
        factory.assert_not_called()
        self.assertEqual((summary['跳过'], summary['未处理'], summary['exit_code']), (4, 0, 0))
        self.assertEqual(self.source.read_bytes(), original)

    def test_report_parameter_change_retries_except_human_confirmation(self):
        self.execute()
        frame = pd.read_excel(self.output)
        frame.loc[0, self.module.STATUS] = self.module.CONFIRMED
        frame.loc[0, '简短结论'] = '保留用户确认'
        frame.to_excel(self.output, index=False)
        original = self.output.read_bytes()
        summary, evaluate = self.execute(input_file=self.output, output_file=self.root / 'again.xlsx',
                                         company_profile='新画像')
        self.assertEqual((summary['成功'], summary['跳过']), (3, 1))
        self.assertEqual([r['row_index'] for r in evaluate.call_args.args[0]], [1, 2, 3])
        self.assertEqual(self.output.read_bytes(), original)
        self.assertIn('保留用户确认', pd.read_excel(self.root / 'again.xlsx')['简短结论'].tolist())

    def test_changed_content_order_model_or_profile_does_not_reuse_checkpoint(self):
        self.execute()
        for kwargs in ({'model_name': 'synthetic-model-2'}, {'company_profile': 'synthetic-profile-2'},
                       {'base_url': 'https://example.invalid/v1'}, {'batch_size': 2}):
            with self.subTest(kwargs=kwargs):
                summary, evaluate = self.execute(**kwargs)
                self.assertEqual(summary['成功'], 4)
                self.assertGreater(evaluate.call_count, 0)
                self.execute()  # restore the baseline parameter binding
        frame = pd.read_excel(self.source)
        frame = frame.iloc[::-1].copy()
        frame.to_excel(self.source, index=False)
        summary, _ = self.execute()
        self.assertEqual(summary['成功'], 4)
        frame.loc[0, 'Node Path'] = 'changed'
        frame.to_excel(self.source, index=False)
        summary, _ = self.execute()
        self.assertEqual(summary['成功'], 4)

    def test_identical_copy_cannot_reuse_another_sources_explicit_checkpoint(self):
        self.execute()
        copy = self.root / 'other-source.xlsx'
        copy.write_bytes(self.source.read_bytes())
        summary, evaluate = self.execute(input_file=copy)
        self.assertEqual(summary['成功'], 4)
        self.assertEqual(evaluate.call_count, 1)

    def test_malformed_checkpoint_metadata_is_ignored_per_row(self):
        self.execute()
        state = json.loads(self.checkpoint.read_text(encoding='utf-8'))
        state['rows']['0'][self.module.ERROR] = ['invalid', 'metadata']
        state['rows']['1'][self.module.STATUS] = self.module.SUCCESS
        state['rows']['1']['判断信心'] = '低'
        self.checkpoint.write_text(json.dumps(state, ensure_ascii=False), encoding='utf-8')
        summary, evaluate = self.execute()
        self.assertEqual([x['row_index'] for x in evaluate.call_args.args[0]], [0, 1])
        self.assertEqual((summary['成功'], summary['跳过']), (2, 2))

    def test_huge_numeric_and_overlong_model_text_do_not_lose_valid_partial_rows(self):
        for field, bad in (('总分', 10 ** 500), ('产品相邻性', 10 ** 500),
                           ('简短结论', 'x' * 32768)):
            with self.subTest(field=field):
                invalid = valid_result(self.module, 0)
                invalid[field] = bad
                valid, errors, _ = self.module.validate_batch(
                    {'results': [invalid, valid_result(self.module, 1)]}, [0, 1])
                self.assertEqual(set(valid), {1})
                self.assertEqual(set(errors), {0})

    def test_changed_report_row_is_reevaluated_but_other_successes_skip(self):
        self.execute()
        frame = pd.read_excel(self.output)
        frame.loc[0, 'Node Path'] = 'changed row'
        frame.to_excel(self.output, index=False)
        summary, evaluate = self.execute(input_file=self.output, output_file=self.root / 'again.xlsx')
        self.assertEqual((summary['成功'], summary['跳过']), (1, 3))
        self.assertEqual(evaluate.call_args.args[0][0]['Node Path'], 'changed row')

    def test_legacy_success_manual_review_and_confirmation_are_distinct(self):
        frame = pd.read_excel(self.source)
        for field, value in valid_result(self.module, 0).items():
            if field != 'row_index':
                frame[field] = [value, '', '', '']
        frame.loc[1, '匹配度等级'] = '人工审核'
        frame.loc[1, '简短结论'] = '旧的模型失败描述'
        frame[self.module.STATUS] = ['', '', '人工确认', '待人工审核']
        frame.loc[2, '简短结论'] = '用户明确确认，不覆盖'
        frame.to_excel(self.source, index=False)
        summary, evaluate = self.execute()
        self.assertEqual([x['row_index'] for x in evaluate.call_args.args[0]], [1, 3])
        self.assertEqual(summary['跳过'], 2)
        saved = pd.read_excel(self.output).set_index('Node ID')
        self.assertEqual(saved.loc[2, '简短结论'], '用户明确确认，不覆盖')
        self.assertEqual(saved.loc[2, '处理状态'], '人工确认')

    def test_low_confidence_is_pending_review_not_completed(self):
        summary, _ = self.execute(lambda items, *args: {'results': [valid_result(self.module, x['row_index'], confidence='低') for x in items]})
        self.assertEqual((summary['成功'], summary['待审'], summary['exit_code']), (0, 4, 2))
        self.assertEqual(set(pd.read_excel(self.output)['处理状态']), {'待人工审核'})
        summary, evaluate = self.execute()
        self.assertEqual(evaluate.call_count, 1)
        self.assertEqual(summary['成功'], 4)

    def test_structure_and_row_id_validation_preserves_only_valid_rows(self):
        for payload in (None, [], {}, {'results': {}}, {'results': []}):
            with self.subTest(payload=payload):
                valid, errors, _ = self.module.validate_batch(payload, [0, 1])
                self.assertEqual(valid, {})
                self.assertEqual(set(errors), {0, 1})
        cases = [True, False, 0.2, 0.0, '0', -1, 10, None]
        for bad in cases:
            with self.subTest(bad=bad):
                invalid = valid_result(self.module, bad)
                valid, errors, issues = self.module.validate_batch({'results': [invalid, valid_result(self.module, 1)]}, [0, 1])
                self.assertEqual(set(valid), {1})
                self.assertEqual(set(errors), {0})
                self.assertTrue(issues)
        duplicate = valid_result(self.module, 0)
        valid, errors, _ = self.module.validate_batch({'results': [duplicate, duplicate, valid_result(self.module, 1), 'bad']}, [0, 1, 2])
        self.assertEqual(set(valid), {1})
        self.assertIn('重复', errors[0])
        self.assertIn(2, errors)

    def test_required_fields_enums_and_scores_are_strict(self):
        for field in self.module.OUTPUT_FIELDS:
            result = valid_result(self.module, 0)
            del result[field]
            self.assertTrue(self.module.validate_result(result), field)
        for field in self.module.ENUMS:
            result = valid_result(self.module, 0)
            result[field] = '人工审核'
            self.assertTrue(self.module.validate_result(result), field)
        for field in self.module.SCORE_FIELDS:
            for value in (True, -1, 5.1, float('nan'), float('inf'), float('-inf'), '4', None):
                with self.subTest(field=field, value=value):
                    result = valid_result(self.module, 0)
                    result[field] = value
                    self.assertTrue(self.module.validate_result(result))
        for field, value in (('总分', 23), ('百分制得分', 99), ('总分', float('nan')), ('百分制得分', True)):
            result = valid_result(self.module, 0)
            result[field] = value
            self.assertTrue(self.module.validate_result(result))

    def test_illegal_xml_model_text_is_retryable_not_poisoned_success(self):
        for text in ('invalid\x00text', 'invalid\ud800text', 'invalid\ufffftext'):
            with self.subTest(text=repr(text)):
                invalid = valid_result(self.module, 0)
                invalid['简短结论'] = text
                valid, errors, _ = self.module.validate_batch(
                    {'results': [invalid, valid_result(self.module, 1)]}, [0, 1])
                self.assertEqual(set(valid), {1})
                self.assertEqual(set(errors), {0})

    def test_fake_openai_valid_response_uses_effective_parameters(self):
        response = {'results': [valid_result(self.module, i) for i in range(4)]}
        self.fake.chat.completions.create.side_effect = None
        self.fake.chat.completions.create.return_value = types.SimpleNamespace(
            choices=[types.SimpleNamespace(message=types.SimpleNamespace(
                content='```json\n' + json.dumps(response, ensure_ascii=False) + '\n```'))])
        with redirect_stdout(self.log):
            summary = self.module.run_screening(self.source, self.output, self.checkpoint,
                                               client=self.fake, model_name='fake-effective-model',
                                               company_profile='合成画像，不调用真实API')
        self.assertEqual(summary['成功'], 4)
        call = self.fake.chat.completions.create.call_args.kwargs
        self.assertEqual(call['model'], 'fake-effective-model')
        self.assertIn('合成画像', call['messages'][1]['content'])
        self.assert_source_unchanged()

    def test_invalid_scores_become_failure_not_business_conclusion(self):
        invalid = valid_result(self.module, 0)
        invalid['产品相邻性'] = 9
        summary, _ = self.execute(lambda *args: {'results': [invalid, valid_result(self.module, 1)]})
        saved = pd.read_excel(self.output).set_index('Node ID')
        self.assertEqual((summary['成功'], summary['失败']), (1, 3))
        self.assertEqual(saved.loc[0, '处理状态'], '可重试失败')
        self.assertTrue(pd.isna(saved.loc[0, '匹配度等级']))
        self.assertTrue(pd.isna(saved.loc[0, '简短结论']))

    def test_exception_and_logs_never_expose_request_secret(self):
        secret = 'sk-THIS-MUST-NOT-APPEAR'
        self.fake.chat.completions.create.side_effect = RuntimeError(secret)
        with redirect_stdout(self.log):
            summary = self.module.run_screening(self.source, self.output, self.checkpoint, client=self.fake)
        self.assertEqual(summary['失败'], 4)
        self.assertNotIn(secret, self.log.getvalue())
        self.assertNotIn(secret, self.checkpoint.read_text(encoding='utf-8'))

    def test_source_and_checkpoint_aliases_rejected_before_api_or_writes(self):
        for kwargs in ({'output_file': self.source}, {'output_file': str(self.source).upper()},
                       {'checkpoint_file': self.source}, {'checkpoint_file': self.output}):
            with self.subTest(kwargs=kwargs):
                with self.assertRaises(ValueError):
                    self.execute(**kwargs)
        self.assert_source_unchanged()
        self.assertFalse(self.output.exists())
        self.assertFalse(self.checkpoint.exists())

    def test_cooperative_stop_saves_first_batch_then_resume(self):
        summary, evaluate = self.execute(batch_size=2, should_stop=lambda: self.checkpoint.exists())
        self.assertEqual((summary['成功'], summary['未处理'], summary['exit_code']), (2, 2, 130))
        self.assertTrue(self.output.exists())
        self.assertEqual(evaluate.call_count, 1)
        summary, evaluate = self.execute(batch_size=2)
        self.assertEqual([r['row_index'] for r in evaluate.call_args.args[0]], [2, 3])
        self.assertEqual(summary['跳过'], 2)
        self.assert_source_unchanged()

    def test_keyboard_interrupt_saves_completed_rows(self):
        calls = []
        def evaluate(items, *args):
            calls.append(items)
            if len(calls) == 2:
                raise KeyboardInterrupt
            return {'results': [valid_result(self.module, x['row_index']) for x in items]}
        summary, _ = self.execute(evaluate, batch_size=2)
        self.assertEqual((summary['成功'], summary['未处理'], summary['exit_code']), (2, 2, 130))
        self.assertEqual(len(json.loads(self.checkpoint.read_text(encoding='utf-8'))['rows']), 2)
        self.assert_source_unchanged()

    def test_stop_before_request_does_not_initialize_client(self):
        with patch.dict('sys.modules', {'openai': types.SimpleNamespace(OpenAI=MagicMock(side_effect=AssertionError('no client')))}), redirect_stdout(self.log):
            summary = self.module.run_screening(self.source, self.output, self.checkpoint, should_stop=lambda: True)
        self.assertEqual((summary['未处理'], summary['exit_code']), (4, 130))
        self.assertTrue(self.output.exists())

    def test_output_replace_failure_keeps_old_and_completed_checkpoint(self):
        self.output.write_bytes(b'old target')
        real_replace = os.replace
        def replace(source, target):
            if Path(target) == self.output:
                raise PermissionError('locked')
            return real_replace(source, target)
        with patch('core.report_io.os.replace', side_effect=replace):
            summary, _ = self.execute()
        self.assertEqual(summary['exit_code'], 1)
        self.assertEqual(self.output.read_bytes(), b'old target')
        recovered = list(self.root.glob('output.pending-*.xlsx'))
        self.assertEqual(len(recovered), 1)
        self.assertEqual(len(pd.read_excel(recovered[0])), 4)
        summary, evaluate = self.execute()
        evaluate.assert_not_called()
        self.assertEqual(summary['跳过'], 4)
        self.assert_source_unchanged()

    def test_checkpoint_replace_failure_stops_requests_but_exports_recovery(self):
        self.checkpoint.parent.mkdir()
        self.checkpoint.write_text('{"old":true}', encoding='utf-8')
        real_replace = os.replace
        def replace(source, target):
            if Path(target) == self.checkpoint:
                raise PermissionError('locked')
            return real_replace(source, target)
        with patch('core.report_io.os.replace', side_effect=replace):
            summary, evaluate = self.execute(batch_size=2)
        self.assertEqual((summary['成功'], summary['未处理'], summary['exit_code']), (2, 2, 1))
        self.assertEqual(evaluate.call_count, 1)
        self.assertEqual(json.loads(self.checkpoint.read_text()), {'old': True})
        self.assertTrue(self.output.exists())
        self.assertTrue(list(self.checkpoint.parent.glob('checkpoint.pending-*.json')))
        self.assert_source_unchanged()

    def test_corrupt_checkpoint_is_ignored(self):
        self.checkpoint.parent.mkdir()
        self.checkpoint.write_text('not JSON', encoding='utf-8')
        summary, _ = self.execute()
        self.assertEqual(summary['成功'], 4)
        self.assertIn('断点无法读取', self.log.getvalue())

    def test_extraneous_model_rows_do_not_report_all_success(self):
        summary, _ = self.execute(lambda items, *args: {'results': [valid_result(self.module, x['row_index']) for x in items] + [valid_result(self.module, 999)]})
        self.assertEqual(summary['成功'], 4)
        self.assertEqual((summary['协议异常'], summary['exit_code']), (1, 2))

    def test_cli_same_source_and_checkpoint_fail_before_reading(self):
        aliases = [str(self.source), str(self.source).upper(), './source.xlsx', 'absent/../source.xlsx']
        for alias in aliases:
            for key in ('CATEGORY_OUTPUT_FILE', 'CATEGORY_CHECKPOINT_FILE'):
                env = {'CATEGORY_INPUT_FILE': str(self.source), 'CATEGORY_OUTPUT_FILE': str(self.output),
                       'CATEGORY_CHECKPOINT_FILE': str(self.checkpoint), 'CATEGORY_STOP_FILE': '', key: alias}
                with self.subTest(key=key, alias=alias), patch.dict(os.environ, env), chdir(self.root), redirect_stdout(self.log), patch.object(pd, 'read_excel', side_effect=AssertionError('reject first')):
                    self.assertEqual(self.module.main(), 1)
        self.assert_source_unchanged()
        self.assertFalse(self.output.exists())

    def test_cli_stop_file_saves_without_request(self):
        stop = self.root / 'stop.json'
        stop.write_text('{"stop":true}', encoding='utf-8')
        env = {'CATEGORY_INPUT_FILE': str(self.source), 'CATEGORY_OUTPUT_FILE': str(self.output),
               'CATEGORY_CHECKPOINT_FILE': str(self.checkpoint), 'CATEGORY_STOP_FILE': str(stop)}
        with patch.dict(os.environ, env), redirect_stdout(self.log):
            self.assertEqual(self.module.main(), 130)
        self.assertTrue(self.output.exists())
        self.assertTrue(self.checkpoint.exists())
        self.assert_source_unchanged()

    def test_100_row_source_unchanged_on_stop_and_write_error(self):
        self.original = source_book(self.source, 100)
        summary, _ = self.execute(should_stop=lambda: self.checkpoint.exists())
        self.assertEqual(summary['exit_code'], 130)
        self.assert_source_unchanged()
        self.output.write_bytes(b'old')
        with patch.object(pd.DataFrame, 'to_excel', side_effect=OSError('disk full')):
            summary, _ = self.execute()
        self.assertEqual(summary['exit_code'], 1)
        self.assertEqual(self.output.read_bytes(), b'old')
        self.assert_source_unchanged()


class CategoryUIProtectionTests(unittest.TestCase):
    def setUp(self):
        from core.config_manager import ConfigManager
        from modules.market_planning import ModuleWidget
        self.app = support.qt_app()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(support.cleanup_qt_widgets)
        self.root = Path(self.tmp.name)
        self.runner = MagicMock()
        self.runner.is_running.return_value = False
        self.runner.is_any_running.return_value = False
        self.env = patch.dict(os.environ, {'CATEGORY_STATE_DIR': str(self.root / 'state'), 'CATEGORY_CHECKPOINT_FILE': ''})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.widget = ModuleWidget(ConfigManager(str(self.root / 'config.json')), self.runner)
        self.widget.cat_input.setText(str(self.root / 'source.xlsx'))

    def tearDown(self):
        from PyQt6.QtCore import QThreadPool
        QThreadPool.globalInstance().waitForDone(5000)

    def test_ui_default_output_and_checkpoint_are_distinct(self):
        from core.report_io import same_file
        with patch.object(self.widget, '_run', return_value=True) as run:
            self.assertTrue(self.widget._run_category())
        env = run.call_args.args[2]
        self.assertEqual(env['CATEGORY_OUTPUT_FILE'], str(self.root / 'source_类目筛选结果.xlsx'))
        self.assertTrue(env['CATEGORY_CHECKPOINT_FILE'].endswith('.json'))
        self.assertFalse(same_file(env['CATEGORY_OUTPUT_FILE'], env['CATEGORY_CHECKPOINT_FILE']))
        self.assertIn(env['CATEGORY_OUTPUT_FILE'], self.widget.cat_output_hint.text())
        self.assertIn(env['CATEGORY_OUTPUT_FILE'], self.widget._result_sources[self.widget.run_btn3]()[0])

    def test_ui_output_aliases_and_checkpoint_source_do_not_start_runner(self):
        for alias in (str(self.root / 'source.xlsx'), str(self.root / 'source.xlsx').upper(), './source.xlsx', 'absent/../source.xlsx'):
            with self.subTest(alias=alias), chdir(self.root), patch.object(self.widget, '_run') as run:
                self.widget.cat_output.setText(alias)
                self.assertFalse(self.widget._run_category())
                run.assert_not_called()
        self.widget.cat_output.clear()
        with patch.dict(os.environ, {'CATEGORY_CHECKPOINT_FILE': str(self.root / 'source.xlsx')}), patch.object(self.widget, '_run') as run:
            self.assertFalse(self.widget._run_category())
        run.assert_not_called()
        self.runner.run_script.assert_not_called()

    def test_stop_button_only_writes_cooperative_signal_and_waits(self):
        with patch.object(self.widget, '_run', return_value=True) as run:
            self.widget._run_category()
        env = run.call_args.args[2]
        self.widget.stop_btn3.click()
        self.assertTrue(Path(env['CATEGORY_STOP_FILE']).exists())
        self.assertFalse(self.widget.stop_btn3.isEnabled())
        self.runner.cancel.assert_not_called()
        self.assertIn('安全保存', self.widget.log3.text.toPlainText())
        run.call_args.kwargs['finished_callback'](130)
        self.assertIn('已停止', self.widget.run_btn3._task_bar.summary.text())
