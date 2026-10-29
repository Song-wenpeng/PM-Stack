"""Atomic saves and shared reference-date parsing; synthetic files only."""
from tests import support

from contextlib import chdir
from datetime import date, datetime
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from core import report_io as rio


class ReferenceDateTests(unittest.TestCase):
    """Shared reference-date parsing: explicit formats only, never a fuzzy guess."""

    def test_blank_falls_back_to_run_day(self):
        for raw in ('', '   ', None):
            with self.subTest(raw=repr(raw)):
                stamp, source = rio.resolve_reference_date(raw, today=date(2026, 10, 28))
                self.assertEqual(stamp, datetime(2026, 10, 28))
                self.assertIn('默认', source)

    def test_explicit_formats_are_accepted(self):
        cases = [('2026-07-31', (2026, 7, 31)), ('2026/07/31', (2026, 7, 31)),
                 ('2026.07.31', (2026, 7, 31)), ('2026-7-1', (2026, 7, 1))]
        for text, expected in cases:
            with self.subTest(text=text):
                stamp, source = rio.resolve_reference_date(text)
                self.assertEqual((stamp.year, stamp.month, stamp.day), expected)
                self.assertIn('显式指定', source)

    def test_ambiguous_or_invalid_values_are_rejected(self):
        for bad in ('31/07/2026', '2026-13-45', 'not-a-date', '15', '2026年7月31日'):
            with self.subTest(bad=bad):
                with self.assertRaises(ValueError) as caught:
                    rio.resolve_reference_date(bad, env_name='MARKET_ANALYSIS_DATE')
                message = str(caught.exception)
                self.assertIn('MARKET_ANALYSIS_DATE', message)
                self.assertIn(bad, message)


class ReportIOTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.target = self.root / 'target.json'
        self.target.write_text('{"old": true}', encoding='utf-8')
        self.old = self.target.read_bytes()

    def test_aliases_and_hardlinks_are_rejected(self):
        aliases = [str(self.target).upper(), self.target.name, self.root / '.' / self.target.name,
                   self.root / 'nonexistent' / '..' / self.target.name]
        for alias in aliases:
            with self.subTest(alias=alias), chdir(self.root):
                with self.assertRaises(ValueError):
                    rio.require_distinct(self.target, alias)
        linked = self.root / 'hardlink.json'
        os.link(self.target, linked)
        with self.assertRaises(ValueError):
            rio.require_distinct(self.target, linked)

    def test_source_guard_runs_before_writer(self):
        with patch('core.report_io.tempfile.mkstemp') as create:
            with self.assertRaises(ValueError):
                rio.write_json(self.target, {'new': True}, protected_paths=(self.target,))
        create.assert_not_called()
        self.assertEqual(self.target.read_bytes(), self.old)

    def test_transient_replace_locks_recover_with_bounded_backoff(self):
        real_replace = os.replace
        calls = []

        def locked(source, target):
            calls.append((Path(source), Path(target)))
            if len(calls) < 3:
                raise PermissionError('scanner')
            real_replace(source, target)

        with patch('core.report_io.os.replace', side_effect=locked), patch('core.report_io.time.sleep') as sleep:
            rio.write_json(self.target, {'new': True})
        self.assertEqual(len(calls), 3)
        self.assertEqual(sleep.call_count, 2)
        self.assertTrue(all(a.parent == b.parent for a, b in calls))
        self.assertEqual(json.loads(self.target.read_text()), {'new': True})

    def test_permanent_lock_keeps_old_and_validated_recovery_file(self):
        with patch('core.report_io.os.replace', side_effect=PermissionError('locked')) as replace, patch('core.report_io.time.sleep'):
            with self.assertRaises(rio.SafeWriteError) as caught:
                rio.write_json(self.target, {'new': True})
        self.assertEqual(replace.call_count, len(rio.RETRY_DELAYS) + 1)
        self.assertTrue(caught.exception.validated)
        self.assertEqual(self.target.read_bytes(), self.old)
        self.assertEqual(json.loads(Path(caught.exception.temporary).read_text()), {'new': True})

    def test_writer_and_validation_failure_never_replace(self):
        for phase in ('writer', 'validator'):
            with self.subTest(phase=phase):
                def write(path):
                    path.write_text('partial', encoding='utf-8')
                    if phase == 'writer':
                        raise OSError('disk full')
                def validate(path):
                    raise ValueError('invalid payload')
                with patch('core.report_io.os.replace') as replace:
                    with self.assertRaises(rio.SafeWriteError) as caught:
                        rio.atomic_write(self.target, write, validate)
                replace.assert_not_called()
                self.assertFalse(caught.exception.validated)
                self.assertEqual(self.target.read_bytes(), self.old)

    def test_transient_writer_lock_recovers(self):
        calls = []
        def write(path):
            calls.append(path)
            if len(calls) < 3:
                raise PermissionError('scanner')
            path.write_text('{}', encoding='utf-8')
        with patch('core.report_io.time.sleep'):
            rio.atomic_write(self.target, write, lambda p: json.loads(p.read_text()))
        self.assertEqual(len(calls), 3)
        self.assertEqual(self.target.read_text(), '{}')

    def test_excel_all_sheets_are_readable_and_text_is_not_formula(self):
        target = self.root / 'report.xlsx'
        rio.write_excel(target, {'结果': pd.DataFrame({'结论': ['=HYPERLINK("https://example.invalid")', '正常']}),
                                '参数': pd.DataFrame({'参数': ['测试']})})
        with target.open('rb') as stream:
            book = load_workbook(stream)
            self.assertEqual(book.sheetnames, ['结果', '参数'])
            self.assertEqual(book['结果']['A2'].data_type, 's')
            self.assertEqual(book['结果'].max_row, 3)
            book.close()

    def test_invalid_excel_candidate_is_not_replaced(self):
        target = self.root / 'report.xlsx'
        target.write_bytes(b'OLD')
        with patch.object(pd.DataFrame, 'to_excel', return_value=None):
            with self.assertRaises(rio.SafeWriteError):
                rio.write_excel(target, {'结果': pd.DataFrame({'x': [1]})})
        self.assertEqual(target.read_bytes(), b'OLD')

    def test_excel_silent_cell_corruption_is_rejected(self):
        target = self.root / 'report.xlsx'
        target.write_bytes(b'old result')
        real_to_excel = pd.DataFrame.to_excel

        def corrupt(frame, writer, *args, **kwargs):
            real_to_excel(frame, writer, *args, **kwargs)
            writer.sheets[kwargs['sheet_name']]['A2'] = 'silently changed'

        with patch.object(pd.DataFrame, 'to_excel', new=corrupt):
            with self.assertRaises(rio.SafeWriteError):
                rio.write_excel(target, {'结果': pd.DataFrame({'结论': ['expected']})})
        self.assertEqual(target.read_bytes(), b'old result')

    def test_excel_truncated_text_does_not_replace_target(self):
        target = self.root / 'report.xlsx'
        target.write_bytes(b'old result')
        with self.assertRaises(rio.SafeWriteError):
            rio.write_excel(target, {'结果': pd.DataFrame({'结论': ['x' * 32768]})})
        self.assertEqual(target.read_bytes(), b'old result')

    def test_json_nonfinite_is_rejected_without_touching_target(self):
        with self.assertRaises(ValueError):
            rio.write_json(self.target, {'score': float('nan')})
        self.assertEqual(self.target.read_bytes(), self.old)
