"""Execute both stages through the real runner; stub scripts avoid paid calls."""
from tests import support
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock
from openpyxl import Workbook, load_workbook
from core.config_manager import ConfigManager
from core.runner import ScriptRunner
from modules.comment_analysis import ModuleWidget


class FormalCommentConcurrencyTests(unittest.TestCase):
    def test_real_runner_chains_stages_and_report_history(self):
        app = support.qt_app()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scripts = root / 'scripts'
            scripts.mkdir()
            (scripts / 'comments_step_1_fast.py').write_text('''
import os
import shutil
assert os.environ['STEP1_CONCURRENCY'] == '8'
assert os.environ['STEP1_MAX_FAILURE_RATE'] == '0'
shutil.copyfile(os.environ['RAW_INPUT_FILE'], os.environ['STEP1_OUTPUT_FILE'])
print('进度: 1/1 (100%)')
''', encoding='utf-8')
            (scripts / 'comments_step_2.py').write_text('''
import os
from openpyxl import Workbook
book = Workbook()
book.active.title = 'VOC看板'
book.active['A1'] = '测试报告 · 1 条评论'
sheet = book.create_sheet('摘要汇总')
sheet.append(['维度', '类别', '提及次数', '占比', '评论数'])
sheet.append(['负向关注点', '耐用性', 1, 1.0, 1])
book.save(os.environ['STEP2_OUTPUT_FILE'])
''', encoding='utf-8')
            source = root / 'reviews.xlsx'
            book = Workbook()
            book.active.append(['内容', '星级'])
            book.active.append(['Broke after a week', 1])
            book.save(source)
            manager = ConfigManager(str(root / 'config.json'))
            widget = ModuleWidget(manager, ScriptRunner())
            try:
                template_path = manager.resolve_config('product_configs.json')
                with mock.patch.object(manager, 'resolve_config', return_value=template_path), mock.patch.object(manager, 'get_script_dir', return_value=str(scripts)), mock.patch.object(manager, 'get_env_dict', return_value={'API_KEY': 'not-used'}):
                    widget.product_combo.setCurrentIndex(widget.product_combo.findData('general'))
                    widget._add_inputs([str(source)])
                    self._wait(app, lambda: widget._worker is None)
                    widget._start_analysis()
                    self._wait(app, lambda: bool(widget.jobs.records) and widget.jobs.records[0]['status'] == 'done')
                    self._wait(app, lambda: not widget.runner.is_running() and widget._worker is None)
                    job = widget.jobs.records[0]
                    self.assertTrue(Path(job['report']).exists())
                    self.assertEqual(widget.result_table.item(0, 1).text(), '耐用性')
                    self.assertEqual(widget._report_path, job['report'])
                    self.assertIn('1/1', Path(job['log']).read_text(encoding='utf-8'))
            finally:
                if widget._worker:
                    widget._worker.wait()
                support.cleanup_qt_widgets()

    def test_summary_merge_preserves_row_ratings(self):
        # Importing the script requires its configuration environment; execute
        # only the pure merge function with its documented dependencies.
        import ast
        import pandas as pd
        path = Path(__file__).resolve().parents[1] / 'scripts/comments_step_2.py'
        module = ast.parse(path.read_text(encoding='utf-8'))
        function = next(node for node in module.body if isinstance(node, ast.FunctionDef) and node.name == 'merge_sheets')
        scope = {'pd': pd, 'List': list, 'SKIP_FAILED_ROWS': True}
        exec(compile(ast.Module(body=[function], type_ignores=[]), str(path), 'exec'), scope)
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'reviews.xlsx'
            book = Workbook()
            book.active.title = '评论'
            book.active.append(['内容', '星级'])
            book.active.append(['Good', 5])
            book.active.append(['Bad', 1])
            sheet = book.create_sheet('2 star')
            sheet.append(['内容'])
            sheet.append(['Poor'])
            book.save(source)
            result = scope['merge_sheets'](str(source), [])
            self.assertEqual(result['星级'].tolist(), [5, 1, '2 star'])

    def _wait(self, app, predicate):
        deadline = time.monotonic() + 10
        while not predicate() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.01)
        app.processEvents()
        self.assertTrue(predicate())


if __name__ == '__main__':
    unittest.main()
